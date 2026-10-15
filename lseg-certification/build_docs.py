# -*- coding: utf-8 -*-
"""
Génère les 2 documents de préparation à la Certification LSEG Workspace :
  1) REVISION  — fiche de révision complète (français)
  2) MOCK EXAM — examen blanc 40 questions (questions EN, corrigé FR)
Sorties : PDF (mise en page soignée) + DOCX (éditable), même contenu.
Usage : python3 build_docs.py
"""
import re
from fpdf import FPDF
from fpdf.enums import XPos, YPos, MethodReturnValue, TableBordersLayout, TableCellFillMode
from fpdf.fonts import FontFace
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------------------------------------------------------------- couleurs
NAVY   = (16, 42, 78)
PURPLE = (108, 62, 234)
PUR_BG = (240, 236, 253)
AMB_BG = (255, 247, 224)
GRN_BG = (233, 246, 238)
GRAY   = (70, 70, 70)
LIGHT  = (150, 150, 150)
GREEN  = (20, 118, 66)
RED    = (168, 32, 32)
ROW_BG = (245, 247, 251)

USABLE = 186  # largeur utile A4 - marges 12/12
DJV  = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
DJVB = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"

# ================================================================ PDF
class LSEGPdf(FPDF):
    def __init__(self, title, tag):
        super().__init__(orientation="P", unit="mm", format="A4")
        self.doc_title, self.doc_tag = title, tag
        self.set_auto_page_break(auto=True, margin=16)
        self.set_margins(12, 15, 12)
        self.add_font("DejaVu", "", DJV)
        self.add_font("DejaVu", "B", DJVB)

    def header(self):
        if self.page_no() > 1:
            self.set_y(7)
            self.set_font("DejaVu", "B", 7.5)
            self.set_text_color(*NAVY)
            self.cell(0, 4, self.doc_title, new_x=XPos.RIGHT, new_y=YPos.TOP)
            self.set_text_color(*PURPLE)
            self.cell(0, 4, self.doc_tag, align="R", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            self.set_draw_color(*PURPLE)
            self.set_line_width(0.35)
            self.line(12, 12.2, 198, 12.2)
            self.set_y(15)

    def footer(self):
        self.set_y(-11)
        self.set_font("DejaVu", "", 7.5)
        self.set_text_color(*LIGHT)
        self.cell(0, 5, f"{self.doc_tag}  ·  page {self.page_no()}/{{nb}}", align="C")

    def cover(self, badge, title, subtitle, meta):
        self.add_page()
        self.set_fill_color(*NAVY)
        self.rect(0, 0, 210, 56, style="F")
        self.set_fill_color(*PURPLE)
        self.rect(0, 56, 210, 1.6, style="F")
        self.set_xy(14, 12)
        self.set_font("DejaVu", "B", 9)
        self.set_text_color(178, 160, 255)
        self.cell(0, 6, badge, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        self.set_xy(14, 20)
        self.set_font("DejaVu", "B", 21)
        self.set_text_color(255, 255, 255)
        self.multi_cell(182, 9.6, title, align="L", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        self.set_xy(14, self.get_y() + 2)
        self.set_font("DejaVu", "", 10.5)
        self.set_text_color(205, 198, 245)
        self.multi_cell(180, 5.4, subtitle, align="L", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        self.set_y(63)
        self.set_font("DejaVu", "", 8.3)
        self.set_text_color(*GRAY)
        self.multi_cell(USABLE, 4.6, meta)
        self.ln(2)

    def _need(self, h):
        if self.will_page_break(h):
            self.add_page()

    def h1(self, text):
        self._need(20)
        self.ln(2.5)
        y = self.get_y()
        self.set_fill_color(*PURPLE)
        self.rect(12, y + 0.6, 2.4, 7.2, style="F")
        self.set_xy(17.5, y)
        self.set_font("DejaVu", "B", 13.5)
        self.set_text_color(*NAVY)
        self.multi_cell(USABLE - 6, 7.2, text, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        self.ln(1.2)

    def h2(self, text):
        self._need(14)
        self.ln(1.6)
        self.set_font("DejaVu", "B", 11.2)
        self.set_text_color(*PURPLE)
        self.multi_cell(USABLE, 5.8, text, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        self.ln(0.6)

    def p(self, text):
        self.set_font("DejaVu", "", 9.3)
        self.set_text_color(*GRAY)
        self.multi_cell(USABLE, 5.0, text, markdown=True, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        self.ln(1.2)

    def bullets(self, items):
        for it in items:
            self._need(10)
            y = self.get_y()
            self.set_xy(14, y)
            self.set_font("DejaVu", "B", 9.6)
            self.set_text_color(*PURPLE)
            self.cell(5.5, 5.1, "▪", new_x=XPos.RIGHT, new_y=YPos.TOP)
            self.set_x(20)
            self.set_font("DejaVu", "", 9.3)
            self.set_text_color(*GRAY)
            self.multi_cell(USABLE - 8, 5.1, it, markdown=True,
                            new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        self.ln(1.2)

    def box(self, title, body, bg=PUR_BG, accent=PURPLE, tcol=NAVY):
        pad = 3.4
        w = USABLE
        if title:
            self.set_font("DejaVu", "B", 9.6)
            th = self.multi_cell(w - 2 * pad, 5.1, title, dry_run=True,
                                 output=MethodReturnValue.HEIGHT)
        else:
            th = 0
        self.set_font("DejaVu", "", 9.2)
        bh = self.multi_cell(w - 2 * pad, 5.1, body, markdown=True, dry_run=True,
                             output=MethodReturnValue.HEIGHT)
        H = th + bh + 2 * pad
        self._need(H + 4)
        y0 = self.get_y()
        self.set_fill_color(*bg)
        self.set_draw_color(*accent)
        self.set_line_width(0.3)
        self.rect(12, y0, w, H, style="DF", round_corners=True, corner_radius=2)
        if title:
            self.set_xy(12 + pad, y0 + pad)
            self.set_font("DejaVu", "B", 9.6)
            self.set_text_color(*tcol)
            self.multi_cell(w - 2 * pad, 5.1, title, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        self.set_xy(12 + pad, y0 + pad + th)
        self.set_font("DejaVu", "", 9.2)
        self.set_text_color(*GRAY)
        self.multi_cell(w - 2 * pad, 5.1, body, markdown=True,
                        new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        self.set_y(y0 + H + 3)

    def tbl(self, headers, rows, widths=None, fs=8.7):
        self.set_font("DejaVu", "", fs)
        head = FontFace(family="DejaVu", emphasis="B", color=(255, 255, 255),
                        fill_color=NAVY)
        with self.table(col_widths=widths, text_align="LEFT", line_height=4.7,
                        borders_layout=TableBordersLayout.ALL, headings_style=head,
                        cell_fill_color=ROW_BG, cell_fill_mode=TableCellFillMode.ROWS,
                        padding=1.3, markdown=True) as tb:
            if headers:
                r = tb.row()
                for c in headers:
                    r.cell(c)
            for row in rows:
                r = tb.row()
                for c in row:
                    r.cell(c)
        self.ln(2.5)

    def q(self, n, text, options):
        self._need(22)
        self.ln(1.0)
        self.set_font("DejaVu", "B", 9.5)
        self.set_text_color(*NAVY)
        self.multi_cell(USABLE, 5.2, f"Q{n}.  {text}", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        self.ln(0.4)
        for opt in options:
            self._need(9)
            self.set_xy(20, self.get_y())
            self.set_font("DejaVu", "", 9.2)
            self.set_text_color(*GRAY)
            self.multi_cell(USABLE - 8, 5.0, "☐   " + opt,
                            new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        self.ln(1.4)

    def ans(self, label, text, warn=False):
        self._need(10)
        y = self.get_y()
        self.set_xy(14, y)
        self.set_font("DejaVu", "B", 9.8)
        self.set_text_color(*GREEN)
        self.cell(6, 5.0, "✓", new_x=XPos.RIGHT, new_y=YPos.TOP)
        self.set_x(20.5)
        self.set_font("DejaVu", "", 9.2)
        self.set_text_color(*GRAY)
        self.multi_cell(USABLE - 9, 5.0, f"**{label}**  —  {text}", markdown=True,
                        new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        if warn:
            self.set_xy(20.5, self.get_y())
            self.set_font("DejaVu", "B", 8.6)
            self.set_text_color(150, 100, 0)
            self.multi_cell(USABLE - 9, 4.4, "⚠ Réponse à haute probabilité (pool candidats) : "
                            "croise-la avec la fiche si le libellé diffère le jour J.",
                            new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        self.ln(1.0)

    def pagebreak(self):
        self.add_page()


def render_pdf(path, title, tag, badge, stitle, meta, blocks):
    pdf = LSEGPdf(title, tag)
    pdf.cover(badge, title, stitle, meta)
    for b in blocks:
        kind = b[0]
        if kind == "h1":
            pdf.h1(b[1])
        elif kind == "h2":
            pdf.h2(b[1])
        elif kind == "p":
            pdf.p(b[1])
        elif kind == "bullets":
            pdf.bullets(b[1])
        elif kind == "box":
            pdf.box(b[1], b[2], *b[3:])
        elif kind == "table":
            pdf.tbl(b[1], b[2], b[3] if len(b) > 3 else None)
        elif kind == "q":
            pdf.q(b[1], b[2], b[3])
        elif kind == "ans":
            pdf.ans(b[1], b[2], len(b) > 3 and b[3])
        elif kind == "pagebreak":
            pdf.pagebreak()
    pdf.output(path)
    print("PDF écrit :", path)


# ================================================================ DOCX
def _shade(cell, hexfill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexfill)
    tcPr.append(shd)


def _runs(par, text, bold=False, color=None, size=None):
    for part in re.split(r"(\*\*.+?\*\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2])
            r.bold = True
        else:
            r = par.add_run(part)
            r.bold = bold
        if color:
            r.font.color.rgb = color
        if size:
            r.font.size = size


def render_docx(path, title, tag, badge, stitle, meta, blocks):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)
    for sec in doc.sections:
        sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.6)
        sec.left_margin, sec.right_margin = Cm(1.8), Cm(1.8)

    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.rows[0].cells[0]
    _shade(cell, "102A4E")
    p0 = cell.paragraphs[0]
    r = p0.add_run(badge + "\n")
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0xB2, 0xA0, 0xFF)
    r2 = p0.add_run(title)
    r2.bold = True
    r2.font.size = Pt(20)
    r2.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p1 = cell.add_paragraph()
    r3 = p1.add_run(stitle)
    r3.font.size = Pt(10)
    r3.font.color.rgb = RGBColor(0xCD, 0xC6, 0xF5)
    doc.add_paragraph(meta)

    navy = RGBColor(0x10, 0x2A, 0x4E)
    purple = RGBColor(0x6C, 0x3E, 0xEA)
    gray = RGBColor(0x46, 0x46, 0x46)
    green = RGBColor(0x14, 0x76, 0x42)

    for b in blocks:
        kind = b[0]
        if kind == "h1":
            par = doc.add_paragraph()
            par.space_before = Pt(10)
            _runs(par, b[1], bold=True, color=navy, size=Pt(14))
        elif kind == "h2":
            par = doc.add_paragraph()
            _runs(par, b[1], bold=True, color=purple, size=Pt(11.5))
        elif kind == "p":
            par = doc.add_paragraph()
            _runs(par, b[1], color=gray)
        elif kind == "bullets":
            for it in b[1]:
                par = doc.add_paragraph(style="List Bullet")
                _runs(par, it, color=gray)
        elif kind == "box":
            tb = doc.add_table(rows=1, cols=1)
            tb.style = "Table Grid"
            c = tb.rows[0].cells[0]
            _shade(c, "F0ECFD")
            pt = c.paragraphs[0]
            _runs(pt, b[1], bold=True, color=navy)
            pb = c.add_paragraph()
            _runs(pb, b[2], color=gray)
            doc.add_paragraph()
        elif kind == "table":
            headers, rows = b[1], b[2]
            tb = doc.add_table(rows=1 + len(rows), cols=len(headers) if headers else len(rows[0]))
            tb.style = "Table Grid"
            ri = 0
            if headers:
                for ci, htxt in enumerate(headers):
                    c = tb.rows[0].cells[ci]
                    _shade(c, "102A4E")
                    _runs(c.paragraphs[0], htxt, bold=True,
                          color=RGBColor(0xFF, 0xFF, 0xFF), size=Pt(9))
                ri = 1
            for row in rows:
                for ci, ctxt in enumerate(row):
                    _runs(tb.rows[ri].cells[ci].paragraphs[0], ctxt,
                          color=gray, size=Pt(9))
                ri += 1
            doc.add_paragraph()
        elif kind == "q":
            par = doc.add_paragraph()
            par.space_before = Pt(8)
            _runs(par, f"Q{b[1]}.  {b[2]}", bold=True, color=navy)
            for opt in b[3]:
                po = doc.add_paragraph("☐   " + opt)
                po.paragraph_format.left_indent = Cm(0.7)
                for r_ in po.runs:
                    r_.font.color.rgb = gray
        elif kind == "ans":
            par = doc.add_paragraph()
            r_ = par.add_run("✓  ")
            r_.bold = True
            r_.font.color.rgb = green
            _runs(par, b[1], bold=True, color=navy)
            rr = par.add_run("  —  ")
            _runs(par, b[2], color=gray)
            if len(b) > 3 and b[3]:
                pw = doc.add_paragraph("⚠ Réponse à haute probabilité (pool candidats) : "
                                       "croise-la avec la fiche si le libellé diffère le jour J.")
                for r__ in pw.runs:
                    r__.font.color.rgb = RGBColor(0x96, 0x64, 0x00)
                    r__.font.size = Pt(8.5)
        elif kind == "pagebreak":
            doc.add_page_break()
    doc.save(path)
    print("DOCX écrit :", path)


# ================================================================ CONTENU 1 : FICHE DE RÉVISION
REV_META = ("Source : modules officiels du LSEG Learning Centre + « Certification sample questions — LSEG Workspace » "
            "(PDF officiel, mise à jour 11/24) + questions remontées par les candidats. "
            "Préparation 100 % hors plateforme : aucune installation Excel nécessaire.")

REV_BLOCKS = [
    ("box", "L'EXAMEN EN BREF — à connaître par cœur",
     "•  **30 questions** à choix multiples · **60 minutes** · **≥ 80 %** pour valider (maximum 6 erreurs) · **2 tentatives** (1 rattrapage, pas besoin d'être immédiat) · résultat **immédiat** · certificat **PDF** à partager (LinkedIn/CV).\n"
     "•  **Accès** : LSEG Learning Centre (compte avec e-mail pro/académique + code OTP) → Certification → LSEG Workspace → complète les **modules eLearning + quiz** (prérequis) → bouton **« Access certification »** → e-mail avec **lien + token** → 60 min une fois lancé.\n"
     "•  **Modules** : Getting Started · News · Charting · Company Analysis · Industry Analysis · LSEG Data in Excel (+ Datastream intégré). Gratuit.\n"
     "•  **Stratégie** : tout est « produit » (chemins de menus, comportements, apps) — aucune formule à calculer. D'où l'intérêt d'une fiche : chaque ligne ci-dessous = une réponse possible le jour J."),

    ("h1", "MODULE 1 [M1] — Getting Started : navigation Workspace"),
    ("p", "Workspace = suite d'**apps flottantes** organisées en **workspaces** (onglets). L'examen adore les chemins exacts et les comportements par défaut : mémorise-les."),
    ("bullets", [
     "**Search box** (en haut) : propulsée par **Autosuggest** — trouve des apps, instruments, personnes, et même des **réponses (Answers)**. Astuce officielle : taper « **Nom + OV** » ouvre la **Company Overview** (ex. « Apple OV ») ; « TOPNEWS » = Top News ; « NEWSD » = News Discovery ; « INDUS » = Industry Overviews.",
     "**App Library** : catalogue de toutes les apps par catégorie/asset class. Toute app ouverte (menu, Library ou Search) rejoint **automatiquement le workspace actif**.",
     "**Workspaces** : jusqu'à **10 workspaces ouverts en même temps** ; onglet = clic droit → **Save** pour sauvegarder/partager. Une app ouverte s'ajoute au workspace actif.",
     "**Active symbol** : une app nouvellement ouverte affiche par défaut **l'instrument actif** = le **dernier instrument consulté** (pas le 1er de ta liste !).",
     "**Apps** : « **snap** » (attacher/déplacer ensemble comme un bloc), « **stack** » (empiler en onglets), redimensionner librement → **VRAI** : on peut attacher les apps pour les déplacer d'un bloc et les empiler en onglets.",
     "**Linking** : chaque app a un **Link Channel** = petit **cercle coloré en haut à gauche**. Apps de **même couleur** = même instrument partagé (tu tapes un code dans l'une, toutes suivent).",
     "**Règle d'or (question Vrai/Faux officielle → VRAI)** : une **nouvelle app groupée avec des apps déjà liées devient automatiquement liée** (elle rejoint la couleur du groupe).",
     "**Aide / FAQs** : menu **« ? » → Get Help & Support**, via **Autosuggest** dans la Search box, et via l'appli **MyRefinitiv** → question type : « **All of these options** ».",
     "**My Briefcase** : icône **en haut à droite de la toolbar** ; collecte news/documents/research pour lecture **offline** (desktop ou **mobile**, synchronisé).",
    ]),
    ("box", "✓ À RETENIR [M1] — News Monitor : les 2 accès officiels",
     "1) Taper « **News** » dans la **Search box** → choisir le **News Component** dans l'**Autosuggest**.  2) **Menu > News & Research**.\n"
     "**Pièges officiels** : taper « **Monitor** » ne renvoie PAS le composant News · **F2 n'ouvre PAS les news** (c'est la Quote List) · « Menu > Search Tools » est un autre menu."),
    ("table", ["Touche", "Fonction", "Touche", "Fonction"],
     [["**F2**", "Quote List (liste de cotations)", "**F5**", "Refresh (rafraîchir la page)"],
      ["**F4**", "Quote (cotation)", "**F6**", "Max/Min (réduire/agrandir l'app)"],
      ["**F9**", "News (ouvrir les news)", "", ""]],
     [18, 75, 18, 75]),

    ("h1", "MODULE 2 [M2] — News & Briefcase"),
    ("p", "News Monitor = app de news ultra-personnalisable : **Reuters exclusif** + milliers de sources premium (WSJ, MarketWatch…), couverture ~**48 langues**."),
    ("bullets", [
     "**Recherche/filtres News Monitor** : **companies**, **portfolios / lists / monitors**, **langues**, sources, marchés, secteurs, régions, topics, classifications business → la réponse officielle est « **All of these options** ».",
     "**Filtres avancés + exclusions** : sélectionner une source, un marché… et **exclure** explicitement un critère (ex. exclure une source).",
     "**Expressions booléennes** (AND / OR / NOT) pour des résultats très ciblés ; les mots-clés sont **surlignés** dans les résultats (highlighter) ; filtrage complémentaire avec le highlight tool.",
     "**Timeline toggle** : bascule entre ordre **chronologique inverse** et critères de **date/heure choisis**.",
     "**News filter réutilisable** : dans News Monitor → **clic droit → File > Favorites > Add to Favorites** (on le retrouvera plus tard).",
     "**My Briefcase** : « **Add to Briefcase** » au survol d'un article (icône en haut à droite de l'article) → lecture **offline** desktop/mobile. Limite officielle : **100 documents maximum ajoutés en même temps** (pas 25, pas 50, pas 200).",
     "**Codes utiles** : NEWS (News Monitor) · TOPNEWS (Top News) · NEWSD (News Discovery). Alertes configurables sur mots-clés, instruments, sociétés.",
    ]),
    ("box", "✓ À RETENIR [M2]",
     "News Monitor searchable = **All of the above** · Briefcase = **100 docs** · touche news = **F9** (F2 = Quote List) · filtre réutilisable = **clic droit > File > Favorites > Add to Favorites** · Briefcase = lecture offline **desktop + mobile**."),

    ("h1", "MODULE 3 [M3] — Charting"),
    ("bullets", [
     "**Chart app** : tous types de graphiques (ligne, chandelier, barres), périodes et **fréquences** multiples.",
     "**Studies/indicateurs** : moyennes mobiles, RSI, MACD, bandes de Bollinger… pour identifier **tendances** et signaux techniques.",
     "**Comparaison** : superposer plusieurs instruments/indices (normalisation en %).",
     "**Templates & styles** : sauvegarder la mise en forme du chart pour la réutiliser.",
     "**Split mode** : diviser l'app chart en **plusieurs vues** (ex. même instrument par **intervalle** différent : 1J, 1M, 1A).",
     "**Link Channel** : le chart suit l'instrument actif des apps liées — tu scrolles ton **Monitor**, le chart se met à jour tout seul.",
     "**Export** : image / Excel / PowerPoint pour les présentations.",
    ]),

    ("h1", "MODULE 4 [M4] — Company Analysis"),
    ("bullets", [
     "**Company Overview** : taper « **Nom + OV** » (ex. « Airbus OV ») ou le nom dans la Search box → page résumé : prix, profil, financials, estimates, news, research.",
     "**Apps clés** : Quote Overview · **Financials** (comptes **WorldScope**) · **Estimates** (consensus **I/B/E/S**) · **Ownership** (actionnariat) · **Debt** · **ESG** (scores **Asset4**) · Peers/Compare · Filings & Research.",
     "**Peers (concurrents)** — question officielle type « All of these » : déterminés par le **chevauchement de couverture analystes**, la **classification business (TRBC/GICS)** et les **concurrents que la société déclare dans ses filings**.",
     "**Screener** : filtrer tout l'univers d'instruments (place, secteur, market cap, ratios…) et sauvegarder en **listes** réutilisables (Monitor, portfolios, Excel).",
     "**Advanced Research Search** : la **Quick Search** filtre par **Report Type, Contributor et Industry** → « **All of these options** ».",
    ]),

    ("h1", "MODULE 5 [M5] — Industry Analysis"),
    ("bullets", [
     "**Industry Overviews** (code **INDUS**), **Country Overviews**, **Market Monitors** : analyse **top-down** secteur/pays, info market-moving, comparaisons pour l'allocation.",
     "**Aggregates** : agrège les fondamentaux d'un **groupe de sociétés** (somme/moyenne de cap, P/E, croissance…). Classifications supportées : **GICS ET TRBC** (pas SIC) → réponse officielle « **GICS AND TRBC** ».",
     "**TRBC** (LSEG Business Classifications) : 5 niveaux de granularité, 250 000+ sociétés, 130 pays — plus fin que GICS sur certaines niches.",
     "**Groupes personnalisés** : l'app Aggregates permet de créer ses propres groupes, dont des **tranches de market cap** → VRAI (pool candidats).",
     "**Portfolios & lists (Monitor)** : comparer secteurs, géographies et pondérations d'un portefeuille.",
     "**SUSFIN — Sustainable Finance Landing Page** : research & white papers + section « **What's New** » + liste complète des apps/vues durables → « **All of these answers** ».",
     "Datastream complète : **9 000 charts pré-construits**, séries sectorielles et économiques.",
    ]),

    ("h1", "MODULE 6 [M6] — LSEG Data in Excel & Datastream (LE module à pièges)"),
    ("p", "Deux briques : l'**add-in Excel** (ruban LSEG/Refinitiv) et **Datastream for Office**. Comme tu ne peux pas t'entraîner sur Excel, apprends ces comportements par cœur — ils tombent tels quels."),
    ("h2", "6.1 — Add-in Excel & Formula Builder"),
    ("bullets", [
     "**Formula Builder** : assistant graphique pour construire des formules sans coder. Récupère **temps réel + fondamentale + historique** → réponse officielle « **All of these options** ».",
     "**Fonctions clés** : **=TR** (données fondamentales & estimates, statiques ou historiques) · **=RHistory** (séries temporelles) · **=RSearch** (recherche d'instruments).",
     "**Templates pré-construits** : menu **LSEG/Refinitiv > Templates** (pas dans une cellule !).",
     "**Template Library** : bibliothèque pour **partager tes fichiers Excel avec les AUTRES utilisateurs**.",
     "**My Files** : tes **propres** fichiers, **accessibles de partout** (cloud perso).",
     "**⚠ PIÈGE OFFICIEL : la fonction « TEMP » n'existe pas.** On ne peut PAS appeler un template depuis une cellule avec =TEMP — c'est l'énoncé FAUX de la question officielle.",
    ]),
    ("h2", "6.2 — Datastream : requests, Navigator, listes"),
    ("bullets", [
     "**Base Datastream** : **120 ans** d'historique · **175 pays** · **35 M+** instruments/indicateurs (dont 8,5 M d'indicateurs éco). Contenus exclusifs : **WorldScope** (comptes), **Asset4** (ESG), **I/B/E/S** (estimates), Reuters Polls.",
     "**Static Request** = photo à un instant T (latest value). **Time Series Request** = historique → exige en plus **dates + fréquence** (D/W/M/Q/Y).",
     "**Plusieurs séries en Static** : taper les codes **séparés par des VIRGULES** ou utiliser un **code de liste** (ex. L#…, LN…) → réponse officielle. **Pas de « + »**, et on **PEUT** taper les codes directement (Navigator optionnel).",
     "**Navigator** : outil de recherche des **séries et datatypes (mnemonics)** ; une multi-sélection s'insère **automatiquement séparée par des virgules**.",
     "**Listes en static** : **1 à 10 datatypes** retournés par série de la liste.",
     "**Request Table** : batch de requêtes rafraîchissables — types : **TS** (time series), **S** (static), **TSL** (time series sur liste), **L** (liste), **CH** (chart).",
     "**Requête embarquée (embedded) dans Excel** — bénéfice officiel : les données **se mettent à jour dynamiquement** ensuite. Les autres propositions (« jamais modifiable », « plus de données », « charts uniquement en embedded ») sont fausses.",
     "**DSGRID** : formule grille multi-séries/multi-datatypes — séries et datatypes **séparés par des virgules** (ex. =DSGRID(\"MKS\",\"EPS\",…)).",
    ]),
    ("h2", "6.3 — Library, Collections & datatypes Equities"),
    ("bullets", [
     "**Ajouter des charts de la Library à une collection — LES 2 BONNES FAÇONS (question « which two »)** : ① **une seule chart** → ouvrir le **flyout** à côté de la chart > Collections > **Add to collection** ; ② **plusieurs charts** → cocher les **cases à cocher** puis icône **Add to collection** de la toolbar.",
     "**Pièges (faux)** : pas de **clic droit sur la chart > Add to collection** ; pas de **clic droit sur un dossier > Export > To collection**.",
     "**Datatypes EQUITIES — disponibles** : **WorldScope** (comptes) ✓ · **Asset4 ESG** ✓ · **I/B/E/S estimates** ✓. **NON disponible : « Attribution data types »** → réponse officielle à « Which of the following data is NOT available for Equities? ».",
     "**API / Codebook** : accès programmatique (Python/R, LSEG Data Library) aux mêmes données pour des workflows custom — cité dans le module « Data Retrieval : Office & API ».",
    ]),

    ("h1", "CHIFFRES & RACCOURCIS À CONNAÎTRE PAR CŒUR"),
    ("table", ["Élément", "Valeur exacte"],
     [["Format de l'examen", "**30 questions · 60 minutes · ≥ 80 % · 2 tentatives**"],
      ["Briefcase", "**100 documents maximum** ajoutés en même temps"],
      ["Workspaces ouverts simultanément", "**10**"],
      ["F2 / F4 / F5 / F6 / F9", "Quote List / Quote / Refresh / Max-Min / **News**"],
      ["Accès News Monitor", "« News » + Autosuggest **OU** Menu > News & Research"],
      ["Multi-séries (Static Request)", "**Virgule** entre les codes, ou **code de liste** (jamais « + »)"],
      ["Datatypes par série (liste en static)", "**1 à 10**"],
      ["Aggregates — classifications", "**GICS ET TRBC** (pas SIC)"],
      ["Datastream", "**120 ans · 175 pays · 35 M+ séries · 9 000 charts pré-construits**"],
      ["TRBC", "**5 niveaux** de granularité, 250 000+ sociétés"],
      ["Équities NON disponible", "**Attribution data types** (WorldScope/Asset4/IBES : dispo)"],
      ["Embedded Excel request", "**Mise à jour dynamique** des données"]],
     [78, 108]),

    ("h1", "LES 12 PIÈGES — CHECKLIST ANTI-ERREUR"),
    ("bullets", [
     "**« Which two… » → cocher EXACTEMENT 2 cases** (ex. accès News Monitor ; charts → collection). Une seule case = 0 point.",
     "**« …is false? » → chercher l'intrus fabriqué** : la **fonction TEMP** n'existe pas dans l'add-in Excel. Les 3 autres énoncés (Template Library / menu Templates / My Files) sont vrais.",
     "**Vrai/Faux linking → VRAI** : app groupée avec des apps liées = **liaison automatique**.",
     "**F2 ≠ News** : F2 = Quote List. La touche news = **F9** ; F5 = Refresh ; F6 = Max/Min.",
     "**Taper « Monitor » ≠ News Monitor** : il faut taper « **News** » et prendre le News Component dans l'Autosuggest.",
     "**Briefcase = 100** (les distracteurs 25/50/200 sont là pour toi).",
     "**Virgule, jamais « + »** pour plusieurs séries — et on **peut** taper les codes sans passer par Navigator.",
     "**Static ≠ Time Series** : la static est une photo ; la time series demande **dates + fréquence**. « Plusieurs séries impossible en static » = FAUX.",
     "**Template Library (partager avec les autres) ≠ My Files (tes fichiers, partout)** — ne les inverse pas.",
     "**Quand 3 propositions sont plausibles → « All of the above/these »** : c'est LA réponse officielle sur News Monitor (recherche), Formula Builder, SUSFIN, FAQs, peers, Quick Search. Mais vérifie chaque item avant de cocher.",
     "**Équities** : WorldScope/Asset4/I/B/E/S disponibles ; **Attribution data types** non.",
     "**Embedded request** = mise à jour dynamique — les 3 autres « bénéfices » sont des anti-bénéfices.",
    ]),

    ("pagebreak",),
    ("h1", "CORRIGÉ OFFICIEL COMMENTÉ — sample questions LSEG (maj 11/24)"),
    ("p", "Les questions officielles d'entraînement publiées par LSEG, avec la bonne réponse et le pourquoi. Le pool réel les reprend à l'identique ou presque : vise **0 erreur** ici avant de réserver ton créneau."),
    ("h2", "Partie « LSEG Workspace »"),
    ("q", 1, "The Aggregates app supports which business classifications schema?",
     ["Global Industry Classification Standard (GICS) only",
      "LSEG Business Classifications (TRBC) only",
      "Standard Industrial Classification (SIC) only",
      "Global Industry Classification Standard (GICS) AND LSEG Business Classifications (TRBC)"]),
    ("ans", "Réponse : GICS AND TRBC", "l'app Agrégats groupe par les schémas **GICS et TRBC** (jamais par SIC). [M5]"),
    ("q", 2, "Which of the following can be searched for in the News Monitor App?",
     ["Companies, Portfolios, & Lists", "Languages, News Source, and Geography",
      "Markets, Topics, and Business Classification", "All of these options"]),
    ("ans", "Réponse : All of these options", "la recherche News Monitor couvre sociétés, portefeuilles/listes, langues, sources, géographie, marchés, topics, classifications. [M2]"),
    ("q", 3, "In Excel, what type of data can be retrieved using Formula Builder?",
     ["Real time data", "Fundamental data", "Historical data", "All of these options"]),
    ("ans", "Réponse : All of these options", "temps réel + fondamentale + historique. [M6]"),
    ("q", 4, "In which two ways can you access the News Monitor app?",
     ["Type the word 'News' into the Search box and select the News Component from Autosuggest",
      "Menu > News & Research", "Press the 'F2' key", "Menu > Search Tools",
      "Type the word 'Monitor' into the Search box and select the News Component from Autosuggest"]),
    ("ans", "Réponses (2) : « News » + Autosuggest  ET  Menu > News & Research",
     "F2 = Quote List (la touche news = F9) ; « Monitor » ne renvoie pas le composant ; Search Tools est un autre menu. [M1/M2]"),
    ("q", 5, "Which of the following statements regarding the Excel add-in is false?",
     ["You can share your Excel files with other users by uploading them to the template library",
      "You can access pre-built templates by clicking 'Templates' in the LSEG menu",
      "You can access pre-built templates from within a cell using the TEMP function",
      "You can upload your own Excel files to 'My Files' so you can access them from anywhere"]),
    ("ans", "Réponse : la TEMP function", "elle n'existe pas — les 3 autres sont vraies. [M6]"),
    ("q", 6, "True or false — If a new app is grouped with apps that are already linked together, the new app will be automatically linked.",
     ["True", "False"]),
    ("ans", "Réponse : TRUE", "le groupement (snap) propage la couleur du Link Channel à la nouvelle app. [M1]"),
    ("h2", "Partie « Datastream »"),
    ("q", 7, "Which two of the following are ways to add a chart or charts from the library to a collection?",
     ["To add a single chart: open the flyout next to the chart, go to 'Collections', select 'Add to collection'",
      "To add multiple charts: click the checkboxes next to the charts, click the 'Add to collection' icon on the toolbar",
      "Right-click on a chart and select 'Add to collection'",
      "Right-click on a folder in the Library > Export > To collection"]),
    ("ans", "Réponses (2) : le flyout (chart unique)  ET  cases à cocher + icône toolbar (plusieurs)",
     "pas de clic droit sur la chart, pas d'Export dossier → collection. [M6]"),
    ("q", 8, "Which of the following is a benefit of embedding a Datastream for Office – Excel request?",
     ["The data can be dynamically updated going forward", "The data request can never be changed",
      "An embedded request allows a user to download more data than a manual request",
      "An embedded request gives the option to download data in chart format whilst a manual request does NOT"]),
    ("ans", "Réponse : The data can be dynamically updated going forward", "seul vrai bénéfice listé. [M6]"),
    ("q", 9, "How can you download data for more than one series in a static request?",
     ["In the Series/List field, type the codes for all required series, separated by a comma, or use a list code",
      "In the Series/List field, type the codes for all required series, separated by a plus sign",
      "It is NOT possible to type in the codes, you must search for them in Navigator",
      "A static request does not allow you to download data for more than one series"]),
    ("ans", "Réponse : codes séparés par une virgule, ou code de liste",
     "le « + » n'existe pas, et les codes se tapent librement (Navigator optionnel). [M6]"),
    ("q", 10, "Which of the following data is NOT available for equities?",
     ["Attribution data types", "Asset4 environmental, social and governance (ESG) data",
      "WorldScope Company Accounts data", "I/B/E/S estimates"]),
    ("ans", "Réponse : Attribution data types", "WorldScope, Asset4 ESG et I/B/E/S sont disponibles pour les équities. [M6]"),
    ("h2", "Pool réel complémentaire (remontées de candidats)"),
    ("ans", "Briefcase — max de documents ajoutés en même temps : 100", "options officielles 25/50/100/200 → **100**. [M2]", True),
    ("ans", "« Quick Search » de l'Advanced Research Search : All of these options", "filtre **Report Type + Contributor + Industry**. [M4]"),
    ("ans", "Peers d'une société : All of these answers are correct", "**couverture analystes + classification (TRBC/GICS) + concurrents déclarés dans les filings**. [M4]"),
    ("ans", "SUSFIN (Sustainable Finance Landing Page) : All of these answers are correct", "research/white papers + « What's New » + liste des apps durables. [M5]"),
    ("ans", "Accès aux FAQs Workspace : All of these options", "menu « ? » Get Help & Support + Autosuggest + appli MyRefinitiv. [M1]"),
    ("ans", "Ce qu'on peut changer dans News Monitor : Story layout", "le style/graisse de police n'est pas paramétrable dans l'app ; le **story layout** (densité d'affichage) l'est. [M2]", True),
    ("ans", "Groupes de market cap personnalisés dans Aggregates : TRUE", "l'app accepte des groupes custom, dont des tranches de capitalisation. [M5]", True),

    ("h1", "PLAN D'ATTAQUE 5 JOURS"),
    ("box", "Objectif : arrivé le jour J avec les 30 réponses « déjà vues »",
     "**J1-J2** — complète les modules eLearning + quiz du Learning Centre (obligatoire pour débloquer l'examen) en prenant des notes sur CE document.\n"
     "**J3** — relis cette fiche 2 fois ; récite les boxes « ✓ À RETENIR » et le tableau des chiffres.\n"
     "**J4** — MOCK EXAM (l'autre document) en conditions : 45 min, sans notes, puis corrigé ligne à ligne.\n"
     "**J5 (veille)** — relecture : tableau « Chiffres & raccourcis » + « Les 12 pièges » + corrigé officiel. Puis réserve ton créneau (token valable une fois lancé : 60 min chrono).\n"
     "**Règle d'or** : lis DEUX fois chaque question (« which two », « false »), élimine, puis coche."),
]

# ================================================================ CONTENU 2 : MOCK EXAM
MOCK_META = ("Questions en anglais = format exact du jour J. Corrigé détaillé en français avec renvois [M1]-[M6] vers la fiche de révision. "
             "Les questions marquées ★ reprennent (quasi) à l'identique les sample questions officielles ou le pool candidats : points quasi garantis.")

MOCK_BLOCKS = [
    ("box", "INSTRUCTIONS — simulation en conditions réelles",
     "•  **40 QCM** · 1 seule bonne réponse SAUF mention « Which two » (coche-en exactement 2) · durée conseillée : **45 minutes** sans notes (le vrai examen : 30 Q / 60 min, ~2 min/question).\n"
     "•  Le seuil de réussite réel est **80 %** (24/30) — mais tu prépares 30/30, donc vise **≥ 38/40** ici.\n"
     "•  Note tes réponses dans la grille page suivante, puis corrige avec le corrigé détaillé. Chaque erreur → relis le module [Mx] indiqué.\n"
     "•  ★ = question issue des sample questions officielles ou du pool remonté par les candidats."),

    ("h2", "SECTION A — Workspace & Navigation [M1]"),
    ("q", 1, "You open a new app from the App Library. What will it display by default?",
     ["The first instrument in your Monitor list",
      "The instrument shown in the largest open app",
      "The active symbol — the most recent financial instrument you focused on",
      "Nothing: you must always type an instrument code"]),
    ("q", 2, "What happens when you open an application from the Workspace menu?",
     ["It opens as part of your active workspace",
      "It replaces your current workspace",
      "It opens in your web browser",
      "You must restart Workspace first"]),
    ("q", 3, "How many workspaces can you have open at the same time?",
     ["3", "5", "10", "Unlimited"]),
    ("q", 4, "How do you save or share your current workspace?",
     ["File > Save As in Windows",
      "Right-click the workspace tab and choose Save",
      "Upload it to the Template Library",
      "Add it to My Briefcase"]),
    ("q", 5, "Two apps are tuned to the same Link Channel colour. What happens?",
     ["They share the same instrument: change it in one, the other follows",
      "They merge into a single app",
      "They display the same news stories only",
      "Nothing — colours are purely cosmetic"]),
    ("q", 6, "How do you link two apps together?",
     ["Click the Link Channel circle (top-left of each app) and select the same colour",
      "Right-click an app title bar > Link > OK",
      "Drag one app onto the other while holding Shift",
      "Settings > Linking > Enable all apps"]),
    ("q", 7, "Typing “Apple OV” into the Search box opens:",
     ["Apple's quote line only", "Apple's Company Overview", "Apple's News Monitor", "A Screener pre-filtered on Apple"]),
    ("q", 8, "Where is My Briefcase located?",
     ["In the App Library > Files", "Top right of the toolbar", "Inside the Excel add-in only", "In the bottom status bar"]),
    ("q", 9, "How can you access LSEG Workspace help and FAQs?",
     ["Through “Get Help & Support” via the “?” menu", "Via Autosuggest in the Workspace Search box",
      "Using the MyRefinitiv application", "All of these options"]),
    ("q", 10, "True or False: you can snap apps together so they move as a single entity, and stack them as tabs.",
     ["True", "False"]),

    ("h2", "SECTION B — News & Briefcase [M2]"),
    ("q", 11, "★ What is the maximum number of documents that can be added to your Briefcase at one time?",
     ["25", "50", "100", "200"]),
    ("q", 12, "★ Which of the following can be searched for in the News Monitor app?",
     ["Companies", "Portfolios / Lists / Monitors", "Languages", "All of the above"]),
    ("q", 13, "Which function key gives you quick access to news?",
     ["F2", "F5", "F6", "F9"]),
    ("q", 14, "You built a complex Boolean news search and want to reuse it tomorrow. What do you do?",
     ["Save it as a news filter: right-click > File > Favorites > Add to Favorites",
      "Upload it to the Template Library",
      "Add the articles to My Briefcase",
      "Save your workspace"]),
    ("q", 15, "The highlighter in News Monitor is used to:",
     ["Highlight keywords from your search within the results",
      "Highlight cells in Excel",
      "Change the app colour theme",
      "Flag stories for compliance review"]),
    ("q", 16, "The Timeline toggle lets you display stories:",
     ["In reverse chronological order or by the date/time criteria you have chosen",
      "Only in alphabetical order",
      "By number of readers",
      "As a Gantt chart"]),
    ("q", 17, "The main purpose of My Briefcase is to:",
     ["Store Excel formulas",
      "Collect news and documents to read later, offline, on desktop or mobile",
      "Share charts with colleagues",
      "Save Datastream requests"]),
    ("q", 18, "You want Reuters news but WITHOUT a specific source. Which feature?",
     ["Source filters with exclusion criteria", "My Briefcase", "The Link Channel", "The highlighter"]),

    ("h2", "SECTION C — Company, Industry & Charts [M3-M5]"),
    ("q", 19, "In Company Analysis, a company's peers are identified using:",
     ["Analyst coverage overlap",
      "Business classification (TRBC / GICS)",
      "The competitors the company identifies in its filings",
      "All of the above"]),
    ("q", 20, "★ “Quick Search” in the Advanced Research Search app allows filtering across:",
     ["Report type only", "Contributor only", "Industry only", "All of these options"]),
    ("q", 21, "★ The Sustainable Finance landing page app (SUSFIN) contains:",
     ["Research and white papers",
      "A “What's New” section highlighting new offerings",
      "A full list of sustainability-related LSEG Workspace apps and views",
      "All of these answers"]),
    ("q", 22, "★ The Aggregates app supports which business classification schema?",
     ["Standard Industrial Classification (SIC) only",
      "LSEG Business Classifications (TRBC) only",
      "Global Industry Classification Standard (GICS) only",
      "GICS AND LSEG Business Classifications (TRBC)"]),
    ("q", 23, "★ True or False: it is possible to define custom market-cap groups in the Aggregates app.",
     ["True", "False"]),
    ("q", 24, "“Split mode” in the Chart app lets you:",
     ["Split one chart into several views, e.g. by time interval",
      "Split the screen between two users",
      "Export part of a chart to Excel",
      "Separate dividends from prices"]),
    ("q", 25, "You scroll through your Monitor list and want the chart to follow automatically. Which feature?",
     ["Link Channel — same colour on both apps", "My Briefcase", "The Timeline toggle", "The Template Library"]),
    ("q", 26, "The Screener app is used to:",
     ["Filter the universe of instruments by criteria (country, sector, market cap…) and build lists",
      "Aggregate fundamentals of a peer group",
      "Store reusable news filters",
      "Embed data requests in Excel"]),

    ("h2", "SECTION D — Excel, Datastream & API [M6]"),
    ("q", 27, "★ In Excel, what type(s) of data can be retrieved using the Formula Builder?",
     ["Real-time data", "Fundamental data", "Historical data", "All of the above"]),
    ("q", 28, "★ Which of the following statements regarding the Excel add-in is FALSE?",
     ["You can share your Excel files with other users by uploading them to the template library",
      "You can access pre-built templates by clicking 'Templates' in the LSEG menu",
      "You can access pre-built templates from within a cell using the TEMP function",
      "You can upload your own Excel files to 'My Files' so you can access them from anywhere"]),
    ("q", 29, "The Template Library is used to:",
     ["Share your Excel files with other users", "Store your personal passwords",
      "Build charts from templates", "Create new workspaces"]),
    ("q", 30, "'My Files' lets you:",
     ["Upload your own Excel files and access them from anywhere",
      "Share files with every user by default",
      "Edit Reuters news stories",
      "Create new datatypes"]),
    ("q", 31, "★ How can you download data for more than one series in a Static Request?",
     ["In the Series/List field, type the codes for all required series, separated by a comma, or use a list code",
      "In the Series/List field, type the codes for all required series, separated by a plus sign",
      "It is NOT possible to type the codes — you must search for them in Navigator",
      "A Static Request does not allow more than one series; only Time Series requests do"]),
    ("q", 32, "Compared with a Static Request, a Time Series request additionally requires:",
     ["A date range and a frequency", "A chart template", "Administrator rights", "Nothing — they are identical"]),
    ("q", 33, "Which tool do you use to look up series codes and datatype mnemonics?",
     ["Datastream Navigator", "The News Monitor", "My Briefcase", "The Aggregates app"]),
    ("q", 34, "When you select several series in the Navigator, they are inserted into the request field:",
     ["Separated by commas automatically", "Separated by semicolons", "On separate worksheets", "As a single block without separators"]),
    ("q", 35, "With a list used in a Static Request, how many datatypes can be returned for each series in the list?",
     ["Exactly 1", "Between 1 and 10", "Up to 100", "Unlimited"]),
    ("q", 36, "★ Which of the following is a benefit of embedding a Datastream for Office (Excel) request?",
     ["The data can be dynamically updated going forward",
      "The data request can never be changed",
      "An embedded request allows a user to download more data than a manual request",
      "Only embedded requests can download data in chart format"]),
    ("q", 37, "★ Which TWO of the following are ways to add a chart or charts from the Library to a collection?",
     ["Single chart: open the flyout next to the chart, go to Collections, select Add to collection",
      "Multiple charts: tick the checkboxes next to the charts, then click the Add to collection icon on the toolbar",
      "Right-click a chart and select Add to collection",
      "Right-click a folder in the Library > Export > To collection"]),
    ("q", 38, "★ Which of the following data is NOT available for Equities?",
     ["Attribution data types", "Asset4 ESG data", "WorldScope Company Accounts data", "I/B/E/S estimates"]),
    ("q", 39, "The 500+ ESG metrics (environmental, social, governance) available in Datastream come from:",
     ["WorldScope", "Asset4", "I/B/E/S", "GICS"]),
    ("q", 40, "In the Excel add-in, which function is designed for time-series requests?",
     ["=TR", "=RHistory", "=RSearch", "=TEMP"]),

    ("pagebreak",),
    ("h2", "GRILLE DE RÉPONSES — à remplir pendant la simulation"),
    ("table", None,
     [[f"Q{i}", "", f"Q{i+10}", "", f"Q{i+20}", "", f"Q{i+30}", ""] for i in range(1, 11)],
     [15, 31, 15, 31, 15, 31, 15, 31]),

    ("pagebreak",),
    ("h1", "CORRIGÉ DÉTAILLÉ (avec renvois vers la fiche)"),
    ("ans", "Q1 → C", "une nouvelle app affiche **l'instrument actif** (dernier instrument consulté). [M1]"),
    ("ans", "Q2 → A", "toute app ouverte depuis le menu rejoint **automatiquement le workspace actif**. [M1]"),
    ("ans", "Q3 → C", "**10 workspaces** ouverts simultanément maximum. [M1]"),
    ("ans", "Q4 → B", "**clic droit sur l'onglet workspace > Save** (pas de File > Save Windows ni de Briefcase). [M1]"),
    ("ans", "Q5 → A", "même couleur de Link Channel = **instrument partagé** : changer dans l'une met à jour l'autre. [M1]"),
    ("ans", "Q6 → A", "le **cercle coloré en haut à gauche** (Link Channel) : choisir la même couleur. [M1]"),
    ("ans", "Q7 → B", "« Nom + OV » ouvre la **Company Overview**. [M1/M4]"),
    ("ans", "Q8 → B", "My Briefcase = **en haut à droite de la toolbar**. [M1/M2]"),
    ("ans", "Q9 → D", "**All of these options** : menu « ? », Autosuggest, appli MyRefinitiv. [M1]"),
    ("ans", "Q10 → A", "**VRAI** — snap (bloc unique) et stack (onglets) sont des comportements natifs. [M1]"),
    ("ans", "Q11 → C", "**100 documents** maximum en même temps (pas 25/50/200). [M2]"),
    ("ans", "Q12 → D", "**All of the above** : companies, portfolios/lists/monitors, langues (+ sources, marchés, topics…). [M2]"),
    ("ans", "Q13 → D", "**F9** = news. F2 = Quote List, F5 = Refresh, F6 = Max/Min. [M1/M2]"),
    ("ans", "Q14 → A", "news filter réutilisable = **clic droit > File > Favorites > Add to Favorites** (pas la Template Library !). [M2]"),
    ("ans", "Q15 → A", "le highlighter **surligne les mots-clés** de la recherche dans les résultats. [M2]"),
    ("ans", "Q16 → A", "Timeline = **ordre chronologique inverse ou critères date/heure choisis**. [M2]"),
    ("ans", "Q17 → B", "Briefcase = collecter news/documents pour lecture **offline desktop/mobile**. [M2]"),
    ("ans", "Q18 → A", "filtres de sources **avec exclusions** — c'est une fonction native des filtres avancés. [M2]"),
    ("ans", "Q19 → D", "**All of the above** : couverture analystes + classification + concurrents des filings. [M4]"),
    ("ans", "Q20 → D", "**All of these options** : Report Type, Contributor, Industry. [M4]"),
    ("ans", "Q21 → D", "**All of these answers** : white papers + What's New + liste d'apps durables. [M5]"),
    ("ans", "Q22 → D", "**GICS AND TRBC** (jamais SIC seul). [M5]"),
    ("ans", "Q23 → A", "**VRAI** — groupes personnalisés possibles, dont tranches de market cap. [M5]"),
    ("ans", "Q24 → A", "Split mode = **plusieurs vues d'un chart, ex. par intervalle**. [M3]"),
    ("ans", "Q25 → A", "**Link Channel** (même couleur Monitor + Chart). [M3]"),
    ("ans", "Q26 → A", "Screener = **filtrer l'univers et construire des listes** (l'agrégation, c'est Aggregates). [M4]"),
    ("ans", "Q27 → D", "**All of the above** : real-time + fundamental + historical. [M6]"),
    ("ans", "Q28 → C", "la **fonction TEMP n'existe pas** — énoncé FAUX. Template Library = partage, menu Templates = pré-construits, My Files = perso. [M6]"),
    ("ans", "Q29 → A", "Template Library = **partager tes fichiers avec les autres utilisateurs**. [M6]"),
    ("ans", "Q30 → A", "My Files = **tes propres fichiers, accessibles de partout**. [M6]"),
    ("ans", "Q31 → A", "**virgules ou code de liste** — jamais « + », codes tapables sans Navigator. [M6]"),
    ("ans", "Q32 → A", "Time Series = **dates + fréquence** en plus ; la static est une photo (latest value). [M6]"),
    ("ans", "Q33 → A", "**Datastream Navigator** = recherche de séries et datatypes (mnemonics). [M6]"),
    ("ans", "Q34 → A", "le Navigator insère les séries **automatiquement séparées par des virgules**. [M6]"),
    ("ans", "Q35 → B", "**1 à 10 datatypes** par série de liste en static. [M6]"),
    ("ans", "Q36 → A", "embedded = **mise à jour dynamique** des données (seul vrai bénéfice). [M6]"),
    ("ans", "Q37 → A et B (les deux !)", "question « which two » : **flyout** pour une chart, **cases à cocher + icône toolbar** pour plusieurs. [M6]"),
    ("ans", "Q38 → A", "**Attribution data types** non disponibles pour les équities ; WorldScope/Asset4/I/B/E/S disponibles. [M6]"),
    ("ans", "Q39 → B", "**Asset4** = la source ESG de Datastream (WorldScope = comptes, I/B/E/S = estimates). [M6]"),
    ("ans", "Q40 → B", "**=RHistory** = séries temporelles (=TR = fondamentale/estimates, =RSearch = recherche d'instruments, =TEMP n'existe pas). [M6]"),

    ("h2", "BILAN — que faire de ton score ?"),
    ("box", "Barème d'auto-évaluation",
     "**38-40 / 40** : incollable. Relecture « Les 12 pièges » la veille et go pour réserver l'examen (souviens-toi : 30 Q / 60 min / token envoyé par e-mail).\n"
     "**34-37 / 40** : très bien — identifie les modules [Mx] de tes erreurs et relis ces sections seulement, puis refais les questions ratées 24 h plus tard.\n"
     "**< 34 / 40** : reprends la fiche module par module + le corrigé officiel commenté, puis refais CE mock dans 48 h. Ne réserve pas l'examen avant un sans-faute approximatif.\n"
     "**Rappel du format réel** : 30 questions, 60 minutes, ≥ 80 % (max 6 erreurs), 2 tentatives, certificat PDF immédiat."),
]

# ================================================================ MAIN
if __name__ == "__main__":
    import os
    out = os.path.dirname(os.path.abspath(__file__))

    rev_title = "Certification LSEG Workspace"
    rev_pdf = os.path.join(out, "REVISION-Certification-LSEG-Workspace.pdf")
    rev_docx = os.path.join(out, "REVISION-Certification-LSEG-Workspace.docx")
    render_pdf(rev_pdf, rev_title, "Fiche de révision — Certification LSEG Workspace",
               "FICHE DE RÉVISION · OBJECTIF 100 %",
               "Foundational Certification — révision complète hors plateforme (sans installer Excel ni Workspace)",
               REV_META, REV_BLOCKS)
    render_docx(rev_docx, rev_title, "Fiche de révision",
                "FICHE DE RÉVISION · OBJECTIF 100 %",
                "Foundational Certification — révision complète hors plateforme",
                REV_META, REV_BLOCKS)

    mock_title = "Mock Exam — Certification LSEG Workspace"
    mock_pdf = os.path.join(out, "MOCK-EXAM-Certification-LSEG-Workspace.pdf")
    mock_docx = os.path.join(out, "MOCK-EXAM-Certification-LSEG-Workspace.docx")
    render_pdf(mock_pdf, mock_title, "Mock Exam — Certification LSEG Workspace",
               "MOCK EXAM · 40 QUESTIONS",
               "Simulation en conditions réelles — questions en anglais, corrigé détaillé en français, 45 minutes",
               MOCK_META, MOCK_BLOCKS)
    render_docx(mock_docx, mock_title, "Mock Exam",
                "MOCK EXAM · 40 QUESTIONS",
                "Simulation en conditions réelles",
                MOCK_META, MOCK_BLOCKS)
    print("Terminé.")
