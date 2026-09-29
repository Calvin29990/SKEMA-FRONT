# -*- coding: utf-8 -*-
"""CV Calvin MINANG - version Word editable (meme contenu que le PDF)."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1B, 0x3A, 0x5C)
ACCENT = RGBColor(0x2E, 0x5C, 0x8A)
MID = RGBColor(0x4A, 0x4A, 0x4A)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(1.1); sec.bottom_margin = Cm(1.0)
sec.left_margin = Cm(1.5); sec.right_margin = Cm(1.5)

st = doc.styles['Normal']
st.font.name = 'Calibri'; st.font.size = Pt(9.5)
st.paragraph_format.space_after = Pt(0)
st.paragraph_format.space_before = Pt(0)

def p(text=None, size=9.5, bold=False, italic=False, color=None, align=None,
      before=0, after=0, style=None):
    par = doc.add_paragraph(style=style)
    par.paragraph_format.space_before = Pt(before)
    par.paragraph_format.space_after = Pt(after)
    if align is not None: par.alignment = align
    if text:
        r = par.add_run(text); r.bold = bold; r.italic = italic
        r.font.size = Pt(size)
        if color: r.font.color.rgb = color
    return par

def runs(par, parts, size=9.5):
    for txt, bold in parts:
        r = par.add_run(txt); r.bold = bold; r.font.size = Pt(size)
    return par

def hline(par):
    pPr = par._p.get_or_add_pPr()
    b = OxmlElement('w:pBdr'); bot = OxmlElement('w:bottom')
    bot.set(qn('w:val'), 'single'); bot.set(qn('w:sz'), '8')
    bot.set(qn('w:space'), '1'); bot.set(qn('w:color'), '1B3A5C')
    b.append(bot); pPr.append(b)

def section(title):
    par = p(title.upper(), size=10, bold=True, color=NAVY, before=7, after=1)
    hline(par)

def row2(left, right, size=9.3):
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    t.columns[0].width = Cm(12.4); t.columns[1].width = Cm(5.6)
    c0, c1 = t.rows[0].cells
    c0.width = Cm(12.4); c1.width = Cm(5.6)
    pa = c0.paragraphs[0]; pa.paragraph_format.space_after = Pt(0)
    runs(pa, [(left, True)], size=size)
    pb = c1.paragraphs[0]; pb.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pb.paragraph_format.space_after = Pt(0)
    r = pb.add_run(right); r.bold = True; r.font.size = Pt(8.6); r.font.color.rgb = ACCENT
    return t

def bullet(text, size=8.8):
    return p(text, size=size, style='List Bullet', after=1)

# ---------- HEADER ----------
p("CALVIN MINANG", size=19, bold=True, color=NAVY, after=1)
p("GLOBAL MARKETS SALES  ·  FX, RATES & EQUITIES", size=9.8, bold=True, color=ACCENT, after=1)
p("Available from mid-December 2026  ·  Paris, France", size=8.8, italic=True, color=MID, after=2)
p("+33 7 52 97 58 09  ·  calvin.minang@skema.edu  ·  linkedin.com/in/calvin-minang", size=8.6)
p("github.com/Calvin29990  ·  calvin-minibloomberg.netlify.app", size=8.6)

# ---------- PROFILE ----------
section("Profile")
p("Final-year student on SKEMA Business School's Programme Grande Ecole (Master grade) and MSc "
  "Corporate Financial Management. Multi-asset reporting and market monitoring; B2B prospecting, "
  "partner CRM and business development. Target: Global Markets Sales Desk - FX Sales, Rates Sales, "
  "Cross-Asset Sales and equity derivatives.", size=9, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

# ---------- EDUCATION ----------
section("Education")
row2("SKEMA Business School", "Jul 2022 - Dec 2026")
p("Programme Grande Ecole (M2) - Dual degree: MSc Corporate Financial Management", size=8.8)
p("Campus Belo Horizonte, Brazil", size=8.8)
p("Coursework: FX & Arbitrage · Fixed Income · Valuation & Risk Management · Market Risk · Credit Risk · Financial Modelling", size=8.8, after=3)
row2("Ipesup Paris", "2021 - 2022")
p("Preparatory classes for French business schools (CPGE ECE) - intensive mathematics and economics", size=8.8)

# ---------- EXPERIENCE ----------
section("Professional Experience")
exps = [
 ("BPCE Assurances", "Jan - May 2024",
  "Multi-Asset Reporting Analyst (internship) · Paris",
  "Scale: €103bn of life-insurance assets under management, direct business",
  ["Daily production of multi-asset reporting for investment committees: FX pairs, equity indices, rates and credit spreads.",
   "Dashboards and Bloomberg BQL extractions on FX forwards and swaps; monitoring of the €STR and SOFR rate curves.",
   "Excel/VBA automation and market materials on trends and volatility."]),
 ("SKEMA Business School", "Oct - Dec 2025",
  "CRM & Partner Data Assistant · Executive Education", None,
  ["Data cleaning, de-duplication and validation of the partner database.",
   "Updated 537 contact records and 102 company records across the Dubai and Belo Horizonte campuses."]),
 ("FinStart.io", "Jun - Aug 2023", "Sales (internship) · Paris", None,
  ["B2B prospecting targeting independent consultants in finance, accounting, risk and compliance: approx. 100 emails per business day, 500 per week, with full response tracking.",
   "Sourcing and screening of profiles via APEC and LinkedIn; platform invitations."]),
 ("SKEMA Job Service Paris", "Sep 2022 - Aug 2023",
  "Board Member - HR & Partnerships (fixed-term contract)", None,
  ["B2B prospecting and partnership development for the student staffing offering: pitch, needs qualification, interviews and HR events with SKEMA, EY and Stade de France.",
   "Team results H1 2023: €19,100 revenue, 125 students placed across 35 assignments, 5 partners including 4 new (EY, Roadbook, Auverprime, CDEFM)."]),
]
for company, dates, role, note, bl in exps:
    row2(company, dates)
    p(role, size=8.8)
    if note: p(note, size=8.2, italic=True, color=MID)
    for b in bl: bullet(b)
    p("", size=4)

# ---------- PROJECTS ----------
section("Selected Projects")
projs = [
 ("CalvinX - Cross-Asset Market Terminal", "JavaScript",
  "Bloomberg-style dashboard covering FX, equity indices, rates, commodities and volatility, with regime detection and Excel export. Live demo: calvin-minibloomberg.netlify.app"),
 ("ShockLab - Portfolio Stress-Testing Engine", "Python · SQL · Machine Learning",
  "Replays real market analogues (Lehman, COVID, Soleimani) to quantify the P&L impact on a multi-asset portfolio over the 20 days following each shock; SQL event database and interpretable ML."),
 ("Black-Scholes Option Pricer - Options & Greeks", "Excel · VBA",
  "European option valuation with Delta, Gamma and Vega; used to locate the gamma peak of an option book during the Petrobras oil shock and validated against the March 2020 crash."),
]
for title, tech, desc in projs:
    row2(title, tech)
    p(desc, size=8.8, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=4)

# ---------- SKILLS ----------
section("Skills & Certifications")
skills = [
 ("Sales Desk", "B2B prospecting · pitch delivery · needs qualification · response tracking · CRM"),
 ("Markets", "FX (forwards, swaps) · Rates · Fixed Income · equity derivatives · volatility · cross-asset monitoring · market risk (VaR)"),
 ("Tools", "Bloomberg Terminal (BQL) · Excel/VBA · Python (pandas, scikit-learn) · SQL (BigQuery) · R · JavaScript"),
 ("Certifications", "FRM candidate (GARP) · AMF certification in progress · Citi Markets Sales & Trading (Forage) · QuantInsti (Python for Trading; Options Trading Strategies in Python)"),
]
for k, v in skills:
    par = doc.add_paragraph(); par.paragraph_format.space_after = Pt(1)
    runs(par, [(k + "  ", True), (v, False)], size=8.8)

section("Languages")
p("French (native)  ·  English (professional)  ·  Spanish (B2)  ·  Portuguese (basic)", size=8.8)

doc.save("CV_Calvin_MINANG_EN.docx")
print("OK -> CV_Calvin_MINANG_EN.docx")
