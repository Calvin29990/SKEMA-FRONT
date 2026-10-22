#!/usr/bin/env python3
"""
Génère le diaporama de soutenance au format PowerPoint (.pptx),
éditable dans PowerPoint, Keynote, LibreOffice ou Google Slides.

    python3 -m venv .venv && .venv/bin/pip install python-pptx
    .venv/bin/python outils/build-pptx.py

Sortie : export/Soutenance-DBA-MINANG-Serge-Patrick.pptx  (26 diapositives, 16:9)

Charte : encre / bronze. Aucun contenu inventé : les textes sont ceux du
diaporama HTML (index.html), eux-mêmes issus du plan général harmonisé.
Les notes de l'orateur sont incluses dans le volet « Notes » de chaque diapositive.
"""
from __future__ import annotations

import math
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "export" / "Soutenance-DBA-MINANG-Serge-Patrick.pptx"
MAP = ROOT / "medias" / "gabon.png"
LOGO = ROOT / "medias" / "logo-upl.png"

# ---------------------------------------------------------------- charte
BG = RGBColor(0x0B, 0x10, 0x15)
BG2 = RGBColor(0x12, 0x1C, 0x26)
CARD = RGBColor(0x16, 0x21, 0x2C)
CARD_HI = RGBColor(0x1E, 0x2A, 0x36)
LINE = RGBColor(0x27, 0x34, 0x3F)
INK = RGBColor(0xF3, 0xEF, 0xE7)
INK2 = RGBColor(0xD5, 0xD0, 0xC5)
MUTED = RGBColor(0x93, 0xA5, 0xB4)
MUTED2 = RGBColor(0x70, 0x83, 0x92)
BRONZE = RGBColor(0xC8, 0xA1, 0x54)
BRONZE2 = RGBColor(0xE3, 0xC4, 0x85)
BRONZE_DIM = RGBColor(0x2A, 0x24, 0x18)
TEAL = RGBColor(0x5F, 0xB3, 0xA1)

SERIF = "Georgia"
SANS = "Segoe UI"
MONO = "Consolas"

# ---------------------------------------------------------------- géométrie
W, H = 13.333, 7.5
M = 0.72
CW = W - 2 * M
FOOT_Y = 6.92
BODY_TOP = 1.62

warnings: list[str] = []


# ---------------------------------------------------------------- outils
def est_lines(text: str, width_in: float, size_pt: float, factor: float = 0.5) -> int:
    """Estimation du nombre de lignes (garde-fou anti-débordement)."""
    cpl = max(8, int(width_in * 72 / (size_pt * factor)))
    total = 0
    for para in text.split("\n"):
        total += max(1, math.ceil(len(para) / cpl))
    return total


def guard(tag: str, text: str, width_in: float, height_in: float, size_pt: float,
          factor: float = 0.5, line: float = 1.22) -> None:
    need = est_lines(text, width_in, size_pt, factor) * size_pt * line / 72
    if need > height_in + 0.02:
        warnings.append(f"{tag}: texte probablement trop long ({need:.2f} in pour {height_in:.2f} in)")


def rgb_alpha(shape, alpha: float) -> None:
    """Ajoute une transparence au remplissage (0 = opaque, 1 = invisible)."""
    solid = shape.fill._xPr.find(qn("a:solidFill"))
    if solid is None:
        return
    clr = solid[0]
    srgb = clr.find(qn("a:srgbClr"))
    if srgb is None:
        return
    for tag in ("a:alpha",):
        el = srgb.find(qn(tag))
        if el is not None:
            srgb.remove(el)
    el = srgb.makeelement(qn("a:alpha"), {"val": str(int((1 - alpha) * 100000))})
    srgb.append(el)


def noline(shape) -> None:
    shape.line.fill.background()


def rect(slide, x, y, w, h, fill=None, line=None, radius=None, line_w=0.9, shape_type=None):
    if shape_type is None:
        shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    if radius and shp.adjustments:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        noline(shp)
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    shp.text_frame.word_wrap = True
    return shp


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return box


def para(tf, text="", size=13, color=INK2, bold=False, font=SANS, italic=False,
         align=PP_ALIGN.LEFT, space_before=0, space_after=0, line=1.18, first=False,
         bullet=None, bullet_color=None, caps=False, tracking=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.line_spacing = line
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    if bullet:
        r = p.add_run()
        r.text = bullet + "  "
        r.font.size = Pt(size * 0.8)
        r.font.bold = True
        r.font.name = font
        r.font.color.rgb = bullet_color or BRONZE
    if text:
        r = p.add_run()
        r.text = text.upper() if caps else text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = font
        r.font.color.rgb = color
        if tracking:
            r.font._rPr.set("spc", str(int(tracking * 100)))
    return p


def rich(tf, chunks, size=13, font=SANS, line=1.18, space_before=0, space_after=0,
         align=PP_ALIGN.LEFT, first=False, bullet=None):
    """Un paragraphe composé de plusieurs fragments (texte, couleur, gras)."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.line_spacing = line
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    if bullet:
        r = p.add_run()
        r.text = bullet + "  "
        r.font.size = Pt(size * 0.8)
        r.font.bold = True
        r.font.name = font
        r.font.color.rgb = BRONZE
    for text, color, bold in chunks:
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = font
        r.font.color.rgb = color
    return p


def notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def new_slide(prs, title_tag: str):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = rect(s, 0, 0, W, H, BG2)
    # dégradé diagonal discret
    try:
        bg.fill.gradient()
        bg.fill.gradient_angle = 135
        st = bg.fill.gradient_stops
        st[0].color.rgb = BG2
        st[0].position = 0.0
        st[1].color.rgb = BG
        st[1].position = 1.0
    except Exception:
        pass
    # filet bronze en pied de diapositive
    bar = rect(s, 0, H - 0.032, W, 0.032, BRONZE)
    rgb_alpha(bar, 0.45)
    s._tag = title_tag  # type: ignore[attr-defined]
    return s


def footer(slide, src: str, n: int) -> None:
    line = rect(slide, M, FOOT_Y - 0.14, CW, 0.008, LINE)
    rgb_alpha(line, 0.55)
    box = textbox(slide, M, FOOT_Y - 0.04, CW - 0.75, 0.5)
    tf = box.text_frame
    rich(tf, [("Sources : ", MUTED, True), (src, MUTED2, False)], size=7.4, first=True, line=1.1)
    pg = textbox(slide, W - M - 0.72, FOOT_Y - 0.04, 0.72, 0.3)
    para(pg.text_frame, f"{n} / 26", size=8, color=MUTED, font=MONO, align=PP_ALIGN.RIGHT, first=True)


def head(slide, kicker: str, title: str, under: str | None = None) -> float:
    y = 0.52
    k = textbox(slide, M, y, CW, 0.22)
    para(k.text_frame, kicker, size=9.5, color=BRONZE2, bold=True, caps=True, tracking=1.6, first=True)
    ty = y + 0.26
    n = len(title)
    size = 24.5 if n < 72 else (22.5 if n < 92 else 20.5)
    tbox = textbox(slide, M, ty, CW - 0.2, 1.0)
    tf = tbox.text_frame
    p = tf.paragraphs[0]
    p.line_spacing = 1.05
    r = p.add_run()
    r.text = title
    r.font.size = Pt(size)
    r.font.name = SERIF
    r.font.bold = False
    r.font.color.rgb = INK
    guard("titre", title, CW - 0.3, 1.0, size, factor=0.47, line=1.06)
    end = ty + size / 72 * 1.06 * est_lines(title, CW - 0.3, size, 0.47)
    if under:
        u = textbox(slide, M, end + 0.04, CW - 1.2, 0.4)
        para(u.text_frame, under, size=11.5, color=MUTED, first=True, line=1.3)
        end += 0.34
    return end + 0.34


def card(slide, x, y, w, h, title: str | None = None, body: str | None = None,
         bullets: list[str] | None = None, accent=False, tight=False, body_size=11.0,
         title_size=9.0, bullet_size=10.7, body_color=INK2, tag: str = "") -> None:
    fill = BRONZE_DIM if accent else CARD
    rect(slide, x, y, w, h, fill, BRONZE if accent else LINE, radius=0.055)
    pad = 0.14 if tight else 0.19
    inner = textbox(slide, x + pad, y + (0.11 if tight else 0.15), w - 2 * pad, h - 0.2,
                    anchor=MSO_ANCHOR.MIDDLE)
    tf = inner.text_frame
    first = True
    if title:
        para(tf, title, size=title_size, color=BRONZE2, bold=True, caps=True,
             tracking=0.7, first=True, space_after=2.5, line=1.15)
        guard(f"carte[{title}]", title, w - 2 * pad, 0.3, title_size)
        first = False
    if body:
        para(tf, body, size=body_size, color=body_color, first=first, line=1.26)
        guard(f"corps[{title or tag}]", body, w - 2 * pad, h - 0.45, body_size)
        first = False
    if bullets:
        for b in bullets:
            para(tf, b, size=bullet_size, color=body_color, first=first, line=1.24,
                 bullet="■", space_before=(7.0 if not first else 0.0))
            first = False
        guard(f"puces[{title or tag}]", "\n".join(bullets), w - 2 * pad - 0.15,
              h - 0.42, bullet_size, factor=0.52)


def stat(slide, x, y, w, h, value: str, label: str, sub: str | None = None) -> None:
    rect(slide, x, y, w, h, CARD, LINE, radius=0.09)
    tf = textbox(slide, x + 0.14, y + 0.08, w - 0.28, h - 0.16, anchor=MSO_ANCHOR.MIDDLE).text_frame
    p = tf.paragraphs[0]
    p.line_spacing = 1.0
    r = p.add_run()
    r.text = value
    r.font.size = Pt(26)
    r.font.name = SERIF
    r.font.color.rgb = INK
    para(tf, label, size=8.8, color=MUTED, space_before=3, line=1.2)
    if sub:
        para(tf, sub.upper(), size=6.8, color=MUTED2, space_before=2.5, tracking=1.2)


def table(slide, x, y, w, widths, rows, header=True, size=10.2, row_h=0.33,
          head_size=7.8, align_last_right=False) -> float:
    """Tableau sobre : en-tête en petites capitales, filets fins, aucune bordure lourde."""
    n_rows, n_cols = len(rows), len(rows[0])
    total = sum(widths)
    widths = [wd / total * w for wd in widths]
    shape = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y), Inches(w), Inches(row_h * n_rows))
    tbl = shape.table
    tbl.first_row = header
    tbl.horz_banding = False
    for j, wd in enumerate(widths):
        tbl.columns[j].width = Inches(wd)
    for i in range(n_rows):
        tbl.rows[i].height = Inches(row_h)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = Inches(0.1)
            cell.margin_right = Inches(0.1)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if header and i == 0:
                cell.fill.fore_color.rgb = CARD_HI
            elif i % 2 == 0:
                cell.fill.fore_color.rgb = CARD
            else:
                cell.fill.fore_color.rgb = BG2
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.line_spacing = 1.12
            r = p.add_run()
            r.text = val
            r.font.name = SANS
            if header and i == 0:
                r.font.size = Pt(head_size)
                r.font.bold = True
                r.font.color.rgb = MUTED
                r.font._rPr.set("spc", "110")
                r.text = val.upper()
            else:
                r.font.size = Pt(size)
                r.font.color.rgb = INK2 if j else INK
                if j == 0 and len(val) <= 5:
                    r.font.name = MONO
                    r.font.color.rgb = BRONZE2
                    r.font.bold = True
            if align_last_right and j == n_cols - 1:
                p.alignment = PP_ALIGN.RIGHT
    return y + row_h * n_rows


def arrow(slide, x1, y1, x2, y2, color=BRONZE, width=1.1) -> None:
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = color
    conn.line.width = Pt(width)
    ln = conn.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)


# ---------------------------------------------------------------- diapositives
def slide_cover(prs):
    s = new_slide(prs, "Couverture")
    k = textbox(s, M, 0.62, 7.4, 0.24)
    para(k.text_frame, "Soutenance de thèse doctorale · DBA", size=9.5, color=BRONZE2,
         bold=True, caps=True, tracking=1.7, first=True)

    t = textbox(s, M, 1.02, 8.3, 2.5).text_frame
    p = t.paragraphs[0]
    p.line_spacing = 1.04
    r = p.add_run()
    r.text = "La gouvernance augmentée des établissements privés d'enseignement supérieur au Gabon"
    r.font.size = Pt(33)
    r.font.name = SERIF
    r.font.color.rgb = INK
    guard("couverture", r.text, 8.3, 2.3, 32, factor=0.47, line=1.08)

    sub = textbox(s, M, 3.28, 7.9, 1.0).text_frame
    para(sub, "Articulation entre intelligence artificielle et intelligence émotionnelle au service "
              "de la performance globale, de la responsabilité sociétale et du développement durable, "
              "en situation de polycrise.", size=12.5, color=INK2, first=True, line=1.4)
    para(sub, "Titre indicatif — plan général harmonisé, septembre 2026.", size=8.6,
         color=MUTED2, space_before=5)

    # logo + mention DBA
    try:
        s.shapes.add_picture(str(LOGO), Inches(9.62), Inches(0.86), height=Inches(1.62))
        tag = textbox(s, 9.62, 2.62, 2.95, 0.7).text_frame
        para(tag, "DOCTORATE IN BUSINESS ADMINISTRATION", size=7.2, color=MUTED2,
             bold=True, tracking=1.3, first=True, line=1.4)
    except Exception:
        warnings.append("logo introuvable : medias/logo-upl.png")

    # bloc auteur
    bar = rect(s, M, 5.24, 0.03, 1.02, BRONZE)
    who = textbox(s, M + 0.18, 5.2, 5.6, 1.15).text_frame
    para(who, "Serge Patrick MINANG", size=17, color=INK, font=SERIF, first=True, line=1.15)
    para(who, "Ingénieur · MBA · doctorant DBA — praticien-chercheur", size=9.6, color=MUTED, space_before=3)
    para(who, "[à compléter par l'auteur : université d'inscription · direction de thèse · date de soutenance]",
         size=7.6, color=MUTED2, space_before=6, line=1.3)

    right = textbox(s, 8.5, 5.3, 4.1, 1.0).text_frame
    para(right, "DOCUMENT DE TRAVAIL · 22 OCTOBRE 2026", size=7.6, color=BRONZE2, bold=True,
         tracking=1.3, align=PP_ALIGN.RIGHT, first=True)
    para(right, "Chapitres 5 et 6 préparatoires : aucun résultat n'est présumé dans ce diaporama.",
         size=8.4, color=MUTED, align=PP_ALIGN.RIGHT, space_before=6, line=1.3)

    footer(s, "plan général harmonisé de la thèse et liste unique des tableaux et figures (septembre 2026). "
              "Provenance détaillée des chiffres et des visuels : page « Sources » du dossier.", 1)
    notes(s, "Script — « Monsieur le Président, Madame, Monsieur les membres du jury, je vous remercie de votre présence. "
             "Je présente une recherche Doctorate in Business Administration consacrée à la gouvernance des établissements "
             "privés d'enseignement supérieur gabonais, et à la manière dont deux capacités — technologiques et humaines — "
             "s'y articulent. » 45 secondes. Le titre est indicatif ; le travail est en cours, chapitres 5 et 6 préparatoires.")


def slide_map(prs):
    s = new_slide(prs, "Le terrain : le Gabon")
    y = head(s, "Terrain et contexte",
             "Neuf provinces, un secteur privé d'enseignement supérieur à gouvernance contrainte")
    # carte
    try:
        s.shapes.add_picture(str(MAP), Inches(7.9), Inches(1.42), height=Inches(5.55))
    except Exception:
        warnings.append("carte introuvable : medias/gabon.png")
    # colonne texte
    tf = textbox(s, M, y + 0.02, 6.9, 4.3).text_frame
    items = [
        ("Cinq établissements enquêtés", ", présentés sous les codes EPAS-A à EPES-E : régime de confidentialité appliqué à la version publique."),
        ("Un ancrage institutionnel", " : l'UPL, siège à la Sablière (Libreville) — Executive MBA ouvert depuis 2022, cours du soir de 17 h à 21 h, appui académique de l'Université de Douala."),
        ("Un secteur sous contraintes", " : cadre juridique et institutionnel, modèles économiques adossés aux droits de scolarité, parties prenantes aux attentes divergentes (chapitre 1)."),
    ]
    first = True
    for bold_part, rest in items:
        rich(tf, [(bold_part, INK, True), (rest.replace("EPAS-A", "EPES-A"), INK2, False)],
             size=11.4, bullet="■", first=first, space_before=0 if first else 9, line=1.3)
        first = False
    para(tf, "Le terrain gabonais n'est pas un décor : il fixe les contraintes de gouvernance que la "
             "thèse prend pour objet (chapitre 1), distinctes de la justification du choix des cas (section 4.3).",
         size=9.4, color=MUTED, space_before=14, line=1.35)
    footer(s, "fond de carte geoBoundaries — GAB ADM1 (GAUL 2001), licence CC BY 4.0 ; repères institutionnels : "
              "upl-gabon.com (consulté le 22/10/2026) ; plan harmonisé, chapitre 1.", 2)
    notes(s, "Situer le terrain sans nommer les établissements : l'anonymat est une exigence éthique du protocole. "
             "Question probable — « Pourquoi cinq cas et non un seul ? » : répondre par la variation (critères C1–C5) "
             "et par la logique de réplication théorique, pas par la représentativité statistique.")


def slide_polycrisis(prs):
    s = new_slide(prs, "La polycrise")
    y = head(s, "Chapitre 1 · § 1.10",
             "La polycrise n'est pas une somme de crises : c'est un problème transversal de gouvernance",
             "Quatre familles de tensions qui se renforcent, et dont les effets se propagent d'un registre à l'autre.")
    items = [
        ("Financière", "Viabilité économique, dépendance aux droits de scolarité, échéanciers, trésorerie."),
        ("Académique", "Qualité des enseignements, encadrement, ressources documentaires, accréditations."),
        ("Humaine et sociale", "Charge des équipes, compétences, climat interne, attentes des étudiants et des familles."),
        ("Institutionnelle", "Agréments, contrôles des tutelles, fiscalité et obligations sociales, concurrence locale."),
    ]
    gap = 0.18
    cw = (CW - 3 * gap) / 4
    for i, (t, b) in enumerate(items):
        card(s, M + i * (cw + gap), y, cw, 1.5, title=t, body=b, tight=True, body_size=9.8)

    y2 = y + 1.72
    hw = (CW - 0.22) * 0.46
    card(s, M, y2, hw, 1.62, title="Effets de cascade",
         body="Un incident de trésorerie dégrade la qualité perçue ; une tension sociale dégrade la "
              "réputation ; une exigence réglementaire nouvelle déplace l'arbitrage budgétaire. "
              "Les crises s'entretiennent.", accent=True, body_size=10.4)
    rx = M + hw + 0.22
    tf = textbox(s, rx + 0.1, y2 + 0.05, CW - hw - 0.3, 1.6).text_frame
    for i, b in enumerate([
        "Conséquence méthodologique : observer des séquences décisionnelles, pas seulement des états.",
        "Conséquence théorique : les modèles classiques de gouvernance supposent une incertitude bornée.",
        "Conséquence managériale : la décision devient un arbitrage entre exigences incommensurables.",
    ]):
        para(tf, b, size=10.4, color=INK2, bullet="■", first=(i == 0), space_before=0 if i == 0 else 7, line=1.28)
    footer(s, "plan harmonisé § 1.10 ; tableau 1.3 (dimensions de la polycrise et implications pour la gouvernance) ; "
              "figure 1.2 (interdépendance des tensions).", 3)
    notes(s, "Insister sur « transversal » : la polycrise est le cadre d'analyse, pas un slogan d'époque. "
             "Ne pas dater ni quantifier des crises gabonaises non documentées dans le manuscrit.")


def slide_problem(prs):
    s = new_slide(prs, "Problème")
    y = head(s, "Problématique", "Un problème managérial réel, un angle scientifique encore peu investi")
    hw = (CW - 0.24) / 2
    card(s, M, y, hw, 3.9, title="Problème managérial",
         body="Les dirigeants d'EPES arbitrent en continu entre viabilité économique, qualité académique, "
              "attentes des parties prenantes, responsabilité sociétale et durabilité — sans commune mesure "
              "unique entre ces exigences.", body_size=11.0, bullets=[
             "Les outils numériques produisent des informations, mais ne tranchent pas l'arbitrage.",
             "La dimension humaine de la décision (relation, confiance, régulation des tensions) reste peu outillée.",
             "L'investissement technologique est parfois plus rapide que la capacité collective à s'en servir.",
         ], bullet_size=10.4)
    x2 = M + hw + 0.24
    card(s, x2, y, hw, 3.9, title="Problème scientifique",
         body="La littérature traite séparément la gouvernance, l'intelligence artificielle, l'intelligence "
              "émotionnelle et la performance durable. Leur articulation est peu documentée, particulièrement "
              "en contexte africain et dans l'enseignement supérieur privé.", body_size=11.0, bullets=[
             "Lacune théorique : l'articulation IA–IE n'est pas modélisée comme un processus.",
             "Lacune contextuelle : les EPES gabonais sont peu étudiés comme cas de gouvernance.",
             "Lacune méthodologique : peu d'études multi-cas longitudinales sur les scènes de décision.",
         ], bullet_size=10.4)
    footer(s, "plan harmonisé — introduction (contexte et justification, problématique) ; § 2.10 "
              "(synthèse critique et lacunes) ; tableau 2.5.", 4)
    notes(s, "Deux problèmes, une seule recherche : le problème managérial justifie l'intérêt, le problème "
             "scientifique justifie la méthode. Question probable — « N'est-ce pas deux thèses en une ? » : "
             "non, l'articulation IA–IE est l'objet unique, la performance durable est le domaine d'effets.")


def slide_question(prs):
    s = new_slide(prs, "Question centrale")
    y = head(s, "Question centrale", "Une seule question, formulée à l'identique dans les six chapitres")
    qh = 2.02
    rect(s, M, y, CW, qh, BRONZE_DIM, None)
    rect(s, M, y, 0.045, qh, BRONZE)
    tf = textbox(s, M + 0.34, y + 0.05, CW - 0.7, qh - 0.1, anchor=MSO_ANCHOR.MIDDLE).text_frame
    para(tf, "« Comment et sous quelles conditions l'articulation entre l'intelligence artificielle et "
             "l'intelligence émotionnelle transforme-t-elle la gouvernance des EPES gabonais en situation de "
             "polycrise, et quels effets produit-elle sur leur performance globale, leur RSE et leur "
             "contribution au développement durable ? »", size=14.6, color=INK, font=SERIF, first=True, line=1.34)
    para(tf, "Plan général harmonisé — § 1.3, repris sans reformulation dans l'ensemble de la thèse.",
         size=8.4, color=MUTED, space_before=8)

    y2 = y + qh + 0.30
    cw = (CW - 2 * 0.18) / 3
    for i, (t, b) in enumerate([
        ("« Comment »", "Un processus à décrire : les scènes de préparation, de délibération, de décision et de suivi."),
        ("« Sous quelles conditions »", "Des conditions organisationnelles et institutionnelles à identifier, non des effets supposés acquis."),
        ("« Quels effets »", "Des effets à documenter, favorables, défavorables ou non établis."),
    ]):
        x = M + i * (cw + 0.18)
        rect(s, x, y2, cw, 1.62, CARD, LINE, radius=0.055)
        tf2 = textbox(s, x + 0.14, y2 + 0.2, cw - 0.28, 1.24, anchor=MSO_ANCHOR.TOP).text_frame
        para(tf2, t, size=9.0, color=BRONZE2, bold=True, caps=True, tracking=0.7, first=True,
             space_after=2.5, line=1.15)
        para(tf2, b, size=10.6, color=INK2, line=1.28)
    footer(s, "plan harmonisé, introduction § 1.3 ; alignement questions–propositions–sources–codes : "
              "tableau 3.4 et annexe G.", 5)
    notes(s, "Lire la question lentement : c'est la phrase que le jury doit pouvoir répéter après la soutenance. "
             "Chaque mot compte — « comment » (processus), « conditions » (contingence), « effets » (prudence).")


def slide_qr(prs):
    s = new_slide(prs, "Quatre questions spécifiques")
    y = head(s, "Opérationnalisation", "Quatre questions spécifiques, quatre objets d'enquête")
    gap = 0.18
    cw = (CW - 3 * gap) / 4
    items = [
        ("QR1", "Gouvernance en polycrise", "Comprendre comment les EPES gouvernent en situation de polycrise : arbitrages, instances, pratiques."),
        ("QR2", "Usages de l'IA", "Caractériser les usages réels de l'intelligence artificielle dans les décisions de gouvernance."),
        ("QR3", "Rôle de l'IE", "Examiner le rôle de l'intelligence émotionnelle dans le jugement et l'action collective."),
        ("QR4", "Articulation et effets", "Analyser les conditions et les effets de leur articulation éventuelle."),
    ]
    for i, (code, t, b) in enumerate(items):
        x = M + i * (cw + gap)
        rect(s, x, y, cw, 2.18, CARD, LINE, radius=0.06)
        tf = textbox(s, x + 0.2, y + 0.2, cw - 0.4, 1.86, anchor=MSO_ANCHOR.TOP).text_frame
        para(tf, code, size=10.6, color=BRONZE2, bold=True, font=MONO, first=True, space_after=6)
        para(tf, t, size=10.4, color=INK, bold=True, caps=True, tracking=0.4, space_after=7, line=1.15)
        para(tf, b, size=10.9, color=INK2, line=1.32)

    card(s, M, y + 2.42, CW, 1.12, title="Ce que le protocole ne fait pas",
         body="Aucune question ne présuppose que l'IA améliore la gouvernance, ni que l'IE soit le facteur "
              "déterminant. Une absence d'articulation — ou une articulation limitée — est une issue possible "
              "et analysable.", accent=True, body_size=11.0, tight=True)
    footer(s, "plan harmonisé, introduction § 1.3 et § 3.4 ; tableau 3.4 (alignement questions / "
              "propositions / collecte / codes).", 6)
    notes(s, "QR1 à QR4 recouvrent les quatre piliers du cadre : contexte, technologie, humain, articulation. "
             "Question probable — « Où est la performance ? » : QR4 renvoie aux effets, traités aux propositions P5 et P6.")


def slide_objectives(prs):
    s = new_slide(prs, "Objectifs")
    y = head(s, "Objectifs", "Un objectif général, quatre objectifs spécifiques alignés sur les questions")
    lw = CW * 0.42
    rect(s, M, y, lw, 3.9, BRONZE_DIM, BRONZE, radius=0.05)
    tf = textbox(s, M + 0.2, y + 0.18, lw - 0.4, 3.6).text_frame
    para(tf, "OBJECTIF GÉNÉRAL", size=8.8, color=BRONZE2, bold=True, tracking=0.7, first=True)
    para(tf, "Comprendre comment et sous quelles conditions l'articulation entre intelligence artificielle "
             "et intelligence émotionnelle transforme la gouvernance des EPES gabonais en situation de "
             "polycrise, et quels effets en résultent sur la performance globale, la RSE et le développement "
             "durable.", size=11.6, color=INK2, space_before=8, line=1.34)
    sep = rect(s, M + 0.2, y + 2.66, lw - 0.4, 0.008, LINE)
    rgb_alpha(sep, 0.5)
    para(tf, "Objectif de compréhension, non de démonstration d'efficacité : le design est qualitatif "
             "et abductif.", size=9.2, color=MUTED, space_before=10, line=1.3)

    tx = M + lw + 0.34
    rows = [
        ("Objectif spécifique", "Question"),
        ("Comprendre la gouvernance des EPES en situation de polycrise", "QR1"),
        ("Caractériser les usages de l'intelligence artificielle", "QR2"),
        ("Examiner le rôle de l'intelligence émotionnelle", "QR3"),
        ("Analyser les conditions et les effets de leur articulation", "QR4"),
    ]
    table(s, tx, y + 0.05, CW - lw - 0.34, [3.2, 0.8], rows, row_h=0.64, size=10.8)
    para(textbox(s, tx, y + 3.34, CW - lw - 0.34, 0.9).text_frame,
         "Intérêts scientifique, managérial et sociétal : annoncés comme des contributions recherchées, "
         "jamais comme des apports déjà démontrés (introduction § 1.6).",
         size=9.2, color=MUTED, first=True, line=1.3)
    footer(s, "plan harmonisé, introduction § 1.4 (objectifs) et § 1.6 (intérêts scientifique, "
              "managérial, sociétal).", 7)
    notes(s, "Souligner la cohérence verticale : objectif général → QR1–QR4 → P1–P7 → guide d'entretien → codes NVivo. "
             "Question probable — « Pourquoi pas d'hypothèses statistiques ? » : posture interprétativiste, "
             "propositions à confronter et non hypothèses causales.")


def slide_lit(prs):
    s = new_slide(prs, "Revue critique")
    y = head(s, "Chapitre 2 · revue critique", "Quatre courants mobilisés, une intersection à construire",
             "La revue n'accumule pas les références : elle organise la convergence vers un objet non traité — l'articulation.")
    gap = 0.2
    cw = (CW - gap) / 2
    items = [
        ("1 · Gouvernance organisationnelle et universitaire",
         "Registres collégial et managérial, spécificités de l'établissement d'enseignement, limites des modèles classiques face à l'incertitude (§ 2.2–2.3)."),
        ("2 · Intelligence artificielle et décision",
         "Contributions, conditions d'usage, limites et risques ; distinction stricte avec la numérisation, l'automatisation et l'analyse de données (§ 2.5)."),
        ("3 · Intelligence émotionnelle et leadership",
         "Modèles, apports et limites : perception, régulation, relations, gestion des tensions dans le jugement (§ 2.6)."),
        ("4 · Performance globale, RSE, durabilité",
         "Dimensions économiques, académiques, humaines, sociales, environnementales ; arbitrages et limites de l'évaluation (§ 2.8–2.9)."),
    ]
    for i, (t, b) in enumerate(items):
        x = M + (i % 2) * (cw + gap)
        yy = y + (i // 2) * 1.12
        card(s, x, yy, cw, 0.98, title=t, body=b, tight=True, body_size=9.6, title_size=8.2)

    card(s, M, y + 2.32, CW, 1.16, title="Le point d'intersection — et la lacune",
         body="Les travaux sur la complémentarité IA–IE (§ 2.7) restent prescriptifs et rarement situés. Aucun cadre "
              "ne décrit l'articulation comme un processus de gouvernance observable dans un établissement, a fortiori "
              "en contexte gabonais. C'est cette lacune que la recherche investit.",
         accent=True, body_size=10.6)
    footer(s, "plan harmonisé, chapitre 2 (§ 2.1–2.10) ; figure 2.1 (convergence des courants) ; tableaux 2.1 à 2.5. "
              "Les références bibliographiques seront celles effectivement mobilisées dans le manuscrit.", 8)
    notes(s, "« Je ne propose pas une théorie de plus, je propose un modèle exploratoire situé, à éprouver. » "
             "Ne citer aucun auteur de mémoire : la bibliographie du manuscrit n'est pas encore arrêtée.")


def slide_model(prs):
    s = new_slide(prs, "Modèle exploratoire")
    y = head(s, "Chapitre 3 · § 3.4", "Le modèle exploratoire : une configuration de processus, pas un outil")

    # contexte
    cx, cy, cw, ch = M + 3.45, y + 0.02, 4.9, 0.42
    rect(s, cx, cy, cw, ch, CARD_HI, LINE, radius=0.12)
    para(textbox(s, cx + 0.1, cy + 0.1, cw - 0.2, 0.24).text_frame,
         "SITUATION DE POLYCRISE — P1", size=8.6, color=BRONZE2, bold=True,
         tracking=1.1, align=PP_ALIGN.CENTER, first=True)
    ytop = cy + ch + 0.3
    arrow(s, cx + cw / 2, cy + ch + 0.02, cx + cw / 2, ytop - 0.02)

    # IA / GA / IE
    bw, bh = 3.0, 1.9
    rect(s, M, ytop, bw, bh, CARD, BRONZE, radius=0.06)
    t = textbox(s, M + 0.15, ytop + 0.14, bw - 0.3, bh - 0.25).text_frame
    para(t, "INTELLIGENCE ARTIFICIELLE", size=7.6, color=BRONZE2, bold=True, tracking=0.6, first=True, space_after=5)
    para(t, "Usages effectifs", size=10.2, color=INK, bold=True, space_after=3)
    para(t, "outils, données, analyses, alertes entrant réellement dans la décision — P2",
         size=9.2, color=MUTED, line=1.28)

    gx = M + bw + 0.55
    gw = CW - 2 * bw - 1.1
    rect(s, gx, ytop - 0.12, gw, bh + 0.24, BRONZE_DIM, BRONZE, radius=0.05, line_w=1.4)
    g = textbox(s, gx + 0.18, ytop + 0.02, gw - 0.36, bh).text_frame
    para(g, "GOUVERNANCE AUGMENTÉE — P4", size=8.0, color=BRONZE2, bold=True, tracking=1.0,
         align=PP_ALIGN.CENTER, first=True, space_after=6)
    para(g, "Configuration des processus", size=10.4, color=INK, bold=True, align=PP_ALIGN.CENTER, space_after=3)
    para(g, "préparation → délibération → décision → suivi", size=9.4, color=MUTED,
         align=PP_ALIGN.CENTER, space_after=6)
    para(g, "articulation délibérée, décision et responsabilité humaines", size=8.6,
         color=BRONZE2, align=PP_ALIGN.CENTER, line=1.25)

    ex = gx + gw + 0.55
    rect(s, ex, ytop, bw, bh, CARD, TEAL, radius=0.06)
    t2 = textbox(s, ex + 0.15, ytop + 0.14, bw - 0.3, bh - 0.25).text_frame
    para(t2, "INTELLIGENCE ÉMOTIONNELLE", size=7.6, color=TEAL, bold=True, tracking=0.6, first=True, space_after=5)
    para(t2, "Jugement et relations", size=10.2, color=INK, bold=True, space_after=3)
    para(t2, "écoute, régulation des tensions, annonces, confiance — P3", size=9.2, color=MUTED, line=1.28)

    arrow(s, M + bw + 0.04, ytop + bh / 2, gx - 0.06, ytop + bh / 2)
    arrow(s, ex - 0.04, ytop + bh / 2, ex - 0.52, ytop + bh / 2)

    # effets
    ey = ytop + bh + 0.42
    ew = 3.3
    e1x = gx - ew - 0.15
    e2x = gx + gw + 0.15
    arrow(s, gx + 0.55, ytop + bh + 0.02, e1x + ew - 0.3, ey - 0.04)
    arrow(s, gx + gw - 0.55, ytop + bh + 0.02, e2x + 0.3, ey - 0.04)
    for xx, title, l1, l2 in [
        (e1x, "PERFORMANCE GLOBALE — P5", "économique · académique · humaine", "sociale · environnementale"),
        (e2x, "RSE ET DURABILITÉ — P6", "engagements et pratiques", "finalités de long terme"),
    ]:
        rect(s, xx, ey, ew, 1.0, CARD, LINE, radius=0.06)
        tt = textbox(s, xx + 0.12, ey + 0.05, ew - 0.24, 0.9, anchor=MSO_ANCHOR.MIDDLE).text_frame
        para(tt, title, size=7.8, color=BRONZE2, bold=True, tracking=0.6, align=PP_ALIGN.CENTER,
             first=True, space_after=4)
        para(tt, l1, size=8.6, color=MUTED, align=PP_ALIGN.CENTER, line=1.3)
        para(tt, l2, size=8.6, color=MUTED, align=PP_ALIGN.CENTER, line=1.3)

    para(textbox(s, M, ey + 1.06, CW, 0.5).text_frame,
         "Conditions organisationnelles et institutionnelles — effet modérateur (P7) · garde-fous : transparence, "
         "responsabilité, prévention des effets défavorables.", size=8.6, color=MUTED2, first=True)
    para(textbox(s, M, ey + 1.3, CW, 0.4).text_frame,
         "Modèle exploratoire : construit avant l'enquête, il ne représente aucune relation comme démontrée. "
         "Sa révision relève du chapitre 6 ; sa confrontation empirique, du chapitre 5.",
         size=8.4, color=MUTED2, first=True)
    footer(s, "plan harmonisé § 3.4–3.5 et figure 3.1 (modèle conceptuel exploratoire). Reprise graphique du "
              "dossier de travail, non une validation empirique.", 9)
    notes(s, "« Ce schéma est une hypothèse de lecture, pas un résultat. Le jury doit pouvoir le contester : c'est prévu. » "
             "Question probable — « Pourquoi la polycrise en amont ? » : parce qu'elle déclenche et contraint les processus, "
             "sans les déterminer.")


def slide_props(prs):
    s = new_slide(prs, "Sept propositions")
    y = head(s, "Chapitre 3 · § 3.6",
             "Sept propositions, formulées pour être confrontées — et possiblement contredites")
    rows = [
        ("Code", "Proposition", "Mécanisme attendu"),
        ("P1", "Polycrise et transformation des pratiques de gouvernance", "La polycrise reconfigure les arbitrages et les instances"),
        ("P2", "Contribution conditionnelle de l'intelligence artificielle", "Détection, anticipation, lecture des signaux"),
        ("P3", "Contribution de l'intelligence émotionnelle au jugement et à l'action collective", "Interprétation, coordination, régulation des tensions"),
        ("P4", "Orchestration IA–IE et émergence de la gouvernance augmentée", "Délibération, arbitrage éthique, suivi et apprentissage"),
        ("P5", "Gouvernance augmentée et performance globale", "Effets à documenter, y compris défavorables"),
        ("P6", "Gouvernance augmentée, RSE et développement durable", "Engagements, pratiques, finalités de long terme"),
        ("P7", "Effet des conditions organisationnelles et institutionnelles", "Conditions de possibilité, freins, effets de contexte"),
    ]
    table(s, M, y, CW, [0.7, 5.2, 3.6], rows, row_h=0.52, size=10.0)
    para(textbox(s, M, y + 0.52 * 8 + 0.2, CW, 0.6).text_frame,
         "Chaque proposition est confrontée au corpus : soutiens, nuances, contradictions et insuffisances "
         "documentaires (tableau 5.5). Aucune n'est présentée comme validée avant l'analyse.",
         size=9.4, color=MUTED, first=True, line=1.3)
    footer(s, "plan harmonisé § 3.6 ; tableau 3.2 (propositions et mécanismes attendus) ; "
              "annexe G (matrice d'alignement).", 10)
    notes(s, "Insister sur le mot « proposition » : en recherche qualitative abductive, on ne teste pas des hypothèses, "
             "on confronte des propositions. Question probable — « P4 n'est-elle pas tautologique ? » : "
             "réponse dans le protocole de codage, on observe des scènes, pas des étiquettes.")


def slide_mechanisms(prs):
    s = new_slide(prs, "Sept mécanismes")
    y = head(s, "Chapitre 3 · § 3.5", "Sept mécanismes à observer dans les scènes de décision",
             "L'articulation n'est pas un état : c'est une chaîne d'opérations cognitives, relationnelles et organisationnelles.")
    items = [
        ("Détection et anticipation", "Voir venir : signaux faibles, scénarios, alertes."),
        ("Interprétation et sens", "Donner un sens situé aux données produites."),
        ("Délibération et arbitrage éthique", "Trancher entre exigences incommensurables."),
        ("Coordination et mobilisation", "Faire agir ensemble des acteurs aux logiques distinctes."),
        ("Suivi et apprentissage", "Capitaliser : ce qui a marché, ce qui a échoué, pourquoi."),
        ("Confiance et transparence", "Rendre identifiable ce que fait la machine, et assumer le reste."),
        ("Prévention des effets défavorables", "Repérer dépendance, biais, perte de jugement."),
    ]
    gap = 0.18
    cw = (CW - 3 * gap) / 4
    for i, (t, b) in enumerate(items):
        x = M + (i % 4) * (cw + gap)
        yy = y + (i // 4) * 1.74
        rect(s, x, yy, cw, 1.56, CARD, LINE, radius=0.07)
        tf = textbox(s, x + 0.17, yy + 0.1, cw - 0.34, 1.36, anchor=MSO_ANCHOR.TOP).text_frame
        para(tf, f"{i + 1:02d}", size=12.5, color=BRONZE2, font=SERIF, first=True, space_after=4)
        para(tf, t, size=10.2, color=INK, bold=True, line=1.16, space_after=4)
        para(tf, b, size=9.4, color=MUTED, line=1.28)

    # 8e case : usage analytique
    x = M + 3 * (cw + gap)
    yy = y + 1.74
    rect(s, x, yy, cw, 1.56, BRONZE_DIM, BRONZE, radius=0.07)
    tt = textbox(s, x + 0.17, yy + 0.1, cw - 0.34, 1.36, anchor=MSO_ANCHOR.TOP).text_frame
    para(tt, "Usage analytique", size=8.6, color=BRONZE2, bold=True, caps=True, tracking=0.6,
         first=True, space_after=6)
    para(tt, "Ces sept mécanismes servent de grille de lecture pour coder les scènes : ils rendent le modèle observable.",
         size=9.4, color=INK2, line=1.3)
    footer(s, "plan harmonisé § 3.5 (mécanismes attendus) ; dictionnaire de codes et arborescence NVivo : annexe F.", 11)
    notes(s, "Diapositive la plus « méthode » du bloc conceptuel : elle prépare le codage. "
             "Question probable — « Comment passez-vous du concept au code ? » : annexe F et annexe G, alignement explicite.")


def slide_design(prs):
    s = new_slide(prs, "Posture et design")
    y = head(s, "Chapitre 4 · § 4.1–4.2",
             "Une posture interprétativiste, une logique abductive, un design multi-cas")
    gap = 0.2
    cw = (CW - 2 * gap) / 3
    cols = [
        ("Positionnement", "Interprétativisme à sensibilité constructiviste : la gouvernance est saisie à travers "
                           "le sens que les acteurs donnent à leurs décisions.",
         "Conséquence : pas de mesure d'efficacité, mais une analyse de processus et de conditions."),
        ("Logique", "Abduction : aller-retour entre cadre théorique et terrain, sans présumer la direction des effets.",
         "Les propositions peuvent être nuancées, déplacées ou abandonnées en cours d'analyse."),
        ("Design", "Étude de cas multiples : cinq EPES (EPES-A à EPES-E), analyses intra-cas puis inter-cas.",
         "Objectif : variation et réplication théorique, non représentativité statistique."),
    ]
    for i, (t, b, s2) in enumerate(cols):
        x = M + i * (cw + gap)
        rect(s, x, y, cw, 3.7, CARD, LINE, radius=0.05)
        tf = textbox(s, x + 0.2, y + 0.2, cw - 0.4, 3.3).text_frame
        para(tf, t, size=8.8, color=BRONZE2, bold=True, caps=True, tracking=0.7, first=True, space_after=8)
        para(tf, b, size=11.0, color=INK2, line=1.34)
        sep = rect(s, x + 0.2, y + 2.42, cw - 0.4, 0.008, LINE)
        rgb_alpha(sep, 0.5)
        para(tf, s2, size=9.2, color=MUTED, space_before=12, line=1.3)
    footer(s, "plan harmonisé § 4.1–4.2 ; tableau 4.1 (positionnement épistémologique et cohérence "
              "du dispositif méthodologique).", 12)
    notes(s, "Assumer la posture : elle est cohérente avec l'objet (processus, sens, conditions), pas un repli faute "
             "de données quantitatives. Question probable — « Quelle généralisation ? » : transférabilité analytique (§ 6.9).")


def slide_cases(prs):
    s = new_slide(prs, "Cas et échantillon")
    y = head(s, "Chapitre 4 · § 4.3", "Cinq cas choisis sur cinq critères, quarante-cinq participants prévus")
    lw = CW * 0.45
    tf = textbox(s, M, y + 0.05, lw, 3.9).text_frame
    para(tf, "CRITÈRES DE SÉLECTION", size=8.6, color=BRONZE2, bold=True, tracking=0.8, first=True, space_after=8)
    for c, lab in [
        ("C1", "ancienneté et trajectoire institutionnelle"),
        ("C2", "diversité de l'offre académique"),
        ("C3", "configuration de gouvernance et modèle de financement"),
        ("C4", "taille, implantation et partenariats"),
        ("C5", "accessibilité du terrain et richesse des données"),
    ]:
        rich(tf, [(c + " — ", BRONZE2, True), (lab, INK2, False)], size=10.4, bullet="■", space_before=7, line=1.28)
    para(tf, "Chiffres prévisionnels : le corpus effectif et les écarts au protocole seront présentés au "
             "chapitre 5 (tableau 5.1).", size=9.2, color=MUTED, space_before=14, line=1.3)

    rx = M + lw + 0.3
    rw = CW - lw - 0.3
    sw = (rw - 0.18) / 2
    stat(s, rx, y, sw, 1.02, "5", "établissements", "EPES-A à EPES-E")
    stat(s, rx + sw + 0.18, y, sw, 1.02, "45", "participants prévisionnels", "échantillon ciblé")
    rows = [
        ("Catégorie", "Effectif", "Code"),
        ("Dirigeants", "15", "DIR"),
        ("Acteurs académiques et enseignants impliqués dans les décisions", "20", "ACA"),
        ("Responsables administratifs, humains, numériques, qualité ou RSE", "10", "ADM · RH · NUM-QRSE"),
        ("Total prévisionnel", "45", ""),
    ]
    table(s, rx, y + 1.2, rw, [4.4, 0.9, 1.9], rows, row_h=0.52, size=9.6)
    para(textbox(s, rx, y + 1.2 + 0.52 * 5 + 0.14, rw, 0.5).text_frame,
         "Recrutement, saturation et suffisance informationnelle documentés (annexes C et H).",
         size=8.8, color=MUTED, first=True, line=1.25)
    footer(s, "plan harmonisé § 4.3 ; tableaux 4.2 (critères et dimensions de variation) et 4.3 (cas pressentis, "
              "échantillonnage prévisionnel) ; annexe E (tableau pseudonymisé).", 13)
    notes(s, "Redire « prévisionnel » : c'est un engagement de protocole, pas un résultat. "
             "Question probable — « Pourquoi 45 et pas 30 ou 60 ? » : suffisance informationnelle visée par "
             "catégories, révisable selon la saturation.")


def slide_collect(prs):
    s = new_slide(prs, "Collecte et analyse")
    y = head(s, "Chapitre 4 · § 4.4 et 4.6", "Cinq instruments, une analyse thématique hybride")
    hw = (CW - 0.24) / 2
    card(s, M, y, hw, 2.55, title="Collecte", bullets=[
        "Entretiens semi-directifs — guide harmonisé de 14 questions, avec relances (annexe A).",
        "Observation directe à dominante non participante des scènes de décision.",
        "Analyse documentaire — documents institutionnels autorisés.",
        "Journal de terrain et mémos réflexifs.",
    ], body_size=10.0)
    card(s, M + hw + 0.24, y, hw, 2.55, title="Analyse", bullets=[
        "Codage hybride : déductif (cadre, propositions) et inductif (thèmes émergents).",
        "Intra-cas (monographies) puis inter-cas (configurations).",
        "Triangulation des sources et des moments de collecte.",
        "Cas négatifs et explications rivales traités explicitement.",
    ], body_size=10.0)

    y2 = y + 2.78
    gap = 0.18
    cw = (CW - 2 * gap) / 3
    for i, (t, b) in enumerate([
        ("Outillage", "NVivo : arborescence de codes versionnée (annexe F)."),
        ("Traçabilité", "Matrices de saturation, de triangulation, d'analyse comparative (annexe H)."),
        ("Restitution", "Verbatims authentiques, anonymisés et contextualisés (annexe I) — jamais reconstitués."),
    ]):
        card(s, M + i * (cw + gap), y2, cw, 1.22, title=t, body=b, tight=True, body_size=9.6)
    footer(s, "plan harmonisé § 4.4, § 4.6 ; tableau 4.4 (sources, instruments, triangulation) ; "
              "annexes A, D, F, H, I.", 14)
    notes(s, "Point fort du protocole : triangulation et traitement des cas négatifs — ce qui protège des conclusions "
             "trop favorables. Question probable — « Comment évitez-vous de trouver ce que vous cherchez ? » : "
             "codage hybride, cas négatifs, explications rivales, mémos.")


def slide_ethics(prs):
    s = new_slide(prs, "Rigueur et éthique")
    y = head(s, "Chapitre 4 · § 4.7–4.8", "Rigueur scientifique et protection des participants")
    hw = (CW - 0.24) / 2
    card(s, M, y, hw, 2.5, title="Critères de rigueur", bullets=[
        "Crédibilité : triangulation, restitution aux acteurs.",
        "Transférabilité : description dense du contexte et des cas.",
        "Fiabilité : protocole tracé, versions datées, journal de terrain.",
        "Confirmabilité : chaîne de preuves auditable, mémos réflexifs.",
    ], body_size=10.2)
    card(s, M + hw + 0.24, y, hw, 2.5, title="Éthique et confidentialité", bullets=[
        "Codes EPES-A à EPES-E en version publique ; participants pseudonymisés.",
        "Consentement libre et éclairé, note d'information, droit de retrait (annexe B).",
        "Dossier sécurisé : table de correspondance, formulaires signés, enregistrements — jamais en ligne.",
        "Verbatims publiés uniquement après anonymisation et contextualisation.",
    ], body_size=10.2)
    card(s, M, y + 2.72, CW, 1.05, title="Double posture assumée",
         body="Praticien-chercheur : la réflexivité sur cette position fait partie du dispositif méthodologique (§ 4.9), "
              "elle n'est pas traitée comme un biais résiduel passé sous silence.",
         accent=True, body_size=10.6, tight=True)
    footer(s, "plan harmonisé § 4.7–4.9 ; tableau 4.5 (rigueur, risques, protection des participants) ; "
              "annexes B, E, J.", 15)
    notes(s, "Anticiper la question classique du jury sur la double posture : la nommer avant qu'elle soit posée. "
             "Question probable — « Votre position de dirigeant ne biaise-t-elle pas les entretiens ? » : "
             "réflexivité, tiers, cas hors UPL, explications rivales.")


def slide_architecture(prs):
    s = new_slide(prs, "Architecture")
    y = head(s, "Architecture", "Trois parties, six chapitres, 77 sections numérotées")
    gap = 0.22
    cw = (CW - 2 * gap) / 3
    parts = [
        ("Partie I · Contexte et fondements", [
            ("Chapitre 1 · 10 §", "Les EPES gabonais face à la polycrise : cadre institutionnel, gouvernance, financement, partenariats, ancrage territorial."),
            ("Chapitre 2 · 10 §", "Revue critique : gouvernance, IA, IE, performance durable, RSE ; synthèse des lacunes."),
        ]),
        ("Partie II · Cadre et méthode", [
            ("Chapitre 3 · 9 §", "Modèle exploratoire, sept propositions, opérationnalisation, alignement des instruments."),
            ("Chapitre 4 · 11 §", "Méthodologie : posture, design multi-cas, instruments, analyse, rigueur, éthique, limites."),
        ]),
        ("Partie III · Résultats et discussion", [
            ("Chapitre 5 · 10 §", "Présentation et analyse des résultats — architecture préparatoire, subordonnée au terrain réel."),
            ("Chapitre 6 · 11 §", "Discussion, réexamen du modèle, contributions, recommandations, limites, perspectives."),
        ]),
    ]
    for i, (part, chapters) in enumerate(parts):
        x = M + i * (cw + gap)
        tag = textbox(s, x, y, cw, 0.25).text_frame
        para(tag, part, size=7.8, color=MUTED, bold=True, caps=True, tracking=0.8, first=True)
        for j, (t, b) in enumerate(chapters):
            yy = y + 0.34 + j * 1.26
            rect(s, x, yy, cw, 1.12, CARD, LINE, radius=0.06)
            tt = textbox(s, x + 0.16, yy + 0.08, cw - 0.32, 0.96, anchor=MSO_ANCHOR.MIDDLE).text_frame
            para(tt, t, size=9.6, color=BRONZE2, bold=True, first=True, space_after=4)
            para(tt, b, size=9.0, color=INK2, line=1.28)

    yflow = y + 3.34
    flow = ["Introduction générale · 8 §", "Chapitres 1 à 6", "Conclusion · 8 points",
            "Bibliographie", "Annexes A à K"]
    x = M
    for i, label in enumerate(flow):
        w = 2.15 if i != 4 else 1.9
        fill = BRONZE_DIM if i == 4 else CARD
        border = BRONZE if i == 4 else LINE
        rect(s, x, yflow, w, 0.4, fill, border, radius=0.2)
        para(textbox(s, x + 0.08, yflow + 0.11, w - 0.16, 0.24).text_frame, label, size=7.6,
             color=BRONZE2 if i == 4 else MUTED, bold=True, align=PP_ALIGN.CENTER, first=True)
        x += w
        if i < len(flow) - 1:
            para(textbox(s, x + 0.04, yflow + 0.1, 0.24, 0.24).text_frame, "→", size=9,
                 color=MUTED2, first=True)
            x += 0.3
    rect(s, M, yflow + 0.52, CW, 0.62, CARD, LINE, radius=0.08)
    para(textbox(s, M + 0.18, yflow + 0.58, CW - 0.36, 0.5, anchor=MSO_ANCHOR.MIDDLE).text_frame,
         "Un seul fil : contexte (1–2) → cadre et méthode (3–4) → résultats et discussion (5–6). "
         "Aucun chapitre autonome de « réexamen du modèle » : la révision est intégrée au chapitre 6.",
         size=9.0, color=MUTED, first=True, line=1.3)
    footer(s, "plan général harmonisé (septembre 2026) — décompte effectué sur le plan lui-même : "
              "8 + (10 + 10 + 9 + 11 + 10 + 11) + 8 = 77.", 16)
    notes(s, "Montrer la symétrie : contexte → cadre → résultats. Le chiffre 77 vient du plan, pas d'une estimation. "
             "Question probable — « Pourquoi la révision du modèle n'est-elle pas un chapitre ? » : pour éviter la "
             "redondance avec les chapitres 5 et 6.")


def slide_supports(prs):
    s = new_slide(prs, "Tableaux et figures")
    y = head(s, "Supports · liste unique synthétisée", "De 115 supports de travail à 40 supports retenus",
             "Une réorganisation éditoriale (≈ 65 % de réduction), non une modification du protocole.")
    labels = ["Introduction", "Chapitre 1", "Chapitre 2", "Chapitre 3", "Chapitre 4", "Chapitre 5", "Chapitre 6"]
    before = [4, 10, 17, 17, 22, 25, 20]
    after = [1, 5, 6, 6, 7, 9, 6]
    x0, bar_x = M, M + 2.55
    max_w = 4.55
    scale = max_w / 25
    step = 0.66
    base = y + 0.12
    for i, lab in enumerate(labels):
        yy = base + i * step
        para(textbox(s, x0, yy + 0.14, 2.4, 0.28).text_frame, lab, size=9.0, color=MUTED,
             align=PP_ALIGN.RIGHT, first=True)
        wb = before[i] * scale
        rb = rect(s, bar_x, yy + 0.02, wb, 0.2, RGBColor(0x3E, 0x4D, 0x5C))
        para(textbox(s, bar_x + wb + 0.1, yy - 0.01, 0.6, 0.26).text_frame, str(before[i]), size=8.4,
             color=MUTED, font=MONO, first=True)
        wa = after[i] * scale
        rect(s, bar_x, yy + 0.3, wa, 0.2, BRONZE)
        para(textbox(s, bar_x + wa + 0.1, yy + 0.27, 0.6, 0.26).text_frame, str(after[i]), size=8.6,
             color=BRONZE2, bold=True, font=MONO, first=True)
    ly = base + 7 * step + 0.1
    rect(s, bar_x, ly, 0.16, 0.12, RGBColor(0x3E, 0x4D, 0x5C))
    para(textbox(s, bar_x + 0.24, ly - 0.05, 2.2, 0.24).text_frame, "Programme de travail (115)", size=8.2,
         color=MUTED, first=True)
    rect(s, bar_x + 2.5, ly, 0.16, 0.12, BRONZE)
    para(textbox(s, bar_x + 2.74, ly - 0.05, 2.2, 0.24).text_frame, "Supports retenus (40)", size=8.2,
         color=MUTED, first=True)

    rx = M + 7.75
    rw = CW - 7.75
    stat(s, rx, y, (rw - 0.18) / 2, 0.98, "27", "tableaux", "série « Tableau »")
    stat(s, rx + (rw - 0.18) / 2 + 0.18, y, (rw - 0.18) / 2, 0.98, "13", "figures", "série « Figure »")
    tf = textbox(s, rx, y + 1.2, rw, 2.6).text_frame
    for i, b in enumerate([
        "Un tableau compare, classe ou documente ; une figure montre des relations ou un processus.",
        "Fusions ciblées, sans fusionner des concepts distincts.",
        "40 supports = une cible éditoriale, pas un quota.",
    ]):
        para(tf, b, size=9.6, color=INK2, bullet="■", first=(i == 0), space_before=0 if i == 0 else 8, line=1.28)
    para(tf, "Trois représentations à ne pas confondre : figure 3.1 (modèle exploratoire), figure 5.3 "
             "(synthèse empirique), figure 6.1 (modèle révisé).", size=8.8, color=MUTED, space_before=12, line=1.3)
    footer(s, "liste unique synthétisée des tableaux et des figures (septembre 2026) ; base de travail : "
              "115 entrées (76 tableaux, 39 figures).", 17)
    notes(s, "Assumer la réduction : un jury reproche plus souvent la surcharge que la sobriété. "
             "Question probable — « Pourquoi avoir supprimé 75 supports ? » : doublons, lisibilité, redondance "
             "entre tableaux de synthèse et texte.")


def slide_status(prs):
    s = new_slide(prs, "État d'avancement")
    y = head(s, "État du travail · 22 octobre 2026", "Ce qui est stabilisé, ce qui reste préparatoire",
             "Un point d'étape sans embellissement : le jury évalue une démarche, pas une promesse.")
    gap = 0.2
    cw = (CW - 2 * gap) / 3
    card(s, M, y, cw, 3.5, title="Stabilisé", bullets=[
        "Plan général harmonisé (3 parties, 6 chapitres, 77 sections).",
        "Question centrale, QR1–QR4, P1–P7.",
        "Architecture méthodologique et éthique.",
        "Liste unique des supports (27 tableaux, 13 figures).",
        "Bibliographie en cours d'harmonisation.",
    ], body_size=9.8)
    card(s, M + cw + gap, y, cw, 3.5, title="En préparation", bullets=[
        "Chapitre 5 : architecture d'analyse arrêtée, contenus subordonnés au corpus réel.",
        "Chapitre 6 : discussion et recommandations indicatives.",
        "Protocole de cas : grilles et saturation à finaliser.",
    ], body_size=9.8)
    x3 = M + 2 * (cw + gap)
    card(s, x3, y, cw, 3.5, title="Règle tenue",
         body="Aucune donnée de terrain n'est inventée, aucun verbatim n'est reconstitué, aucune référence "
              "bibliographique n'est citée sans avoir été lue. Un trou reste un trou : il est signalé comme tel.",
         accent=True, body_size=10.4)
    para(textbox(s, x3 + 0.19, y + 2.85, cw - 0.38, 0.5).text_frame,
         "Cette règle s'applique aussi au présent diaporama et à ses visuels.",
         size=8.6, color=MUTED, first=True, line=1.25)
    footer(s, "dossier de travail de l'auteur, 8 septembre – 22 octobre 2026 ; plan harmonisé ; liste des supports.", 18)
    notes(s, "Diapositive de sincérité : elle protège la soutenance, le jury sait ce qui est attendu de lui. "
             "Question probable — « Quand le terrain est-il réalisé ? » : réponse de l'auteur, à préparer en amont ; "
             "ne rien inventer ici.")


def slide_ch5(prs):
    s = new_slide(prs, "Chapitre 5")
    y = head(s, "Chapitre 5 · architecture préparatoire",
             "Les résultats seront organisés par scènes, cas et propositions",
             "Les intitulés indiquent des objets d'analyse, non des résultats acquis.")
    gap = 0.2
    cw = (CW - gap) / 2
    items = [
        ("5.1 – 5.2 · Corpus et monographies", "Corpus effectif, écarts au protocole, suffisance informationnelle ; analyse intra-cas sous codes."),
        ("5.3 – 5.4 · Polycrise et IA", "Configurations de gouvernance actuelles ; usages numériques et contribution de l'intelligence artificielle."),
        ("5.5 – 5.6 · IE et articulation", "Dimensions humaines de la décision ; articulation IA–IE dans les processus de gouvernance."),
        ("5.7 – 5.8 · Effets et inter-cas", "Effets perçus ou documentés sur la performance, la RSE et la durabilité ; thèmes émergents et configurations."),
    ]
    for i, (t, b) in enumerate(items):
        x = M + (i % 2) * (cw + gap)
        yy = y + (i // 2) * 1.05
        card(s, x, yy, cw, 0.92, title=t, body=b, tight=True, body_size=9.8, title_size=8.2)
    y2 = y + 2.22
    card(s, M, y2, cw, 1.2, title="5.9 · Confrontation des sept propositions",
         body="Soutiens, conditions, nuances, contradictions, explications rivales et insuffisances documentaires — "
              "présentés avec la même exigence, y compris lorsqu'ils affaiblissent le modèle.",
         accent=True, body_size=10.2)
    card(s, M + cw + gap, y2, cw, 1.2, title="Règle d'écriture du chapitre",
         body="Le chapitre ne présuppose ni l'efficacité de l'IA, ni la supériorité d'une configuration de "
              "gouvernance. Une articulation absente reste un résultat.", body_size=10.2)
    footer(s, "plan harmonisé, chapitre 5 (§ 5.1–5.10) ; tableaux 5.1 à 5.6 ; figures 5.1 à 5.3.", 19)
    notes(s, "« Ce chapitre est une architecture : je présente la structure d'analyse, pas des conclusions. » "
             "Question probable — « Comment gérerez-vous un cas où aucune IA n'est utilisée ? » : résultat analysable, "
             "comparaison inter-cas et explications rivales.")


def slide_ch6(prs):
    s = new_slide(prs, "Chapitre 6")
    y = head(s, "Chapitre 6 · version indicative", "Discuter, réviser, recommander — sans jamais surpromettre")
    tf = textbox(s, M, y + 0.05, CW * 0.56, 4.4).text_frame
    for i, chunks in enumerate([
        [("6.1–6.3", BRONZE2, True), (" — Discussion des résultats et des sept propositions au regard de la littérature.", INK2, False)],
        [("6.4", BRONZE2, True), (" — Réexamen du modèle : modèle exploratoire, modèle empirique, outil d'action ; règles de révision ; ", INK2, False), ("issue possible en l'absence de gouvernance augmentée documentée.", INK, True)],
        [("6.5", BRONZE2, True), (" — Contributions scientifiques et managériales susceptibles d'être dégagées.", INK2, False)],
        [("6.6–6.8", BRONZE2, True), (" — Recommandations managériales et institutionnelles, proportionnées aux besoins, ressources et risques documentés ; feuille de route et outils de pilotage.", INK2, False)],
        [("6.9–6.11", BRONZE2, True), (" — Limites, portée des interprétations, perspectives.", INK2, False)],
    ]):
        rich(tf, chunks, size=12.4, bullet="■", first=(i == 0), space_before=0 if i == 0 else 17, line=1.34)

    rw = CW * 0.42
    rx = M + CW - rw
    card(s, rx, y, rw, 2.2, title="La clause de prudence, énoncée d'emblée",
         body="« La gouvernance augmentée est envisagée comme une proposition à examiner empiriquement, et non "
              "comme une solution dont l'efficacité serait acquise. »", accent=True, body_size=9.8)
    para(textbox(s, rx + 0.2, y + 1.78, rw - 0.4, 0.3).text_frame,
         "Formulation du plan général harmonisé, reprise dans le chapitre 6.",
         size=7.8, color=MUTED2, first=True)
    card(s, rx, y + 2.42, rw, 1.55, title="Trois figures, trois statuts",
         body="3.1 modèle exploratoire (avant terrain) · 5.3 synthèse empirique (issue du corpus) · 6.1 modèle révisé "
              "(discussion). Elles ne sont jamais fusionnées.", body_size=9.6)
    footer(s, "plan harmonisé, chapitre 6 (§ 6.1–6.11) ; § 6.4 (issue possible en l'absence de gouvernance "
              "augmentée) ; tableaux 6.1 à 6.4 et figures 6.1–6.2.", 20)
    notes(s, "La clause de prudence est un argument de crédibilité, pas une précaution oratoire. "
             "Question probable — « Et si vous ne trouvez rien ? » : c'est prévu au § 6.4 ; l'absence documentée "
             "est un résultat.")


def slide_contrib(prs):
    s = new_slide(prs, "Contributions")
    y = head(s, "Contributions · conditionnelles",
             "Trois contributions recherchées — au conditionnel, jusqu'à l'analyse")
    gap = 0.2
    cw = (CW - 2 * gap) / 3
    items = [
        ("Scientifique", "Proposer un cadre qui traite l'articulation IA–IE comme un processus de gouvernance "
                         "observable, et non comme une vertu du dirigeant.",
         "Apport théorique visé : un modèle exploratoire situé, révisable."),
        ("Managériale", "Outiller les dirigeants et instances d'EPES : séquences de décision, points de contrôle, "
                        "conditions de prudence dans les usages décisionnels de l'IA.",
         "Apport visé : recommandations proportionnées, pas une méthode universelle."),
        ("Sociétale", "Éclairer la contribution des EPES au développement durable et à la RSE — au Gabon et, plus "
                      "largement, en Afrique centrale.",
         "Apport visé : un langage commun entre gouvernance, qualité et durabilité."),
    ]
    for i, (t, b, s2) in enumerate(items):
        x = M + i * (cw + gap)
        rect(s, x, y, cw, 3.3, CARD, LINE, radius=0.05)
        tf = textbox(s, x + 0.2, y + 0.2, cw - 0.4, 2.9).text_frame
        para(tf, t, size=8.8, color=BRONZE2, bold=True, caps=True, tracking=0.7, first=True, space_after=8)
        para(tf, b, size=10.8, color=INK2, line=1.34)
        sep = rect(s, x + 0.2, y + 2.32, cw - 0.4, 0.008, LINE)
        rgb_alpha(sep, 0.5)
        para(tf, s2, size=9.0, color=MUTED, space_before=12, line=1.3)
    para(textbox(s, M, y + 3.52, CW, 0.7).text_frame,
         "Ces contributions ne seront présentées comme acquises qu'au chapitre 6, après analyse du corpus. "
         "Le protocole prévoit explicitement le cas où elles ne seraient que partielles.",
         size=9.2, color=MUTED, first=True, line=1.3)
    footer(s, "plan harmonisé § 6.5 ; introduction § 1.6 (intérêts scientifique, managérial et sociétal).", 21)
    notes(s, "Le conditionnel n'est pas une faiblesse : c'est la signature d'une thèse qui ne confond pas "
             "intention et résultat.")


def slide_limits(prs):
    s = new_slide(prs, "Limites")
    y = head(s, "Limites · § 6.9", "Limites assumées et conditions de transférabilité")
    hw = (CW - 0.24) / 2
    card(s, M, y, hw, 3.9, title="Limites de la recherche", bullets=[
        "Portée : cinq cas, un secteur, un pays — pas de généralisation statistique.",
        "Données déclaratives : risque d'écart entre discours et pratiques, traité par triangulation.",
        "Temporalité : la polycrise est un processus ; l'enquête en saisit une fenêtre.",
        "Posture : praticien-chercheur, avec les effets de position que cela implique.",
        "Accès : confidentialité et disponibilité des acteurs contraignent le corpus.",
    ], body_size=10.3)
    x2 = M + hw + 0.24
    card(s, x2, y, hw, 3.9, title="Ce que la thèse peut transférer", bullets=[
        "Une grille d'analyse des scènes de décision instrumentées.",
        "Des conditions de possibilité, pas des recettes.",
        "Un vocabulaire commun entre gouvernance, usages de l'IA et dimensions humaines.",
        "Une méthode réutilisable dans d'autres EPES d'Afrique centrale.",
    ], body_size=10.3)
    sep = rect(s, x2 + 0.19, y + 3.15, hw - 0.38, 0.008, LINE)
    rgb_alpha(sep, 0.5)
    para(textbox(s, x2 + 0.19, y + 3.28, hw - 0.38, 0.5).text_frame,
         "Transférabilité analytique : le lecteur juge de la pertinence du transfert, la thèse documente le contexte.",
         size=8.6, color=MUTED, first=True, line=1.25)
    footer(s, "plan harmonisé § 6.9 et § 4.10 (limites méthodologiques et stratégies d'atténuation) ; annexe J.", 22)
    notes(s, "Prendre les limites une par une, sans les minimiser : le jury attend cette diapositive. "
             "Question probable — « Comment atténuez-vous le risque déclaratif ? » : triangulation "
             "entretiens / observation / documents.")


def slide_perspectives(prs):
    s = new_slide(prs, "Perspectives")
    y = head(s, "Perspectives · § 6.10", "Après la thèse : trois prolongements possibles")
    gap = 0.2
    cw = (CW - 2 * gap) / 3
    items = [
        ("1", "Élargir le terrain", "Comparer les EPES gabonais à d'autres établissements d'Afrique centrale : conditions institutionnelles contrastées, mêmes questions."),
        ("2", "Suivre dans le temps", "Reprendre les mêmes scènes de décision à intervalles réguliers : ce qui se maintient, ce qui s'abandonne, ce qui se transforme."),
        ("3", "Outiller et évaluer", "Tester la grille comme outil de pilotage en situation réelle, avec évaluation des effets — y compris des effets défavorables."),
    ]
    for i, (n, t, b) in enumerate(items):
        x = M + i * (cw + gap)
        rect(s, x, y, cw, 2.5, CARD, LINE, radius=0.05)
        nn = textbox(s, x + 0.2, y + 0.18, 0.6, 0.4).text_frame
        para(nn, n, size=15, color=BRONZE2, font=SERIF, first=True)
        tt = textbox(s, x + 0.2, y + 0.72, cw - 0.4, 0.35).text_frame
        para(tt, t, size=10.0, color=INK, bold=True, first=True)
        bb = textbox(s, x + 0.2, y + 1.12, cw - 0.4, 1.2).text_frame
        para(bb, b, size=9.6, color=INK2, first=True, line=1.3)
    card(s, M, y + 2.72, CW, 1.1, title="Ce que ces prolongements supposent",
         body="Un ancrage institutionnel stable, des autorisations de terrain, et la possibilité de publier des "
              "résultats sans exposer les établissements ni les personnes. La confidentialité n'est pas une contrainte "
              "technique : c'est une condition de la recherche.", accent=True, body_size=10.4, tight=True)
    footer(s, "plan harmonisé § 6.10 (perspectives de recherche et prolongements possibles) ; "
              "§ 4.8 (considérations éthiques).", 23)
    notes(s, "Terminer le bloc scientifique sur une ouverture maîtrisée : trois pistes, aucune promesse.")


def slide_numbers(prs):
    s = new_slide(prs, "Repères chiffrés")
    y = head(s, "Repères chiffrés", "Huit chiffres, tous traçables dans le dossier de recherche")
    gap = 0.16
    cw = (CW - 3 * gap) / 4
    tiles = [
        ("3", "parties", "contexte · cadre · résultats"),
        ("6", "chapitres", "plan harmonisé"),
        ("77", "sections numérotées", "8 + 61 + 8"),
        ("4", "questions spécifiques", "QR1 – QR4"),
        ("7", "propositions", "P1 – P7"),
        ("40", "supports retenus", "27 tableaux · 13 figures"),
        ("5", "cas", "EPES-A à EPES-E"),
        ("45", "participants prévus", "15 · 20 · 10"),
    ]
    for i, (v, k, sub) in enumerate(tiles):
        stat(s, M + (i % 4) * (cw + gap), y + (i // 4) * 1.24, cw, 1.1, v, k, sub)
    y2 = y + 2.66
    gap3 = 0.18
    cw3 = (CW - 2 * gap3) / 3
    for i, (t, b) in enumerate([
        ("115", "supports dans le programme de travail initial (76 tableaux, 39 figures)."),
        ("14", "questions du guide d'entretien semi-directif harmonisé (annexe A)."),
        ("11", "annexes, de A à K, dont le dictionnaire de codes et les matrices de saturation."),
    ]):
        x = M + i * (cw3 + gap3)
        rect(s, x, y2, cw3, 1.22, CARD, LINE, radius=0.07)
        tf = textbox(s, x + 0.2, y2 + 0.12, cw3 - 0.4, 1.0, anchor=MSO_ANCHOR.MIDDLE).text_frame
        para(tf, t, size=17, color=BRONZE2, font=SERIF, first=True, line=1.0, space_after=4)
        para(tf, b, size=8.6, color=MUTED, line=1.28)
    para(textbox(s, M, y2 + 1.36, CW, 0.6).text_frame,
         "Aucun de ces nombres n'est une estimation : ils sont décomptés sur le plan général harmonisé et la liste "
         "unique des supports (septembre 2026). Les effectifs de terrain restent prévisionnels jusqu'à la clôture du corpus.",
         size=8.8, color=MUTED, first=True, line=1.3)
    footer(s, "plan général harmonisé (architecture, questions, propositions, annexes) ; liste unique synthétisée "
              "des tableaux et figures ; § 4.3 pour l'échantillon prévisionnel.", 24)
    notes(s, "Diapositive de contrôle : chaque chiffre est vérifiable dans le dossier, elle coupe court aux approximations. "
             "Question probable — « D'où vient le 77 ? » : 8 sections d'introduction + 61 sections de chapitres + "
             "8 points de conclusion.")


def slide_next(prs):
    s = new_slide(prs, "Prochaines étapes")
    y = head(s, "Prochaines étapes", "Ce qui reste à produire, dans l'ordre des blocages")
    rows = [
        ("#", "Chantier", "Contenu attendu", "Statut"),
        ("1", "Protocole de cas (annexes C–D)", "Grilles d'observation, règles de saturation, conventions datées", "En cours"),
        ("2", "Dictionnaire de codes (annexe F)", "Codes déductifs et inductifs, arborescence NVivo, historique des versions", "À finaliser"),
        ("3", "Terrain et corpus (chapitre 5)", "Entretiens, observations, documents autorisés, journal de terrain", "À engager"),
        ("4", "Analyse et rédaction (chapitres 5–6)", "Analyses intra et inter-cas, confrontation des propositions, discussion", "Subordonné au corpus"),
        ("5", "Bibliographie harmonisée", "Références effectivement mobilisées, dédoublonnées, vérifiées", "En cours"),
    ]
    table(s, M, y, CW, [0.5, 3.3, 5.4, 1.9], rows, row_h=0.56, size=9.8)
    card(s, M, y + 0.56 * 6 + 0.26, CW, 1.0, title="Deux décisions qui n'appartiennent qu'à l'auteur",
         body="La validation du fond scientifique du modèle et l'engagement du terrain. Aucune de ces deux décisions "
              "ne peut être déléguée à un outil — ni à une intelligence artificielle.",
         accent=True, body_size=10.6, tight=True)
    footer(s, "dossier de travail de l'auteur (chantiers ouverts au 22 octobre 2026) ; plan harmonisé, annexes A–K.", 25)
    notes(s, "Annoncer le calendrier avec prudence, sans date inventée : dire ce qui est fait, ce qui vient, dans quel ordre.")


def slide_end(prs):
    s = new_slide(prs, "Clôture")
    y = head(s, "Pour conclure", "Une recherche qui outille la décision sans prétendre la remplacer")
    qh = 1.62
    rect(s, M, y, CW, qh, BRONZE_DIM, None)
    rect(s, M, y, 0.045, qh, BRONZE)
    tf = textbox(s, M + 0.34, y + 0.22, CW - 0.7, qh - 0.4).text_frame
    para(tf, "« La meilleure façon de prédire l'avenir, c'est de le créer. »", size=18.5,
         color=INK, font=SERIF, first=True, line=1.3)
    para(tf, "Peter Drucker — devise reprise par l'Université Privée de Libreville.",
         size=8.6, color=MUTED, space_before=6)

    y2 = y + qh + 0.34
    gap = 0.18
    cw = (CW - 2 * gap) / 3
    for i, (t, b) in enumerate([
        ("Merci", "Au jury, pour le temps et l'exigence."),
        ("Échanges", "contact@upl-gabon.com — objet « Recherche doctorale »."),
        ("Statut", "Document de travail, 22 octobre 2026. Aucun résultat n'est présumé."),
    ]):
        card(s, M + i * (cw + gap), y2, cw, 1.35, title=t, body=b, tight=True, body_size=10.2)

    para(textbox(s, M, y2 + 1.62, CW, 0.9).text_frame,
         "Serge Patrick MINANG — ingénieur · MBA · doctorant DBA · Président-Fondateur de l'Université Privée de "
         "Libreville. Cette recherche est conduite à titre personnel ; elle est distincte du site institutionnel de l'UPL.",
         size=9.2, color=MUTED, first=True, line=1.35)
    footer(s, "citation affichée sur upl-gabon.com (consulté le 22/10/2026) ; coordonnées publiques de l'UPL.", 26)
    notes(s, "Remercier, s'arrêter, laisser le jury poser ses questions. Ne pas meubler. "
             "Après la soutenance : consigner les remarques du jury, elles alimentent directement le chapitre 6.")


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)

    slide_cover(prs)
    slide_map(prs)
    slide_polycrisis(prs)
    slide_problem(prs)
    slide_question(prs)
    slide_qr(prs)
    slide_objectives(prs)
    slide_lit(prs)
    slide_model(prs)
    slide_props(prs)
    slide_mechanisms(prs)
    slide_design(prs)
    slide_cases(prs)
    slide_collect(prs)
    slide_ethics(prs)
    slide_architecture(prs)
    slide_supports(prs)
    slide_status(prs)
    slide_ch5(prs)
    slide_ch6(prs)
    slide_contrib(prs)
    slide_limits(prs)
    slide_perspectives(prs)
    slide_numbers(prs)
    slide_next(prs)
    slide_end(prs)

    prs.core_properties.title = "Soutenance DBA — Gouvernance augmentée des EPES gabonais"
    prs.core_properties.author = "Serge Patrick MINANG"
    prs.core_properties.subject = "Diaporama de soutenance doctorale (DBA)"
    prs.core_properties.keywords = "DBA, gouvernance augmentée, IA, intelligence émotionnelle, EPES, Gabon"
    prs.core_properties.comments = (
        "Document de travail du 22/10/2026. Chapitres 5 et 6 préparatoires : aucun résultat présumé. "
        "Provenance des chiffres : page « Sources » du dossier."
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    size_kb = OUT.stat().st_size / 1024
    print(f"écrit : {OUT.relative_to(ROOT)} — {len(prs.slides.__iter__.__self__._sldIdLst)} diapositives, {size_kb:.0f} Ko")
    if warnings:
        print("\nAvertissements de mise en page :")
        for w in warnings:
            print("  · " + w)
    else:
        print("aucun avertissement de mise en page")


if __name__ == "__main__":
    main()
