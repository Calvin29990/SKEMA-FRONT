#!/usr/bin/env python3
"""Remplit le modèle de l'auteur (modele/template-canva.pdf) — sans le redessiner.

Méthode
-------
1. Chaque page retenue du modèle est reprise telle quelle : on efface seulement le
   **texte** du PDF (photos, formes, dégradés, motifs, illustrations restent intacts)
   et l'image obtenue sert de fond de diapositive.
2. Les textes sont reposés en zones de texte PowerPoint **éditables**, exactement à la
   position, à la couleur et à la taille du modèle.
3. Seuls les passages qui portent les données de l'auteur sont réécrits (table OVR,
   repérée par le début du texte d'origine). La carte de l'Afrique de la page 2 est
   remplacée par la carte du Gabon.

Rien d'autre n'est déplacé, ajouté ni supprimé : le design reste celui du modèle.

Usage :  python3 outils/build-pptx-modele.py [--apercu DOSSIER]
Sortie : export/Soutenance-DBA-MINANG-modele-canva-18.pptx
"""

from __future__ import annotations

import json
import math
import re
import sys
from pathlib import Path

import pymupdf
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
PDF = ROOT / "modele" / "template-canva.pdf"
OUT = ROOT / "export" / "Soutenance-DBA-MINANG-modele-canva-18.pptx"
GEOJSON = ROOT / "data" / "gab_adm1.geojson"
CACHE = Path("/tmp/build")

PDF_W, PDF_H = 1440.0, 810.0
SLIDE_W, SLIDE_H = 13.333, 7.5
K = SLIDE_W / PDF_W                 # points PDF -> pouces
FT = K * 72.0                       # points PDF -> points PowerPoint
DPI = 155

# Les 18 pages du modèle, dans l'ordre du modèle (15 à 20 diapositives demandées)
PAGES = [1, 2, 4, 5, 6, 7, 8, 9, 10, 11, 12, 15, 16, 17, 19, 20, 21, 22]

DEJAVU = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
DEJAVU_B = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"


# --- Polices -------------------------------------------------------------------
def pptx_face(pdf_font: str) -> tuple[str, bool]:
    """Police du modèle -> police installée partout, même graisse."""
    if "GlockGrotesk" in pdf_font:
        return "Arial Black", False
    return "Arial", ("Bold" in pdf_font)


def metric_font(pdf_font: str) -> str:
    return DEJAVU_B if ("GlockGrotesk" in pdf_font or "Bold" in pdf_font) else DEJAVU


def slack(pdf_font: str) -> float:
    """Écart de largeur entre la police mesurée (DejaVu) et la police posée (Arial)."""
    if "GlockGrotesk" in pdf_font:
        return 1.10
    if "Bold" in pdf_font:
        return 1.16
    return 1.20


_metrics: dict[tuple[str, int], ImageFont.FreeTypeFont] = {}


def measure(text: str, size_pt: float, pdf_font: str) -> float:
    key = (metric_font(pdf_font), max(1, int(size_pt * 4)))
    if key not in _metrics:
        _metrics[key] = ImageFont.truetype(key[0], key[1])
    return _metrics[key].getlength(text) / 4.0


# --- Contenus réécrits ----------------------------------------------------------
QUESTION = (
    "« Comment et sous quelles conditions l'articulation entre l'intelligence artificielle et "
    "l'intelligence émotionnelle transforme-t-elle la gouvernance des EPES gabonais en situation "
    "de polycrise, et quels effets produit-elle sur leur performance globale, leur RSE et leur "
    "contribution au développement durable ? »"
)
QUESTION_COURTE = (
    "Comment l'articulation entre intelligence artificielle et intelligence émotionnelle "
    "transforme-t-elle la gouvernance des EPES gabonais en situation de polycrise ?"
)

# clé : page du modèle -> {début du texte d'origine : nouveau texte} ; « | » = paragraphe
OVR: dict[int, dict[str, str]] = {
    1: {"Doctorant :": "Doctorant : Serge Patrick MINANG|Soutenance de thèse doctorale (DBA)|"
                       "Directeur de thèse : [à compléter]|Date de soutenance : [à compléter]"},
    4: {"« Comment la gouvernance":
            "Comment l'articulation IA–IE transforme-t-elle la gouvernance des EPES gabonais en "
            "polycrise, et avec quels effets sur la performance, la RSE et la durabilité ?"},
    5: {"Comment la gouvernance augmentée, intégrant IA": QUESTION_COURTE,
        "Dans quelle mesure l'intelligence artificielle":
            "Comment les EPES gabonais gouvernent-ils en situation de polycrise : arbitrages, "
            "instances, pratiques de décision ?",
        "Comment l'intelligence émotionnelle des dirigeants":
            "Quels usages réels de l'intelligence artificielle observe-t-on dans les décisions "
            "de gouvernance ?",
        "Quelles sont les modalités d'articulation":
            "Quel rôle l'intelligence émotionnelle joue-t-elle dans le jugement et l'action "
            "collective des dirigeants ?",
        "Quels dispositifs de gouvernance augmentée":
            "À quelles conditions l'articulation IA–IE produit-elle des effets, et lesquels, sur "
            "la performance, la RSE et la durabilité ?"},
    10: {"Les Sept Propositions": "Les Sept Propositions|de Recherche (P1–P7)"},
    11: {"Étude menée auprès de 5 EPES":
             "Étude menée auprès de 5 EPES (établissements privés d'enseignement supérieur) au "
             "Gabon, avec 45 participants répartis en trois vagues (15 / 20 / 10).",
         "Entretiens semi-directifs individuels (30":
             "Entretiens semi-directifs individuels (45 participants) et observations directes en "
             "contexte organisationnel gabonais pour capter les dynamiques de décision."},
    12: {"Entretiens semi-": "Entretiens semi-directifs|(45 participants)"},
    22: {"Rimberio": "",
         "+241-XX-XX-XX-XX": "contact@upl-gabon.com|+241 02 62 19 78|Libreville, Gabon"},
}

# Numéros que le modèle répète par erreur (il affiche le numéro de page au lieu de 01)
NUM_FIX: dict[int, dict[str, str]] = {
    10: {"P1 :": "01"},
    17: {"Renforcement de la résilience sociétale": "01"},
    19: {"Limites identifiées": "01"},
    20: {"Originalité scientifique": "01"},
    21: {"Jensen & Meckling": "01"},
}

NOTES = {
    1: "Couverture. Annoncer le titre, se nommer, puis le plan en une phrase : contexte, question, "
       "cadre, modèle, méthode, contributions.",
    2: "Contexte : la polycrise gabonaise n'est pas une somme de crises, c'est leur interaction. "
       "Carte du Gabon, neuf provinces ; l'étude porte sur des EPES, non sur l'État.",
    3: "Question centrale : la lire lentement, c'est la phrase que le jury doit pouvoir répéter.",
    4: "QR1 à QR4 : chaque question ouvre un objet d'enquête distinct. Ne pas présupposer que l'IA "
       "améliore la gouvernance.",
    5: "Agence (contrôle, incitations) et parties prenantes (légitimité, confiance) : l'IA outille "
       "le contrôle, l'IE travaille la légitimité.",
    6: "Résilience et capacités dynamiques : détecter, absorber, reconfigurer. C'est le pont entre "
       "le cadre théorique et le modèle.",
    7: "Modèle conceptuel : quatre piliers, quatre effets attendus. Insister sur « attendus » : le "
       "terrain n'est pas clos.",
    8: "Articulation IA–IE : deux registres complémentaires, jamais substituables.",
    9: "Sept propositions, formulées pour être contredites. Si le terrain les infirme, c'est un "
       "résultat.",
    10: "Méthodologie : qualitatif multi-cas. 5 EPES, 45 participants (15 / 20 / 10), triangulation.",
    11: "Protocole : collecte, puis analyse et validation. Assumer le terrain en contexte de polycrise.",
    12: "Contributions scientifiques : elles restent conditionnelles jusqu'à l'analyse.",
    13: "Contributions managériales : outils de décision, compétences IE, coordination institutionnelle.",
    14: "Contributions sociétales : résilience collective et inclusion.",
    15: "Limites et perspectives : les énoncer soi-même désarme la question du jury.",
    16: "Conclusion : trois idées, puis s'arrêter.",
    17: "Bibliographie : les pères fondateurs et les auteurs de méthode, pas une liste exhaustive.",
    18: "Remercier le jury, s'arrêter, laisser venir les questions.",
}


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h.lstrip("#").upper())


# --- Lecture du modèle ----------------------------------------------------------
NUM_RE = re.compile(r"^(?:\d{1,2}|Q\d)$")


def read_blocks(page: pymupdf.Page) -> list[dict]:
    """Blocs du modèle : chaque ligne garde ses runs (texte, position, police, couleur)."""
    blocks = []
    for b in page.get_text("dict")["blocks"]:
        if b["type"] != 0:
            continue
        lines = []
        for l in b["lines"]:
            runs = []
            for s in l["spans"]:
                if not s["text"].strip():
                    continue
                runs.append({"t": s["text"], "x0": s["bbox"][0], "x1": s["bbox"][2],
                             "y0": s["bbox"][1], "y1": s["bbox"][3],
                             "size": round(s["size"], 2),
                             "font": s["font"].split("+")[-1],
                             "color": "#%06x" % s["color"]})
            if runs:
                lines.append(runs)
        if lines:
            blocks.append(lines)
    blocks.sort(key=lambda ls: (round(min(r["y0"] for l in ls for r in l) / 14),
                                min(r["x0"] for l in ls for r in l)))
    return blocks


def line_text(runs: list[dict]) -> str:
    out = ""
    for i, r in enumerate(runs):
        if i and not out.endswith(" ") and not r["t"].startswith(" "):
            out += " "
        out += r["t"]
    return out


def line_box(runs: list[dict]) -> tuple[float, float, float, float]:
    return (min(r["x0"] for r in runs), min(r["y0"] for r in runs),
            max(r["x1"] for r in runs), max(r["y1"] for r in runs))


def split_num(lines: list[dict]) -> tuple[dict | None, list[dict]]:
    """Sépare la pastille « 01 » / « Q1 » du texte qu'elle précède."""
    first = lines[0]
    if len(first) >= 2 and NUM_RE.match(first[0]["t"].strip()):
        numeral = first[0]
        rest = lines[0][1:]
        body = ([rest] if rest else []) + lines[1:]
        return numeral, body
    if len(first) == 1 and NUM_RE.match(first[0]["t"].strip()):
        return first[0], lines[1:]
    return None, lines


def block_text(lines: list[dict]) -> str:
    return re.sub(r"\s+", " ", " ".join(line_text(l) for l in lines)).strip()


def block_box(lines: list[dict]) -> tuple[float, float, float, float]:
    boxes = [line_box(l) for l in lines]
    return (min(b[0] for b in boxes), min(b[1] for b in boxes),
            max(b[2] for b in boxes), max(b[3] for b in boxes))


def lookup(table: dict[int, dict[str, str]], page: int, text: str) -> str | None:
    for prefix, value in table.get(page, {}).items():
        if text.strip().startswith(prefix):
            return value
    return None


def wrap(text: str, width: float, size: float, font: str) -> list[str]:
    out, cur = [], ""
    for word in text.split():
        test = (cur + " " + word).strip()
        if not cur or measure(test, size, font) <= width:
            cur = test
        else:
            out.append(cur)
            cur = word
    if cur:
        out.append(cur)
    return out


# --- Fonds de diapositive --------------------------------------------------------
LABELS = {
    "Estuaire": "Estuaire", "Haut-Ogooue": "Haut-Ogooué", "Moyen-Ogooue": "Moyen-Ogooué",
    "Ngounie": "Ngounié", "Nyanga": "Nyanga", "Ogooue-Ivindo": "Ogooué-Ivindo",
    "Ogooue-lolo": "Ogooué-Lolo", "Ogooue-Maritime": "Ogooué-Maritime",
    "Woleu-Ntem": "Woleu-Ntem",
}
NUDGE = {"Ogooué-Maritime": (-0.03, 0.02), "Ngounié": (0.01, 0.02),
         "Moyen-Ogooué": (0.012, 0.0)}


def carte_gabon_claire(size: tuple[int, int]) -> Image.Image:
    """Carte du Gabon (9 provinces) en version claire, aux couleurs du modèle."""
    data = json.loads(GEOJSON.read_text(encoding="utf-8"))
    polys: list[tuple[str, list, bool]] = []
    for f in data["features"]:
        name = LABELS.get(f["properties"]["ADM1_NAME"], f["properties"]["ADM1_NAME"])
        for poly in f["geometry"]["coordinates"]:
            outer = True
            for ring in poly:                       # 1er anneau = contour, suivants = trous
                polys.append((name, [(float(x), float(y)) for x, y in ring], outer))
                outer = False

    lons = [x for _, r, _ in polys for x, _ in r]
    lats = [y for _, r, _ in polys for _, y in r]
    minx, maxx, miny, maxy = min(lons), max(lons), min(lats), max(lats)
    k = math.cos(math.radians((miny + maxy) / 2))
    W, H = size
    pad = int(min(W, H) * 0.05)
    scale = min((W - 2 * pad) / ((maxx - minx) * k), (H - 2 * pad) / (maxy - miny))
    ox = (W - (maxx - minx) * k * scale) / 2
    oy = (H - (maxy - miny) * scale) / 2

    def proj(lon, lat):
        return ((lon - minx) * k * scale + ox, (maxy - lat) * scale + oy)

    img = Image.new("RGBA", (W, H), (238, 246, 251, 255))
    d = ImageDraw.Draw(img)
    centroids: dict[str, tuple[float, float, float]] = {}
    for name, ring, outer in polys:
        pts = [proj(x, y) for x, y in ring]
        fill = (127, 196, 238, 255) if name == "Estuaire" else (198, 222, 240, 255)
        d.polygon(pts, fill=fill if outer else (238, 246, 251, 255),
                  outline=(255, 255, 255, 255), width=max(2, W // 320))
        if not outer:
            continue
        area = cx = cy = 0.0
        for i in range(len(pts) - 1):
            x0, y0 = pts[i]
            x1, y1 = pts[i + 1]
            cross = x0 * y1 - x1 * y0
            area += cross
            cx += (x0 + x1) * cross
            cy += (y0 + y1) * cross
        if abs(area) > 1e-6:
            centroids[name] = (cx / (3 * area), cy / (3 * area), abs(area) / 2)

    f_lab = ImageFont.truetype(DEJAVU_B, max(9, W // 46))
    f_cit = ImageFont.truetype(DEJAVU_B, max(10, W // 40))
    cities = [("Libreville", 9.45, 0.39), ("Port-Gentil", 8.78, -0.72),
              ("Franceville", 13.58, -1.63)]
    city_pts = [proj(lon, lat) for _, lon, lat in cities]

    for name, (cx, cy, area) in centroids.items():
        if area < 0.02 * W * H:                     # Estuaire : trop étroit, étiqueté par ville
            continue
        if any((cx - x) ** 2 + (cy - y) ** 2 < 26 ** 2 for x, y in city_pts):
            continue
        dx, dy = NUDGE.get(name, (0.0, 0.0))
        d.text((cx + dx * W, cy + dy * H), name.upper(), font=f_lab,
               fill=(0, 53, 94, 225), anchor="mm")

    for (label, _, _), (X, Y) in zip(cities, city_pts):
        d.ellipse([X - 4, Y - 4, X + 4, Y + 4], fill=(0, 53, 94, 255))
        d.text((X + 7, Y), label, font=f_cit, fill=(0, 53, 94, 255), anchor="lm")
    return img


def make_background(page_no: int) -> Path:
    CACHE.mkdir(parents=True, exist_ok=True)
    out = CACHE / f"fond-{page_no:02d}.jpg"
    doc = pymupdf.open(PDF)
    pg = doc[page_no - 1]
    for b in pg.get_text("dict")["blocks"]:
        if b["type"] == 0:
            pg.add_redact_annot(b["bbox"], fill=None)
    pg.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE,
                        graphics=pymupdf.PDF_REDACT_LINE_ART_NONE)
    pix = pg.get_pixmap(dpi=DPI)
    img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    doc.close()

    if page_no == 2:                                # carte : le Gabon, pas l'Afrique
        s = img.width / PDF_W
        x0, y0, x1, y1 = [int(v * s) for v in (48, 250, 618, 765)]
        carte = carte_gabon_claire((x1 - x0, y1 - y0)).convert("RGB")
        masque = Image.new("L", carte.size, 0)
        ImageDraw.Draw(masque).rounded_rectangle([0, 0, carte.width - 1, carte.height - 1],
                                                 radius=int(26 * s), fill=255)
        img.paste(carte, (x0, y0), masque)

    img.save(out, quality=84, subsampling=2)
    return out


# --- Écriture des textes ---------------------------------------------------------
def add_box(slide, x0: float, y0: float, w: float, h: float):
    box = slide.shapes.add_textbox(Inches(x0 * K), Inches((y0 - 1.5) * K),
                                   Inches(w * K + 0.10), Inches(max(h, 8) * K))
    tf = box.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    return tf


def put(tf, text: str, size: float, color: str, font: str, bold: bool, align,
        pitch: float | None = None, first: bool = True, wrap_on: bool = False):
    tf.word_wrap = wrap_on
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    if pitch:
        p.line_spacing = pitch
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.name = font
    r.font.bold = bold
    r.font.color.rgb = rgb(color)


def obstacles(page: pymupdf.Page) -> list[tuple[float, float, float]]:
    """Objets du modèle (images, aplats, formes) qui bornent une zone de texte."""
    out = []
    for info in page.get_image_info():
        x0, y0, x1, y1 = info["bbox"]
        if (x1 - x0) * (y1 - y0) > 4000:
            out.append((y0, x0, x1))
    for d in page.get_drawings():
        r = d["rect"]
        if r.width * r.height > 9000 and r.height > 12:
            out.append((r.y0, r.x0, r.x1))
    return out


def available_height(blocks, page_obstacles, idx, x0, x1, y0) -> float:
    best = PDF_H - 26 - y0
    for j, other in enumerate(blocks):
        if j == idx:
            continue
        ox0, oy0, ox1, _ = block_box(other)
        if oy0 <= y0 + 2 or ox1 < x0 + 4 or ox0 > x1 - 4:
            continue
        best = min(best, oy0 - y0 - 5)
    for oy0, ox0, ox1 in page_obstacles:
        if oy0 <= y0 + 2 or ox1 < x0 + 4 or ox0 > x1 - 4:
            continue
        best = min(best, oy0 - y0 - 5)
    return max(best, 20)


def build(preview: Path | None = None) -> None:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    cp = prs.core_properties
    cp.title = "Soutenance DBA — Gouvernance augmentée des EPES gabonais"
    cp.author = "Serge Patrick MINANG"
    cp.subject = "Mise en page du modèle Canva de l'auteur, remplie avec le contenu de la thèse"
    cp.comments = ("18 diapositives. Aucun élément du modèle n'a été déplacé : les textes sont "
                   "reposés aux coordonnées d'origine.")
    blank = prs.slide_layouts[6]
    report: list[str] = []

    for pos, page_no in enumerate(PAGES, 1):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(make_background(page_no)), 0, 0,
                                 Inches(SLIDE_W), Inches(SLIDE_H))

        doc = pymupdf.open(PDF)
        blocks = read_blocks(doc[page_no - 1])
        page_obstacles = obstacles(doc[page_no - 1])
        doc.close()

        for idx, lines in enumerate(blocks):
            numeral, body = split_num(lines)
            body_txt = block_text(body) if body else ""

            if numeral is not None:
                txt = lookup(NUM_FIX, page_no, body_txt) or numeral["t"].strip()
                font, _ = pptx_face(numeral["font"])
                colour = "#ffffff" if numeral["color"].lower() == "#ffffff" else "#00355e"
                tf = add_box(slide, numeral["x0"], numeral["y0"],
                             numeral["x1"] - numeral["x0"], numeral["y1"] - numeral["y0"])
                put(tf, txt, numeral["size"] * FT, colour, font, True, PP_ALIGN.CENTER)

            if not body:
                continue

            x0, y0, x1, y1 = block_box(body)
            first_run = body[0][0]
            pdf_font, colour = first_run["font"], first_run["color"]
            font, bold = pptx_face(pdf_font)
            base = first_run["size"]
            ys = sorted({round(l[0]["y0"], 1) for l in body})
            pitch = ((ys[1] - ys[0]) / (base * 1.2)) if len(ys) > 1 else None
            centres = [(line_box(l)[0] + line_box(l)[2]) / 2 for l in body]
            centred = (max(centres) - min(centres)) < 4 and (x1 - x0) > 0.5 * PDF_W
            align = PP_ALIGN.CENTER if centred else PP_ALIGN.LEFT

            override = lookup(OVR, page_no, body_txt)
            if override == "":
                continue
            if override is None:
                # Texte du modèle : on garde ses retours à la ligne, la taille est ajustée
                # juste ce qu'il faut pour que la police de substitution ne déborde pas.
                factor = 1.0
                for l in body:
                    txt = line_text(l)
                    width = line_box(l)[2] - line_box(l)[0]
                    if txt.strip():
                        ratio = width / max(measure(txt, base, pdf_font), 1e-6)
                        factor = min(factor, ratio * slack(pdf_font))
                factor = max(0.72, min(1.0, factor))
                size = base * FT * factor
                if factor < 0.965:
                    report.append(f"  page {page_no:2d} « {body_txt[:44]} » "
                                  f"— {base * FT:.1f} -> {size:.1f} pt")
                tf = add_box(slide, x0, y0, x1 - x0, y1 - y0)
                for i, l in enumerate(body):
                    put(tf, line_text(l), size, colour, font, bold, align, pitch,
                        first=(i == 0))
            else:
                # Texte de l'auteur : reflow dans la colonne d'origine, taille qui tient.
                width = x1 - x0
                height = available_height(blocks, page_obstacles, idx, x0, x1, y0)
                size = base * FT
                while size > base * FT * 0.62:
                    total = len(wrap(override.replace("|", " "), width, size / FT,
                                     pdf_font)) * size / FT * 1.34
                    if total <= height:
                        break
                    size *= 0.965
                tf = add_box(slide, x0, y0, width, height)
                for i, part in enumerate(override.split("|")):
                    put(tf, part, size, colour, font, bold, align, pitch,
                        first=(i == 0), wrap_on=True)

        if pos in NOTES:
            slide.notes_slide.notes_text_frame.text = NOTES[pos]

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"écrit : {OUT.relative_to(ROOT)} — {len(PAGES)} diapositives, "
          f"{OUT.stat().st_size / 1e6:.1f} Mo")
    if report:
        print("tailles ajustées pour rester dans la largeur du modèle :")
        print("\n".join(report))
    else:
        print("aucune réduction de taille nécessaire")

    if preview:
        preview.mkdir(parents=True, exist_ok=True)
        for i, page_no in enumerate(PAGES, 1):
            Image.open(CACHE / f"fond-{page_no:02d}.jpg").save(preview / f"fond-{i:02d}.png")
        print(f"fonds exportés dans {preview}")


if __name__ == "__main__":
    build(Path(sys.argv[sys.argv.index("--apercu") + 1]) if "--apercu" in sys.argv else None)
