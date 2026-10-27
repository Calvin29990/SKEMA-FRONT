"""Présentation de groupe 'Why would we invest in IKEA?' - 5 diapositives, calibrée pour 3 minutes.

Style inspiré Canva (blocs colorés, gros chiffres) ; logo IKEA simplifié dessiné en formes natives.
Chaque diapositive porte dans les notes : le script oral, le minutage et les idées détaillées,
pour ne rien perdre de la version précédente.

Usage : python3 build_pptx.py
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "Groupe-IKEA-Why-invest-3min.pptx")

BLUE = RGBColor(0x00, 0x58, 0xA3)
YEL = RGBColor(0xFF, 0xDA, 0x1A)
INK = RGBColor(0x1A, 0x1A, 0x1A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0x5F, 0x63, 0x68)
LIGHT = RGBColor(0xF3, 0xF5, 0xF8)
GREEN = RGBColor(0x1E, 0x7B, 0x4A)
RED = RGBColor(0xB3, 0x26, 0x1E)
FONT = "Arial"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def box(slide, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius
    return shp


def text(slide, x, y, w, h, runs, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, shape=None):
    """runs : str, ou liste de paragraphes (str). Un paragraphe peut être un tuple (texte, {opts})."""
    tb = shape if shape is not None else slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    paras = runs if isinstance(runs, list) else [runs]
    for i, p_item in enumerate(paras):
        opts = {}
        if isinstance(p_item, tuple):
            p_item, opts = p_item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = opts.get("align", align)
        p.space_after = Pt(opts.get("space", 6))
        r = p.add_run()
        r.text = p_item
        f = r.font
        f.name = FONT
        f.size = Pt(opts.get("size", size))
        f.bold = opts.get("bold", bold)
        f.color.rgb = opts.get("color", color)
    return tb


def notes(slide, txt):
    slide.notes_slide.notes_text_frame.text = txt


def logo(slide, x, y, scale=1.0):
    """Logo IKEA simplifié : rectangle bleu + ovale jaune + mot IKEA."""
    w, h = 2.4 * scale, 1.5 * scale
    box(slide, x, y, w, h, BLUE, shape=MSO_SHAPE.RECTANGLE)
    ov = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.12 * scale), Inches(y + 0.36 * scale),
                                Inches(w - 0.24 * scale), Inches(h - 0.72 * scale))
    ov.fill.solid()
    ov.fill.fore_color.rgb = YEL
    ov.line.fill.background()
    text(slide, x, y + 0.38 * scale, w, h - 0.76 * scale, "IKEA", size=int(30 * scale), color=BLUE,
         bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def header(slide, title, n):
    box(slide, 0, 0, 13.333, 0.12, BLUE, shape=MSO_SHAPE.RECTANGLE)
    text(slide, 0.6, 0.45, 11.8, 0.9, title, size=32, color=BLUE, bold=True)
    box(slide, 0.65, 1.35, 1.1, 0.08, YEL, shape=MSO_SHAPE.RECTANGLE)
    text(slide, 12.0, 6.95, 1.0, 0.35, f"{n} / 5", size=10, color=GREY, align=PP_ALIGN.RIGHT)


# ---------- 1. Titre (0:00 - 0:15) ----------
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, 13.333, 7.5, BLUE, shape=MSO_SHAPE.RECTANGLE)
box(s, 0, 0, 0.18, 7.5, YEL, shape=MSO_SHAPE.RECTANGLE)
box(s, 0.9, 0.8, 2.9, 2.0, WHITE, radius=0.1)
logo(s, 1.15, 1.05, scale=1.0)
text(s, 0.9, 3.4, 11.5, 1.6, "Why would we invest in IKEA?", size=54, color=YEL, bold=True)
text(s, 0.9, 5.0, 11.5, 0.6, "Une lecture ESG d'Ingka Group et d'IKEA France", size=22, color=WHITE)
text(s, 0.9, 6.1, 11.5, 0.8,
     ["MSc CFM · SKEMA · Sustainable Finance",
      "Calvin Minang · Jonathan Bouaffon · Liza Amouyal"], size=14, color=YEL)
notes(s, "[0:00 - 0:15]  Titre.\n"
         "Script : « Bonjour, nous sommes… Notre question : pourquoi investirait-on dans IKEA ? »\n\n"
         "Idée à garder : on ne parle pas d'action cotée. Ingka Group n'est pas coté : l'exposition passerait "
         "par la dette ou une voie privée. D'où une question de risque, pas de rendement.")

# ---------- 2. Réponse (0:15 - 0:45) ----------
s = prs.slides.add_slide(BLANK)
header(s, "Notre réponse : un oui conditionnel", 2)
box(s, 0.65, 1.75, 5.1, 4.9, YEL, radius=0.08)
text(s, 0.95, 2.05, 4.6, 1.3, "Oui, sous conditions.", size=36, color=BLUE, bold=True)
text(s, 0.95, 3.45, 4.6, 2.4,
     [("Un cas de réduction du risque,", {"size": 22, "bold": True}),
      ("pas une prime ESG.", {"size": 22, "bold": True}),
      ("Aucune preuve que l'ESG augmente le rendement ici.", {"size": 16})], color=INK)
kpis = [("−39 %", "émissions opérationnelles\nen France depuis 2016"),
        ("184 %", "des besoins électriques de la France\ncouverts par éolien et solaire"),
        ("−30 %", "empreinte du groupe, tous scopes,\ndepuis 2016 (objectif −50 % en 2030)")]
for i, (num, lbl) in enumerate(kpis):
    y = 1.75 + i * 1.7
    box(s, 6.2, y, 6.5, 1.5, LIGHT, radius=0.1)
    box(s, 6.2, y, 0.12, 1.5, BLUE, shape=MSO_SHAPE.RECTANGLE)
    text(s, 6.6, y + 0.2, 2.6, 1.1, num, size=40, color=BLUE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 9.2, y + 0.2, 3.4, 1.1, lbl, size=15, color=INK, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "[0:15 - 0:45]  Réponse en premier.\n"
         "Script : « Notre réponse est un oui conditionnel : c'est un cas de réduction du risque, pas une prime ESG. »\n\n"
         "Chiffres : −39 % d'émissions opérationnelles en France depuis 2016 ; éolien et solaire couvrent 184 % des "
         "besoins électriques de la France ; −30 % pour le groupe, tous scopes, depuis 2016 malgré la croissance ; "
         "objectif −50 % d'ici 2030 (SBTi, scopes 1 à 3), net zéro en 2050.\n"
         "Attention : périmètres différents (France ≠ groupe) ; données auto-déclarées, non testées par nous.")

# ---------- 3. Environnement et social / gouvernance (0:45 - 1:35) ----------
s = prs.slides.add_slide(BLANK)
header(s, "Environnement, social, gouvernance : ce qui tient, ce qui manque", 3)
# barres empreinte
text(s, 0.65, 1.7, 5.6, 0.5, "Où est l'empreinte du groupe ?", size=16, color=BLUE, bold=True)
bars = [("Matières premières", 52, BLUE), ("Usage à domicile", 17, BLUE), ("Autres postes", 31, RGBColor(0xB8, 0xC4, 0xD6))]
for i, (lbl, val, col) in enumerate(bars):
    y = 2.35 + i * 0.85
    text(s, 0.65, y, 4.0, 0.35, lbl, size=13, color=INK)
    box(s, 0.65, y + 0.36, 4.3 * val / 52, 0.36, col, radius=0.06)
    text(s, 0.65 + 4.3 * val / 52 + 0.1, y + 0.33, 1.2, 0.42, f"{val} %", size=14, color=INK, bold=True)
text(s, 0.65, 5.0, 5.6, 1.2,
     "Les matières premières et l'usage pèsent plus que les opérations : la chaîne d'approvisionnement compte davantage.",
     size=14, color=INK)
# colonnes ce qui avance / manque
cols = [("Ce qui avance", GREEN, ["Opérations France : −39 %", "Éolien et solaire : 184 % des besoins",
                                   "Chauffage sans fossile d'ici 2030", "Code IWAY et audits fournisseurs",
                                   "Accessibilité : 200–400 emplois par magasin", "Égalité femmes-hommes : 100/100"]),
        ("Ce qui manque", RED, ["Taxonomie UE : alignement non publié", "Premiers résultats TNFD en 2026",
                                "Objectif SBTi : transition discutable", "Chaîne hors France à vérifier",
                                "Groupe non coté : détention à confirmer"])]
for i, (ttl, col, items) in enumerate(cols):
    x = 6.7 + i * 3.2
    box(s, x, 1.7, 3.0, 5.0, LIGHT, radius=0.08)
    box(s, x, 1.7, 3.0, 0.14, col, shape=MSO_SHAPE.RECTANGLE)
    text(s, x + 0.15, 1.95, 2.7, 0.5, ttl, size=17, color=col, bold=True)
    text(s, x + 0.15, 2.5, 2.75, 4.0, [(f"•  {it}", {"space": 10}) for it in items], size=13, color=INK)
notes(s, "[0:45 - 1:35]  Environnement, social, gouvernance.\n"
         "Script : « Côté environnement, les opérations baissent, mais l'essentiel de l'empreinte est dans la matière "
         "(52 %) et l'usage (17 %). Côté social, l'accessibilité et le code IWAY réduisent le risque ; côté gouvernance, "
         "devoir de vigilance, CSRD et reporting public. Ce qui manque : taxonomie non publiée, premiers résultats TNFD "
         "attendus en 2026, objectif SBTi discutable. »\n\n"
         "Détails à garder : 200–400 emplois locaux par magasin ; environ 200 M€ investis pour baisser les prix de "
         "~2 000 produits ; 7 % d'emploi de personnes handicapées (légal 6 %) ; environ 250 réfugiés visés en 3 ans ; "
         "Fondation IKEA plus de 2 Md€ depuis 2009 ; 1,2 Md€ investis en France 2023-26 ; gestion des fournisseurs "
         "(IWAY, audits) ; SFDR classe des fonds, pas des entreprises.")

# ---------- 4. Mécanismes vers la valeur (1:35 - 2:15) ----------
s = prs.slides.add_slide(BLANK)
header(s, "De l'ESG à la valeur : trois canaux, un effet principal", 4)
cards = [("1", "Flux de trésorerie", "Revenus et marges",
          ["Accessibilité : plus de clients", "Seconde vie et rachat : marge", "Énergie renouvelable : coût"], BLUE, WHITE),
         ("2", "Risque", "Taux d'actualisation  ★ effet principal",
          ["IWAY, audits, reporting", "Moins de risque réglementaire", "et de réputation"], YEL, INK),
         ("3", "Investissements", "Résilience future",
          ["Pompes à chaleur, géothermie", "Dépenses aujourd'hui", "Protection demain"], BLUE, WHITE)]
for i, (num, ttl, sub, items, col, tcol) in enumerate(cards):
    x = 0.65 + i * 4.1
    box(s, x, 1.75, 3.85, 3.9, LIGHT, radius=0.08)
    box(s, x, 1.75, 3.85, 0.85, col, radius=0.08)
    text(s, x + 0.2, 1.85, 3.5, 0.65, f"{num}  {ttl}", size=20, color=tcol, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.2, 2.8, 3.5, 0.5, sub, size=14, color=BLUE, bold=True)
    text(s, x + 0.2, 3.4, 3.5, 2.2, [(f"•  {it}", {"space": 10}) for it in items], size=15, color=INK)
box(s, 0.65, 5.95, 12.0, 0.85, YEL, radius=0.12)
text(s, 0.9, 6.0, 11.5, 0.75,
     "Matières premières 52 % et usage 17 % de l'empreinte : la chaîne compte plus que les −39 % en France.",
     size=16, color=INK, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "[1:35 - 2:15]  Mécanismes.\n"
         "Script : « L'ESG agit par trois canaux : les flux de trésorerie, le risque et les investissements. "
         "Le principal est le risque : IWAY, audits et reporting réduisent les risques réglementaires et de réputation. "
         "Cela joue sur le taux d'actualisation, pas sur le chiffre d'affaires. »\n\n"
         "Détails à garder : canal flux = accessibilité qui élargit la clientèle, seconde vie et rachat qui ajoutent "
         "de la marge, renouvelables qui couvrent une part du coût énergétique. Canal investissements = pompes à chaleur, "
         "géothermie, renouvelables : dépenses aujourd'hui, résilience demain ; effet difficile à chiffrer sans taxonomie publiée.")

# ---------- 5. Décision et limites (2:15 - 2:50) ----------
s = prs.slides.add_slide(BLANK)
header(s, "Règle de décision : oui tant que quatre conditions tiennent", 5)
conds = ["Les allégations vertes restent vérifiables",
         "La taxonomie UE est publiée au prochain rapport",
         "Les premiers résultats TNFD soutiennent la feuille de route nature",
         "Le groupe reste sur la trajectoire −50 % d'ici 2030"]
for i, txt in enumerate(conds):
    y = 1.75 + i * 0.95
    circ = box(s, 0.7, y, 0.7, 0.7, BLUE, shape=MSO_SHAPE.OVAL)
    text(s, 0.7, y, 0.7, 0.7, "abcd"[i], size=18, color=YEL, bold=True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, shape=circ)
    text(s, 1.6, y + 0.05, 5.6, 0.65, txt, size=16, color=INK, anchor=MSO_ANCHOR.MIDDLE)
box(s, 7.7, 1.75, 5.0, 3.7, LIGHT, radius=0.08)
text(s, 7.95, 1.9, 4.6, 0.5, "Limites à garder en tête", size=17, color=BLUE, bold=True)
text(s, 7.95, 2.5, 4.6, 3.0,
     [(f"•  {it}", {"space": 9}) for it in [
         "Aucune valorisation, spread ni rendement",
         "Une seule année de reporting, données auto-déclarées",
         "Chiffres France ≠ périmètre groupe",
         "Voie d'investissement (dette ou privé) à confirmer"]], size=14, color=INK)
box(s, 0.65, 5.9, 12.0, 0.9, BLUE, radius=0.12)
text(s, 0.9, 5.95, 11.5, 0.8,
     [("Verdict : oui conditionnel, un cas de réduction du risque, révisé si une condition tombe.",
       {"color": WHITE, "size": 18, "bold": True})], anchor=MSO_ANCHOR.MIDDLE)
notes(s, "[2:15 - 2:50]  Décision et limites.\n"
         "Script : « Notre règle de décision : oui tant que quatre conditions tiennent. Nous réviserons si l'une tombe. "
         "Nos limites : pas de valorisation ni de rendement, une seule année de reporting, des données auto-déclarées. »\n\n"
         "Détails à garder : (a) allégations vertes vérifiables ; (b) alignement taxonomie publié dans le prochain rapport ; "
         "(c) premiers résultats TNFD cohérents avec la feuille de route nature ; (d) trajectoire −50 % 2030 "
         "(scopes 1 à 3). Limites complémentaires : groupe non coté, détention à vérifier (Stichting INGKA Foundation) ; "
         "chiffres France ≠ groupe.\n\n"
         "Fin : « Merci, nous sommes prêts pour vos questions. » (≈ 2 min 50 au total, à répéter une fois au chrono.)")

prs.core_properties.title = "Why would we invest in IKEA? - présentation de groupe (3 min)"
prs.core_properties.author = "Calvin Minang, Jonathan Bouaffon, Liza Amouyal"
prs.save(OUT)
print("écrit :", OUT)
