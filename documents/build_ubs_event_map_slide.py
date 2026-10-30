from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
MAP = ROOT / "assets" / "ubs-event-map-approved.png"
PPTX_OUT = ROOT / "ubs-event-map-slide.pptx"
PNG_OUT = ROOT / "ubs-event-map-slide-preview.png"

# Palette
BG = "F4F7FA"
NAVY = "173044"
MUTED = "647684"
TEAL = "139B96"
BORDER = "E1E8EE"
WHITE = "FFFFFF"
GOLD = "A77A00"
RED = "C8272C"
DARK_RED = "B82631"

SLIDE_W, SLIDE_H = 13.333333, 7.5


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def add_text(slide, text, x, y, w, h, size, color, bold=False,
             font="Arial", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE,
             margin=0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def rounded_card(slide, x, y, w, h, fill=WHITE, line=BORDER, radius=MSO_SHAPE.ROUNDED_RECTANGLE):
    shape = slide.shapes.add_shape(radius, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(0.8)
    return shape


def add_case(slide, x, y, label, year, amount, color, org=None):
    w, h = 4.12, 1.02
    rounded_card(slide, x, y, w, h)
    accent = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                    Inches(x + 0.16), Inches(y + 0.20),
                                    Inches(0.055), Inches(0.61))
    accent.fill.solid()
    accent.fill.fore_color.rgb = rgb(color)
    accent.line.fill.background()
    add_text(slide, label, x + 0.34, y + 0.10, 2.95, 0.27, 12.5, NAVY, True)
    add_text(slide, year, x + 3.29, y + 0.10, 0.60, 0.27, 9.5, color, True,
             align=PP_ALIGN.RIGHT)
    add_text(slide, amount, x + 0.34, y + 0.44, 2.30, 0.42, 19, color, True)
    if org:
        add_text(slide, org, x + 2.18, y + 0.52, 1.70, 0.25, 9, MUTED, False,
                 align=PP_ALIGN.RIGHT)


def build_pptx():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(BG)

    # Header
    add_text(slide, "RISK CASE STUDY  /  EUROPE", 0.56, 0.23, 5.0, 0.22,
             9, TEAL, True)
    add_text(slide, "MAJOR EUROPEAN TRADING LOSS EVENTS", 0.56, 0.48,
             8.0, 0.43, 24, NAVY, True)
    add_text(slide,
             "Selected benchmark trading-loss events relevant to the UBS 2011 case",
             0.57, 0.98, 8.4, 0.24, 10.5, MUTED)

    # Map panel: retain the approved crop's native aspect ratio.
    rounded_card(slide, 0.45, 1.43, 7.55, 5.66)
    slide.shapes.add_picture(str(MAP), Inches(0.55), Inches(1.50),
                             width=Inches(7.35), height=Inches(5.52))

    # Event sidebar
    rounded_card(slide, 8.20, 1.43, 4.68, 5.66)
    add_text(slide, "BENCHMARK CASES", 8.48, 1.73, 3.9, 0.22,
             9, TEAL, True)
    add_text(slide, "Three historic losses", 8.48, 1.99, 4.1, 0.36,
             16, NAVY, True)

    add_case(slide, 8.48, 2.50, "BARINGS", "1995", "USD -1.4bn", GOLD)
    add_case(slide, 8.48, 3.66, "KERVIEL", "2008", "EUR -4.9bn", DARK_RED,
             "Société Générale")
    add_case(slide, 8.48, 4.82, "ADOBOLI", "2011", "USD -2.3bn", RED, "UBS")

    # Geography note distinguishes the trading desk from the bank headquarters.
    note = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                  Inches(8.48), Inches(6.00), Inches(4.12), Inches(0.82))
    note.fill.solid()
    note.fill.fore_color.rgb = rgb(NAVY)
    note.line.fill.background()
    add_text(slide, "LOCATION MATTERS", 8.70, 6.11, 3.65, 0.19,
             8.5, "75D8CF", True)
    add_text(slide, "Adoboli traded from London; UBS headquarters are in Zurich.",
             8.70, 6.34, 3.65, 0.39, 10.5, WHITE, False)

    prs.core_properties.title = "Major European Trading Loss Events"
    prs.core_properties.subject = "One-slide map of selected European trading-loss events"
    prs.core_properties.author = "SKEMA"
    prs.save(PPTX_OUT)


def build_preview():
    # A faithful high-resolution visual preview of the PowerPoint layout.
    scale = 120
    canvas = Image.new("RGB", (int(SLIDE_W * scale), int(SLIDE_H * scale)), f"#{BG}")
    draw = ImageDraw.Draw(canvas)
    regular_path = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    bold_path = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"

    def font(size, bold=False):
        path = bold_path if bold else regular_path
        return ImageFont.truetype(path, int(size * scale / 72))

    def px(v):
        return int(v * scale)

    def text(s, x, y, size, color, bold=False):
        draw.text((px(x), px(y)), s, font=font(size, bold), fill=f"#{color}")

    def card(x, y, w, h, fill=WHITE, outline=BORDER, radius=12):
        draw.rounded_rectangle((px(x), px(y), px(x+w), px(y+h)),
                               radius=px(radius/72), fill=f"#{fill}",
                               outline=f"#{outline}", width=2)

    text("RISK CASE STUDY  /  EUROPE", 0.56, 0.25, 9, TEAL, True)
    text("MAJOR EUROPEAN TRADING LOSS EVENTS", 0.56, 0.52, 24, NAVY, True)
    text("Selected benchmark trading-loss events relevant to the UBS 2011 case",
         0.57, 1.01, 10.5, MUTED)

    card(0.45, 1.43, 7.55, 5.66)
    img = Image.open(MAP).convert("RGB")
    img = img.resize((px(7.35), px(5.52)), Image.Resampling.LANCZOS)
    canvas.paste(img, (px(0.55), px(1.50)))

    card(8.20, 1.43, 4.68, 5.66)
    text("BENCHMARK CASES", 8.48, 1.77, 9, TEAL, True)
    text("Three historic losses", 8.48, 2.04, 16, NAVY, True)

    cases = [
        (2.50, "BARINGS", "1995", "USD -1.4bn", GOLD, None),
        (3.66, "KERVIEL", "2008", "EUR -4.9bn", DARK_RED, "Société Générale"),
        (4.82, "ADOBOLI", "2011", "USD -2.3bn", RED, "UBS"),
    ]
    for y, label, year, amount, color, org in cases:
        card(8.48, y, 4.12, 1.02)
        draw.rounded_rectangle((px(8.64), px(y+0.20), px(8.695), px(y+0.81)),
                               radius=3, fill=f"#{color}")
        text(label, 8.82, y+0.14, 12.5, NAVY, True)
        text(year, 11.98, y+0.17, 9.5, color, True)
        text(amount, 8.82, y+0.49, 19, color, True)
        if org:
            text(org, 10.77, y+0.58, 8.3, MUTED)

    draw.rounded_rectangle((px(8.48), px(6.00), px(12.60), px(6.82)),
                           radius=px(0.08), fill=f"#{NAVY}")
    text("LOCATION MATTERS", 8.70, 6.12, 8.5, "75D8CF", True)
    text("Adoboli traded from London; UBS", 8.70, 6.37, 9.6, WHITE)
    text("headquarters are in Zurich.", 8.70, 6.55, 9.6, WHITE)
    canvas.save(PNG_OUT, optimize=True)


if __name__ == "__main__":
    build_pptx()
    build_preview()
    print(PPTX_OUT)
    print(PNG_OUT)
