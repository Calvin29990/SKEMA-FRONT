from pathlib import Path
from PIL import Image
import xlsxwriter
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
PPTX_OUT = ROOT / "ubs-treasury-risk-presentation-8-slides.pptx"
XLSX_OUT = ROOT / "ubs-illustrative-liquidity-model.xlsx"
MAP_SLIDE = ASSETS / "approved-map-slide-report-order.png"
NORMAL_HIDDEN_SLIDE = ASSETS / "approved-three-panel-slide.png"
DALTONS = ASSETS / "daltons-illustration.png"
UBS_LOGO = ASSETS / "ubs-logo-slide.png"
SKEMA_LOGO = ASSETS / "skema-logo.png"

SW, SH = 13.333333, 7.5
BG = "F7F7F3"
NAVY = "18232B"
BLACK = "101315"
DARK = "252B2F"
WHITE = "FFFFFF"
MUTED = "68747B"
SOFT = "A9B0B2"
RED = "C82036"
TEAL = "008F72"
GOLD = "C49500"
BLUE = "287FA8"
GREEN = "4F8A64"
GRID = "C9CECB"
LIGHT = "E8ECE8"
PALE_RED = "F6E2E3"
PALE_AMBER = "F8F0D9"
PALE_GREEN = "E6F0E9"
FONT = "Arial"

DATES = ["14 Sep", "15", "16", "17", "18", "19", "20", "30 Sep", "Oct", "Nov", "14 Dec"]
NO_INCIDENT = [1000, 990, 985, 990, 985, 987, 984, 1004, 1104, 1054, 1024]
INCIDENT_BASE = [1000, 915, 805, 775, 765, 772, 774, 769, 909, 849, 824]
COMBINED_STRESS = [1000, 915, 805, 775, 765, 772, 774, 769, 909, 549, 524]
BUFFER = 700


def C(value):
    return RGBColor.from_string(value)


def rect(slide, x, y, w, h, fill, line_color=None, line_width=0.8,
         shape=MSO_SHAPE.RECTANGLE):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = C(fill)
    if line_color:
        sh.line.color.rgb = C(line_color)
        sh.line.width = Pt(line_width)
    else:
        sh.line.fill.background()
    return sh


def line(slide, x1, y1, x2, y2, color=GRID, width=1.0):
    sh = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                    Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    sh.line.color.rgb = C(color)
    sh.line.width = Pt(width)
    return sh


def txt(slide, value, x, y, w, h, size=12, color=NAVY, bold=False,
        align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE, margin=0.02):
    sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    r = p.add_run()
    r.text = value
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = C(color)
    return sh


def add_img_fit(slide, path, x, y, w, h):
    """Fit the complete image without crop or distortion."""
    with Image.open(path) as im:
        iw, ih = im.size
    aspect, box_aspect = iw/ih, w/h
    if aspect > box_aspect:
        dw, dh = w, w/aspect
        dx, dy = x, y+(h-dh)/2
    else:
        dh, dw = h, h*aspect
        dx, dy = x+(w-dw)/2, y
    return slide.shapes.add_picture(str(path), Inches(dx), Inches(dy),
                                    width=Inches(dw), height=Inches(dh))


def header(slide, title, kicker, page, section):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = C(BG)
    slide.shapes.add_picture(str(UBS_LOGO), Inches(.42), Inches(.13),
                             width=Inches(1.38), height=Inches(.49))
    txt(slide, title, 2.02, .15, 9.65, .31, 17, NAVY, True)
    txt(slide, kicker, 2.02, .47, 9.65, .20, 9.3, RED, True)
    txt(slide, section, 2.02, .71, 9.65, .16, 7.4, MUTED)
    slide.shapes.add_picture(str(SKEMA_LOGO), Inches(12.35), Inches(.12),
                             width=Inches(.52), height=Inches(.55))
    line(slide, .42, .97, 12.90, .97, GRID, .9)
    line(slide, .42, 7.17, 12.90, 7.17, GRID, .7)
    txt(slide, f"UBS  ×  SKEMA   /   {page:02d}", 11.40, 7.22, 1.45, .12,
        6.4, MUTED, True, align=PP_ALIGN.RIGHT)


def add_user_slide(prs, image):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = C(BG)
    add_img_fit(s, image, 0, 0, SW, SH)
    return s


def build_model_workbook():
    wb = xlsxwriter.Workbook(str(XLSX_OUT))
    wb.set_properties({
        "title": "UBS 2011 illustrative Treasury liquidity model",
        "subject": "Cash bridge and stress assumptions; not historical UBS Treasury data",
        "author": "Corporate Treasury Management Team",
    })
    navy = "18232B"
    red = "C82036"
    teal = "008F72"
    gold = "C49500"
    muted = "68747B"
    grid = "D9DEDB"
    pale = "F1F4F1"
    fmt_title = wb.add_format({"bold": True, "font_size": 16, "font_color": navy})
    fmt_sub = wb.add_format({"font_color": muted, "font_size": 10, "text_wrap": True})
    fmt_head = wb.add_format({"bold": True, "font_color": "FFFFFF", "bg_color": navy,
                              "border": 1, "border_color": "FFFFFF", "text_wrap": True,
                              "valign": "vcenter"})
    fmt_input = wb.add_format({"bg_color": "FFF2CC", "border": 1, "border_color": grid,
                               "num_format": '#,##0;[Red](#,##0);–'})
    fmt_num = wb.add_format({"border": 1, "border_color": grid,
                             "num_format": '#,##0;[Red](#,##0);–'})
    fmt_num_bold = wb.add_format({"bold": True, "border": 1, "border_color": grid,
                                  "num_format": '#,##0;[Red](#,##0);–'})
    fmt_label = wb.add_format({"border": 1, "border_color": grid})
    fmt_note = wb.add_format({"font_color": muted, "italic": True, "text_wrap": True})
    fmt_section = wb.add_format({"bold": True, "font_color": "FFFFFF", "bg_color": teal})
    fmt_red = wb.add_format({"bold": True, "font_color": red, "border": 1, "border_color": grid,
                             "num_format": '#,##0;[Red](#,##0);–'})

    # README / evidence classification
    ws = wb.add_worksheet("Read me")
    ws.set_column("A:A", 27); ws.set_column("B:B", 94)
    ws.write("A1", "UBS 2011 · Treasury liquidity model", fmt_title)
    ws.merge_range("A2:B2", "Illustrative teaching model built in Python / Excel. It is not UBS historical Treasury data.", fmt_sub)
    notes = [
        ("Central question", "How should UBS Treasury preserve liquidity and meet payment obligations over the three months after concealed trading losses are discovered?"),
        ("Model horizon", "14 Sep–14 Dec 2011; daily in the first week, then period-end snapshots."),
        ("Evidence: public", "Incident dates, unauthorised trading, reported USD 2.3bn loss and control findings."),
        ("Inputs: assumed", "Opening cash USD 1,000m; minimum buffer USD 700m; receipts, payments, settlements, collateral, costs and stress shocks."),
        ("Outputs: calculated", "Scenario closing cash, first buffer breach, and peak funding need. Reproduce from the Cash Bridge formulas."),
        ("Cash balance", "Accessible opening cash + receipts − payments − settlement outflow + signed collateral movement − costs + confirmed funding."),
        ("Funding gap", "max(0, minimum buffer − cash before funding). A trading loss is not automatically a same-day cash payment."),
        ("Use limitation", "Replace all yellow inputs with dated, entity-level Treasury ledger records before treating any funding gap as a UBS result."),
        ("Sources", "FINMA 2012; FSA 2012; UBS Q3 2011 disclosures. Model schema follows the illustrative cash bridge in the presentation."),
    ]
    for i,(a,b) in enumerate(notes,4):
        ws.write(i-1,0,a,fmt_head if i==4 else fmt_label)
        ws.write(i-1,1,b,fmt_sub if i==4 else fmt_label)
        ws.set_row(i-1,36 if len(b)>85 else 26)
    ws.freeze_panes(3,0)

    # Cash bridge with visible formulas and model inputs.
    cash = wb.add_worksheet("Cash Bridge")
    cash.freeze_panes(3,1)
    cash.set_column("A:A", 16)
    cash.set_column("B:K", 18)
    cash.merge_range("A1:K1", "Cash bridge · illustrative USD millions", fmt_title)
    cash.merge_range("A2:K2", "Yellow cells are assumed inputs. Closing cash is formula-driven; positive collateral movement is cash returned to UBS.", fmt_sub)
    headers=["Period","Base receipts","Base payments","Incident settlement\noutflow","Collateral\nmovement (+/−)","Incident costs","No-incident\nclose","Incident-base\nclose","Stress\nadjustment","Stress\nclose","Minimum\nbuffer"]
    for c,h in enumerate(headers): cash.write(2,c,h,fmt_head)
    cash.set_row(2,36)
    rows=[
        ("14 Sep · open",0,0,0,0,0),
        ("15 Sep",120,130,55,-20,0),
        ("16 Sep",90,95,80,-25,0),
        ("17 Sep",110,105,45,10,0),
        ("18 Sep",100,105,20,15,0),
        ("19 Sep",90,88,5,10,0),
        ("20 Sep",95,98,15,20,0),
        ("30 Sep",500,480,50,30,5),
        ("October",1800,1700,50,100,10),
        ("November",1600,1650,20,20,10),
        ("1–14 Dec",800,830,0,10,5),
    ]
    expected_no=[1000,990,985,990,985,987,984,1004,1104,1054,1024]
    expected_base=[1000,915,805,775,765,772,774,769,909,849,824]
    expected_stress=[1000,915,805,775,765,772,774,769,909,549,524]
    for idx,row in enumerate(rows,3):
        period,receipts,payments,settlement,collateral,costs=row
        cash.write(idx,0,period,fmt_label)
        vals=[receipts,payments,settlement,collateral,costs]
        for j,v in enumerate(vals,1): cash.write_number(idx,j,v,fmt_input)
        cash.write_number(idx,10,700,fmt_input)
        if idx==3:
            cash.write_formula(idx,6,"=1000",fmt_num_bold,1000)
            cash.write_formula(idx,7,"=1000",fmt_num_bold,1000)
            cash.write_number(idx,8,0,fmt_input)
            cash.write_formula(idx,9,"=H4+I4",fmt_num_bold,1000)
        else:
            prev=idx-1
            excelrow=idx+1
            cash.write_formula(idx,6,f"=G{excelrow-1}+B{excelrow}-C{excelrow}",fmt_num,expected_no[idx-3])
            cash.write_formula(idx,7,f"=H{excelrow-1}+B{excelrow}-C{excelrow}-D{excelrow}+E{excelrow}-F{excelrow}",fmt_num,expected_base[idx-3])
            stress_adj=-300 if period=="November" else 0
            cash.write_number(idx,8,stress_adj,fmt_input)
            cash.write_formula(idx,9,f"=J{excelrow-1}+B{excelrow}-C{excelrow}-D{excelrow}+E{excelrow}-F{excelrow}+I{excelrow}",fmt_num_bold,expected_stress[idx-3])
    cash.write("A16","Assumptions",fmt_section)
    cash.write("B16","Opening cash",fmt_label); cash.write_number("C16",1000,fmt_input)
    cash.write("D16","Minimum buffer",fmt_label); cash.write_number("E16",700,fmt_input)
    cash.merge_range("F16:K16","Opening cash and buffer are teaching assumptions, not published UBS Treasury data.",fmt_note)
    cash.write("A18","Peak gap",fmt_label)
    cash.write_formula("B18","=MAX(0,E16-MIN(J4:J14))",fmt_red,176)
    cash.write("C18","Peak gap = buffer less minimum stress close",fmt_note)
    # Excel chart
    chart=wb.add_chart({"type":"line"})
    chart.add_series({"name":"No incident","categories":"='Cash Bridge'!$A$4:$A$14","values":"='Cash Bridge'!$G$4:$G$14","line":{"color":teal,"width":2.25},"marker":{"type":"circle","size":4,"border":{"color":teal},"fill":{"color":teal}}})
    chart.add_series({"name":"Incident base","categories":"='Cash Bridge'!$A$4:$A$14","values":"='Cash Bridge'!$H$4:$H$14","line":{"color":"287FA8","width":2.25},"marker":{"type":"circle","size":4,"border":{"color":"287FA8"},"fill":{"color":"287FA8"}}})
    chart.add_series({"name":"Combined stress","categories":"='Cash Bridge'!$A$4:$A$14","values":"='Cash Bridge'!$J$4:$J$14","line":{"color":red,"width":2.5},"marker":{"type":"circle","size":4,"border":{"color":red},"fill":{"color":red}}})
    chart.add_series({"name":"Minimum buffer","categories":"='Cash Bridge'!$A$4:$A$14","values":"='Cash Bridge'!$K$4:$K$14","line":{"color":gold,"width":1.5,"dash_type":"dash"},"marker":{"type":"none"}})
    chart.set_title({"name":"Scenario closing cash vs minimum buffer"})
    chart.set_y_axis({"name":"USD millions","min":500,"max":1200,"major_unit":100,"major_gridlines":{"visible":True,"line":{"color":"D9DEDB"}}})
    chart.set_x_axis({"label_position":"low","num_font":{"size":8}})
    chart.set_legend({"position":"bottom","font":{"size":8}})
    chart.set_size({"width":920,"height":420})
    chart.set_chartarea({"border":{"none":True},"fill":{"color":"FFFFFF"}})
    chart.set_plotarea({"border":{"none":True},"fill":{"color":"FFFFFF"}})
    cash.insert_chart("A20",chart)
    cash.freeze_panes(3,1)

    stress=wb.add_worksheet("Stress Tests")
    stress.set_column("A:A",30); stress.set_column("B:B",17); stress.set_column("C:C",67)
    stress.merge_range("A1:C1","Combined stress · 30 Nov 2011 (illustrative)",fmt_title)
    stress.merge_range("A2:C2","Scenario inputs are assumed; the $2.3bn public loss is not entered as a single cash outflow.",fmt_sub)
    for c,h in enumerate(["Shock","Cash impact (USDm)","Illustrative timing / explanation"]): stress.write(2,c,h,fmt_head)
    shocks=[("Accelerated / increased settlement",-50,"Cash outflow brought forward or increased on 30 Nov."),
            ("Collateral return delayed",-100,"Expected collateral receipt falls beyond the model horizon."),
            ("Additional withdrawals",-100,"Incremental customer outflow in the stress period."),
            ("Funding maturity not renewed",-50,"Scheduled funding rollover unavailable on 30 Nov.")]
    for i,(a,b,c) in enumerate(shocks,3): stress.write(i,0,a,fmt_label); stress.write_number(i,1,b,fmt_input); stress.write(i,2,c,fmt_label)
    stress.write("A8","Total shock",fmt_head); stress.write_formula("B8","=SUM(B4:B7)",fmt_num_bold,-300); stress.write("C8","Additional delayed-liquidity-source timing is tested separately below.",fmt_label)
    stress.write("A10","Incident-base November close",fmt_label); stress.write_number("B10",849,fmt_input)
    stress.write("A11","Stress November close",fmt_label); stress.write_formula("B11","=B10+B8",fmt_red,549)
    stress.write("A12","Buffer",fmt_label); stress.write_number("B12",700,fmt_input)
    stress.write("A13","30 Nov shortfall",fmt_label); stress.write_formula("B13","=MAX(0,B12-B11)",fmt_red,151)
    stress.write("A14","14 Dec stress close",fmt_label); stress.write_number("B14",524,fmt_input)
    stress.write("A15","Peak funding need",fmt_label); stress.write_formula("B15","=MAX(0,B12-B14)",fmt_red,176)
    stress.merge_range("A17:C17","Funding timing test",fmt_section)
    for i,(a,b,c) in enumerate([("Assumed line","USD 200m","Availability / collateral conditions must be verified."),("Trigger","25 Nov","Trigger at least T+2 business days before the modeled breach."),("Illustrative arrival","29 Nov","T+2 assumption; not a historical UBS commitment."),("First breach","30 Nov","If the source arrives later, it cannot prevent the first shortfall.")],18):
        stress.write(i-1,0,a,fmt_label); stress.write(i-1,1,b,fmt_label); stress.write(i-1,2,c,fmt_label)

    risk=wb.add_worksheet("Risk Register")
    risk.set_column("A:A",25); risk.set_column("B:C",15); risk.set_column("D:D",15); risk.set_column("E:E",57)
    risk.merge_range("A1:E1","Residual risk register · illustrative ratings",fmt_title)
    risk.merge_range("A2:E2","Likelihood and impact are analyst assumptions for the 14 Sep–14 Dec horizon; re-score after containment.",fmt_sub)
    for c,h in enumerate(["Risk","Likelihood (1–3)","Impact (1–3)","Score","Why it matters"]): risk.write(2,c,h,fmt_head)
    risk_rows=[("Operational",2,3,"Unreconciled breaks could hide open positions."),
               ("Liquidity",2,3,"Settlement / collateral peak versus approved buffer."),
               ("Market",2,2,"Adverse price movement before positions close."),
               ("Counterparty",2,2,"Collateral return or receivable may be delayed."),
               ("Funding",2,3,"A maturity inside the horizon may not roll."),
               ("Reputational",1,2,"Outflows or tighter terms in a confidence stress."),
               ("Legal / regulatory",2,2,"Investigation and remediation timing / cost.")]
    for i,(name,l,impact,why) in enumerate(risk_rows,3):
        risk.write(i,0,name,fmt_label); risk.write_number(i,1,l,fmt_input); risk.write_number(i,2,impact,fmt_input)
        risk.write_formula(i,3,f"=B{i+1}*C{i+1}",fmt_num_bold,l*impact); risk.write(i,4,why,fmt_label)
    risk.merge_range("A12:E12","Scoring key",fmt_section)
    risk.merge_range("A13:E13","Likelihood: 1 unlikely · 2 possible · 3 likely. Impact: 1 headroom erosion <25% · 2 ≥25% without breach · 3 cash below buffer. The 25% cutoff is proposed, not UBS policy.",fmt_note)

    actions=wb.add_worksheet("CFO Actions")
    actions.set_column("A:A",29); actions.set_column("B:B",30); actions.set_column("C:C",29); actions.set_column("D:D",48); actions.set_column("E:E",45)
    actions.merge_range("A1:E1","CFO action plan · owners, timing and decision output",fmt_title)
    for c,h in enumerate(["Pillar","Owner","Deadline","Expected benefit / deliverable","Constraint / fallback"]): actions.write(2,c,h,fmt_head)
    action_rows=[
        ("Immediate liquidity protection","Group Treasurer","Initial view in 24h; daily to 14 Dec","Verified accessible cash, dated 13-week forecast and pre-positioned funding.","Confirm facility access, collateral eligibility and arrival; fallback only if independently verified."),
        ("Exposure containment","IB COO + Market Risk + Treasury","Start now; daily until reconciled","Independently reconcile open positions, settlements and collateral breaks.","Close-out cost / cash release only when confirmed; keep conservative overlay until then."),
        ("Governance & controls","CRO + IB COO; Internal Audit validates","Interim controls now; programme by 14 Dec","Dual verification, clear ownership and timely escalation of exceptions.","Manual dual sign-off until system controls are remediated."),
    ]
    for i,row in enumerate(action_rows,3):
        for c,val in enumerate(row): actions.write(i,c,val,fmt_label)
        actions.set_row(i,54)
    actions.freeze_panes(3,0)
    wb.close()


def make_deck():
    prs=Presentation()
    prs.slide_width=Inches(SW); prs.slide_height=Inches(SH)
    blank=prs.slide_layouts[6]

    # 1–2: the user's approved slides remain intact as single picture objects.
    add_user_slide(prs,MAP_SLIDE)
    add_user_slide(prs,NORMAL_HIDDEN_SLIDE)

    # 3 — The causal bridge between operational failure and Treasury.
    s=prs.slides.add_slide(blank)
    header(s,"FROM CONTROL FAILURE TO TREASURY IMPACT",
           "SECTION 1.3  ·  FAILED CONTROLS & CASH TRANSMISSION",3,
           "THE BRIDGE FROM OPERATIONAL RISK TO THE CENTRAL TREASURY QUESTION")
    txt(s,"Control failure",.62,1.30,2.20,.22,9,RED,True)
    txt(s,"How can UBS preserve liquidity and meet payment obligations over the next three months?",
        3.05,1.29,9.50,.24,10,NAVY,True)
    steps=[
        ("CONTROL\nFAILURES","verification gaps",RED),
        ("HIDDEN\nPOSITIONS","false offsets",GOLD),
        ("TRADING\nLOSS","USD 2.3bn",RED),
        ("SETTLEMENT\nOBLIGATIONS","dated cash due",BLUE),
        ("COLLATERAL\nCALLS","calls / delayed returns",GOLD),
        ("LIQUIDITY\nPRESSURE","buffer erosion",RED),
        ("FUNDING\nREQUIREMENT","timing gap",TEAL),
    ]
    x0,w,gap,y,h=.52,1.48,.32,2.02,2.0
    for i,(name,detail,color) in enumerate(steps):
        x=x0+i*(w+gap)
        rect(s,x,y,w,h,WHITE,GRID,0.8)
        rect(s,x,y,w,.055,color)
        txt(s,f"{i+1:02d}",x+.12,y+.17,.36,.22,8,color,True)
        txt(s,name,x+.12,y+.55,w-.24,.54,9,NAVY,True,align=PP_ALIGN.CENTER)
        line(s,x+.14,y+1.29,x+w-.14,y+1.29,GRID,.65)
        txt(s,detail,x+.11,y+1.43,w-.22,.28,7.4,MUTED,align=PP_ALIGN.CENTER)
        if i<len(steps)-1:
            rect(s,x+w+.08,y+.83,gap-.13,.20,color,shape=MSO_SHAPE.CHEVRON)
    # A clear three-stage grouping, not a repeated panel stack.
    groups=[(.62,3.65,"OPERATIONAL EVENT","control failure → hidden positions",RED),
            (4.83,3.65,"CASH TRANSMISSION","loss → settlement / collateral",BLUE),
            (9.04,3.65,"TREASURY RESPONSE","liquidity pressure → funding",TEAL)]
    for x,ww,title,sub,color in groups:
        txt(s,title,x,4.53,ww,.20,8,color,True,align=PP_ALIGN.CENTER)
        line(s,x,4.81,x+ww,4.81,color,1.8)
        txt(s,sub,x,4.94,ww,.25,8.1,NAVY,True,align=PP_ALIGN.CENTER)
    rect(s,.62,5.62,12.02,.60,"E9ECE8")
    txt(s,"The USD 2.3bn headline is a loss figure; the Treasury need depends on when cash is due, recoverable collateral and accessible liquidity.",
        .82,5.77,11.62,.28,9.5,NAVY,True,align=PP_ALIGN.CENTER)
    txt(s,"Potential cash channels are analytical, not asserted historical UBS cash movements. Sources: FINMA 2012 §§23–41; FSA 2012 §§26–42.",
        .64,6.55,12.0,.22,7.0,MUTED)

    # 4 — Timeline with sourced event sequence and requested humorous illustration.
    s=prs.slides.add_slide(blank)
    header(s,"INCIDENT TIMELINE","SECTION 1.2  ·  DISCOVERY, ANNOUNCEMENT & RESPONSE",4,
           "DATES ARE MILESTONES; CONTAINMENT AND REMEDIATION CONTINUE AFTER DISCOVERY")
    # Main timeline occupies left ~9 inches; sourced milestones are explicit.
    line(s,.78,2.05,9.28,2.05,GRID,2.2)
    events=[
        (.72,RED,"14 SEP 2011","Discovery","Unauthorised futures positions identified."),
        (2.91,GOLD,"15 SEP 2011","Public announcement","Initial loss estimate communicated."),
        (5.10,BLUE,"FOLLOWING DAYS","Containment & hedging","Positions managed; exposure reviewed."),
        (7.29,TEAL,"FOLLOWING WEEKS","Investigation & remediation","Control failures examined."),
    ]
    for x,color,date,title,desc in events:
        mid=x+.92
        rect(s,mid-.10,1.95,.20,.20,color,shape=MSO_SHAPE.OVAL)
        txt(s,date,x,2.31,1.86,.23,8,color,True)
        txt(s,title,x,2.64,1.86,.42,9.2,NAVY,True)
        txt(s,desc,x,3.13,1.86,.56,7.8,MUTED,valign=MSO_ANCHOR.TOP)
    # Cartoon replaces an archival trader portrait; fitted without distortion.
    rect(s,9.72,1.47,3.12,2.08,WHITE,GRID,.8)
    add_img_fit(s,DALTONS,9.81,1.56,2.94,1.65)
    txt(s,"ILLUSTRATIVE CARTOON",9.84,3.28,2.85,.17,6.8,MUTED,True,align=PP_ALIGN.RIGHT)
    # timeline interpretation / source line
    line(s,.68,4.12,12.68,4.12,GRID,.8)
    txt(s,"TREASURY LENS",.72,4.39,1.55,.20,8,RED,True)
    txt(s,"Discovery establishes the control event; it does not by itself identify the dated cash outflow or funding gap.",
        2.42,4.34,9.94,.33,10,NAVY,True)
    # Small date strip helps M1/M2 presenters anchor the 3-month horizon.
    rect(s,.72,5.12,11.96,.88,PALE_AMBER)
    txt(s,"DECISION WINDOW",.95,5.28,1.72,.20,7.8,GOLD,True)
    txt(s,"14 September",2.72,5.25,2.05,.24,9,NAVY,True)
    line(s,4.72,5.39,9.50,5.39,GOLD,2)
    rect(s,4.62,5.29,.20,.20,GOLD,shape=MSO_SHAPE.OVAL)
    rect(s,9.40,5.29,.20,.20,RED,shape=MSO_SHAPE.OVAL)
    txt(s,"3-month Treasury response horizon",5.10,5.48,3.80,.18,7.2,MUTED,align=PP_ALIGN.CENTER)
    txt(s,"14 December",9.78,5.25,1.96,.24,9,NAVY,True)
    txt(s,"Sources: FINMA 2012; FSA 2012; UBS disclosures. Cartoon is illustrative, not an incident photograph.",
        .72,6.53,11.95,.20,7,MUTED)

    # 5 — Native editable PowerPoint line chart, same data as the Excel model.
    s=prs.slides.add_slide(blank)
    header(s,"TREASURY LIQUIDITY EVOLUTION","SECTION 2.2  ·  CASH BRIDGE",5,
           "ILLUSTRATIVE MODEL · USDm · CASH PATHS ARE NOT HISTORICAL UBS TREASURY DATA")
    txt(s,"Cash balance by scenario",.60,1.24,6.0,.24,10,NAVY,True)
    cats=CategoryChartData(); cats.categories=DATES
    cats.add_series("No incident",NO_INCIDENT)
    cats.add_series("Incident base",INCIDENT_BASE)
    cats.add_series("Combined stress",COMBINED_STRESS)
    cats.add_series("Minimum buffer",[BUFFER]*len(DATES))
    chart=s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(.55), Inches(1.53),
                             Inches(9.2), Inches(4.98), cats).chart
    chart.has_title=False
    chart.has_legend=True
    chart.legend.position=XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout=False
    chart.legend.font.name=FONT
    chart.legend.font.size=Pt(8)
    plot=chart.plots[0]
    plot.has_data_labels=False
    va=chart.value_axis
    va.minimum_scale=500; va.maximum_scale=1200; va.major_unit=100
    va.has_major_gridlines=True
    va.tick_labels.font.name=FONT; va.tick_labels.font.size=Pt(8)
    va.tick_labels.font.color.rgb=C(MUTED)
    ca=chart.category_axis
    ca.tick_labels.font.name=FONT; ca.tick_labels.font.size=Pt(8)
    ca.tick_labels.font.color.rgb=C(MUTED)
    ca.major_tick_mark=2
    colors=[TEAL,BLUE,RED,GOLD]
    for i,ser in enumerate(chart.series):
        ser.format.line.color.rgb=C(colors[i])
        ser.format.line.width=Pt(2 if i<3 else 1.4)
        if i<3:
            ser.marker.style=2
            ser.marker.size=5
            ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb=C(colors[i])
            ser.marker.format.line.color.rgb=C(colors[i])
        else:
            ser.marker.style=1
            ser.format.line.dash_style=4
    # Key results, in a compact analyst-style table rather than decorative cards.
    rect(s,9.95,1.56,2.88,4.88,WHITE,GRID,.8)
    rect(s,9.95,1.56,2.88,.055,RED)
    txt(s,"KEY MODEL OUTPUTS",10.14,1.79,2.50,.22,8.4,NAVY,True)
    rows=[("Opening cash*","$1,000m",TEAL),
          ("Minimum buffer*","$700m",GOLD),
          ("Incident-base minimum","$765m",BLUE),
          ("First stress breach","30 Nov · $549m",RED),
          ("Peak stress funding","$176m · 14 Dec",RED)]
    for i,(label,value,color) in enumerate(rows):
        yy=2.19+i*.69
        txt(s,label,10.14,yy,1.50,.22,7.1,MUTED)
        txt(s,value,11.56,yy-.02,1.10,.30,8.6,color,True,align=PP_ALIGN.RIGHT)
        if i<4: line(s,10.14,yy+.42,12.64,yy+.42,GRID,.6)
    txt(s,"*Teaching assumptions. Base case stays above the buffer; combined stress falls below it.",
        10.14,5.82,2.48,.43,7.0,MUTED)
    txt(s,"Formula: closing cash = opening + receipts − payments − settlement outflow + signed collateral movement − costs. Funding gap = max(0, buffer − cash before funding).",
        .64,6.70,12.0,.28,7.0,MUTED)

    # 6 — A transparent stress table and funding-arrival test.
    s=prs.slides.add_slide(blank)
    header(s,"STRESS TESTS","SECTION 2.3  ·  CASH SHOCKS & FUNDING TIMING",6,
           "COMBINED SCENARIO INPUTS ARE ASSUMPTIONS, NOT HISTORICAL OUTCOMES")
    txt(s,"Combined shock to November incident-base cash",.62,1.22,7.2,.22,10,NAVY,True)
    # Spreadsheet-like data table
    x0,y0= .62,1.63
    widths=[3.05,1.42,4.38]
    headers=["Stress input","Cash impact\n(USDm)","What changes in the scenario"]
    xx=x0
    for i,(head,wid) in enumerate(zip(headers,widths)):
        rect(s,xx,y0,wid,.55,NAVY,NAVY)
        txt(s,head,xx+.10,y0+.07,wid-.20,.40,8.0,WHITE,True,
            align=PP_ALIGN.RIGHT if i==1 else PP_ALIGN.LEFT)
        xx+=wid
    stress_rows=[
        ("Settlement accelerated / increased",-50,"More cash is due on 30 Nov."),
        ("Collateral return delayed",-100,"Expected return shifts beyond the horizon."),
        ("Additional withdrawals",-100,"Incremental customer outflow on 30 Nov."),
        ("Funding rollover not renewed",-50,"A maturing source is unavailable."),
    ]
    for r,row in enumerate(stress_rows):
        yy=y0+.55+r*.58
        fill=WHITE if r%2==0 else "F0F2EF"
        xx=x0
        for i,(val,wid) in enumerate(zip(row,widths)):
            rect(s,xx,yy,wid,.58,fill,GRID,.55)
            txt(s,f"{val:,}" if i==1 else str(val),xx+.10,yy+.08,wid-.20,.40,
                8.2,RED if i==1 else NAVY,i==1,align=PP_ALIGN.RIGHT if i==1 else PP_ALIGN.LEFT)
            xx+=wid
    yy=y0+.55+4*.58
    xx=x0
    total=["TOTAL CASH SHOCK",-300,"Applied to assumed November incident-base close of $849m → $549m." ]
    for i,(val,wid) in enumerate(zip(total,widths)):
        rect(s,xx,yy,wid,.58,PALE_RED,GRID,.55)
        txt(s,f"{val:,}" if i==1 else str(val),xx+.10,yy+.08,wid-.20,.40,
            8.5,RED,True,align=PP_ALIGN.RIGHT if i==1 else PP_ALIGN.LEFT)
        xx+=wid
    # Funding timing table / chronology
    rect(s,10.05,1.63,2.76,3.45,PALE_AMBER,GRID,.8)
    txt(s,"LIQUIDITY SOURCE TEST",10.23,1.84,2.40,.27,8.5,NAVY,True)
    txt(s,"Assumed line",10.23,2.31,1.10,.18,7.2,MUTED)
    txt(s,"$200m · T+2",11.27,2.29,1.20,.22,8.3,NAVY,True,align=PP_ALIGN.RIGHT)
    line(s,10.23,2.62,12.62,2.62,GRID,.7)
    for i,(date,meaning,color) in enumerate([("25 Nov","Trigger",GOLD),("29 Nov","Illustrative arrival",TEAL),("30 Nov","First buffer breach",RED)]):
        yy=2.85+i*.58
        rect(s,10.26,yy+.05,.18,.18,color,shape=MSO_SHAPE.OVAL)
        txt(s,date,10.57,yy, .82,.23,8,NAVY,True)
        txt(s,meaning,11.39,yy,1.20,.25,7.1,MUTED,align=PP_ALIGN.RIGHT)
    txt(s,"Late arrival cannot prevent the first shortfall.",10.23,4.65,2.40,.27,7.5,RED,True)
    # Output strip
    line(s,.62,5.47,12.81,5.47,GRID,.8)
    outputs=[("30 NOV","$549m cash","$151m below buffer",RED),
             ("14 DEC","$524m cash","$176m peak need",RED),
             ("MITIGATION","Trigger by 25 Nov","Verify source / collateral",TEAL)]
    for i,(date,value,note,color) in enumerate(outputs):
        x=.70+i*4.12
        txt(s,date,x,5.70,1.25,.18,7.4,color,True)
        txt(s,value,x,5.98,2.20,.28,11,NAVY,True)
        txt(s,note,x+2.10,5.99,1.84,.28,7.2,MUTED,align=PP_ALIGN.RIGHT)
        if i<2: line(s,x+3.82,5.70,x+3.82,6.43,GRID,.7)
    txt(s,"Assumed facility timing and stress inputs must be validated. A delayed funding source is tested as a timing failure, not counted as cash.",
        .64,6.70,12.0,.24,7.0,MUTED)

    # 7 — Risk matrix in conventional likelihood / impact orientation.
    s=prs.slides.add_slide(blank)
    header(s,"RISK MATRIX","SECTION 3.1  ·  LIKELIHOOD × IMPACT",7,
           "RESIDUAL RATINGS ARE ANALYST JUDGEMENTS FOR THE THREE-MONTH HORIZON")
    mx,my,cell=.98,1.72,1.34
    # high likelihood at top; high impact at right.
    fills=[["F7E4E3","F3D7D1","EDC4C4"],
           ["F5EDDA","EFE2BC","E9D3B0"],
           ["E5EEE6","E9EDDF","F3E9D2"]]
    for r in range(3):
        for c in range(3):
            rect(s,mx+c*cell,my+r*cell,cell-.04,cell-.04,fills[r][c],WHITE,1.0)
            score=(3-r)*(c+1)
            txt(s,str(score),mx+c*cell+.07,my+r*cell+.05,.18,.15,6.8,MUTED)
    for r,label in enumerate(["LIKELY","POSSIBLE","UNLIKELY"]):
        txt(s,label,.16,my+r*cell+.48,.73,.20,7.1,NAVY,True,align=PP_ALIGN.RIGHT)
    for c,label in enumerate(["LOW","MEDIUM","HIGH"]):
        txt(s,label,mx+c*cell,my+3*cell+.02,cell-.04,.19,7.0,NAVY,True,align=PP_ALIGN.CENTER)
    txt(s,"LIKELIHOOD",.18,1.43,.72,.18,6.8,MUTED,True,align=PP_ALIGN.RIGHT)
    txt(s,"IMPACT",mx+.91,my+3*cell+.28,1.20,.18,7.0,MUTED,True,align=PP_ALIGN.CENTER)
    # risk markers, codes match the editable register to the right
    markers=[(2,3,"O",RED),(2,3,"L",BLUE),(2,3,"F",RED),
             (2,2,"M",GOLD),(2,2,"C",TEAL),(2,2,"G",GOLD),(1,2,"R",GREEN)]
    seen={}
    offsets={(2,3):[(-.27,-.18),(0,-.18),(.27,-.18)],
             (2,2):[(-.27,.17),(0,.17),(.27,.17)],(1,2):[(0,0)]}
    for lik,impact,code,color in markers:
        r=3-lik;c=impact-1;key=(lik,impact)
        dx,dy=offsets[key][seen.get(key,0)];seen[key]=seen.get(key,0)+1
        cx=mx+c*cell+.65+dx;cy=my+r*cell+.65+dy
        rect(s,cx-.12,cy-.12,.24,.24,color,shape=MSO_SHAPE.OVAL)
        txt(s,code,cx-.12,cy-.08,.24,.15,6.5,WHITE,True,align=PP_ALIGN.CENTER)
    # Excel-like risk register table
    tx,ty=5.35,1.62
    colw=[1.80,.83,.75,.78,3.25]
    heads=["Risk","Likelihood","Impact","Score","Why it matters"]
    xx=tx
    for h,wid in zip(heads,colw):
        rect(s,xx,ty,wid,.47,NAVY,NAVY)
        txt(s,h,xx+.08,ty+.06,wid-.16,.34,7.4,WHITE,True,
            align=PP_ALIGN.CENTER if h!="Why it matters" else PP_ALIGN.LEFT)
        xx+=wid
    riskrows=[("Operational","2","3","6","Open breaks can hide positions."),
              ("Liquidity","2","3","6","Settlement / collateral vs buffer."),
              ("Market","2","2","4","Adverse move before close-out."),
              ("Counterparty","2","2","4","Delayed collateral / receivable."),
              ("Funding","2","3","6","Maturity may not roll."),
              ("Reputational","1","2","2","Outflows or tighter terms."),
              ("Legal / regulatory","2","2","4","Investigation / remediation." )]
    for r,row in enumerate(riskrows):
        yy=ty+.47+r*.48;xx=tx
        for j,(val,wid) in enumerate(zip(row,colw)):
            fill=WHITE if r%2==0 else "F0F2EF"
            rect(s,xx,yy,wid,.48,fill,GRID,.5)
            color=RED if j==3 and int(val)>=6 else NAVY
            txt(s,val,xx+.08,yy+.06,wid-.16,.34,7.4,color,j==3,
                align=PP_ALIGN.CENTER if j in (1,2,3) else PP_ALIGN.LEFT)
            xx+=wid
    txt(s,"Priorities: operational control, liquidity-buffer breach and funding rollover. Re-score after close-out.",
        5.57,5.68,7.0,.28,8.0,RED,True)
    txt(s,"Illustrative ratings, not UBS / regulator scores. L=likelihood; I=impact; score=L×I. Define impact using the approved buffer; the 25% threshold is proposed, not UBS policy.",
        .62,6.61,12.0,.32,7.0,MUTED)

    # 8 — CFO action plan as a compact decision table, not a card stack.
    s=prs.slides.add_slide(blank)
    header(s,"CFO ACTION PLAN","SECTION 3.2  ·  THREE PILLARS",8,
           "OWNERS · DEADLINES · EXPECTED BENEFITS · PRACTICAL FALLBACKS")
    txt(s,"Decision rule: preserve payment capacity; activate funding before a verified forecast shows a buffer breach.",
        .64,1.20,12.0,.28,10,NAVY,True)
    x0,y0=.62,1.72
    widths=[2.00,2.00,1.55,3.60,2.85]
    heads=["PILLAR","ACCOUNTABLE OWNER","DEADLINE","DELIVERABLE / BENEFIT","CONSTRAINT / FALLBACK"]
    xx=x0
    for title,wid in zip(heads,widths):
        rect(s,xx,y0,wid,.56,NAVY,NAVY)
        txt(s,title,xx+.10,y0+.08,wid-.20,.40,7.4,WHITE,True)
        xx+=wid
    actions=[
        ("01 · Liquidity\nprotection","Group Treasurer","Initial view\nwithin 24h","Dated 13-week cash forecast; verify accessible headroom and stage funding ahead of breach.","Confirm availability, collateral eligibility, lead time and cost; use only verified fallback."),
        ("02 · Exposure\ncontainment","IB COO + Market Risk\n+ Treasury","Immediate; daily\nuntil reconciled","Independent position match; resolve settlement and collateral breaks. Benefit: bound uncertainty.","Close-out cost / cash release only when confirmed; keep conservative overlay until then."),
        ("03 · Governance\n& controls","CRO + IB COO;\nInternal Audit validates","Interim now;\nprogramme by 14 Dec","Dual verification, named break ownership and escalation. Benefit: detect exceptions earlier.","Manual dual sign-off until system controls and access are remediated."),
    ]
    heights=[1.15,1.15,1.15]
    yy=y0+.56
    for r,row in enumerate(actions):
        hh=heights[r];xx=x0;fill=WHITE if r%2==0 else "F0F2EF"
        for j,(value,wid) in enumerate(zip(row,widths)):
            rect(s,xx,yy,wid,hh,fill,GRID,.65)
            color=[RED,BLUE,TEAL][r] if j==0 else NAVY
            txt(s,value,xx+.11,yy+.10,wid-.22,hh-.20,7.7 if j else 8.1,color,j==0,
                valign=MSO_ANCHOR.TOP)
            xx+=wid
        yy+=hh
    # 3-step monitoring strip for presenters
    txt(s,"MONITOR",.65,6.10,1.05,.20,7.5,RED,True)
    checks=[("Accessible cash vs buffer",1.92), ("Unmatched trades / collateral",5.42), ("Funding arrival confirmed",9.15)]
    for label,x in checks:
        rect(s,x,6.02,.16,.16,TEAL,shape=MSO_SHAPE.OVAL)
        txt(s,label,x+.25,5.97,2.92,.26,7.3,NAVY,True)
    txt(s,"Use internal Treasury records for real balances; public sources do not establish UBS’s actual daily funding gap.",
        .64,6.55,12.0,.25,7.0,MUTED)

    prs.core_properties.title="UBS 2011 Rogue Trading · Treasury Response"
    prs.core_properties.subject="Eight-slide, report-aligned UBS Treasury presentation; Excel/Python-style model outputs"
    prs.core_properties.author="Corporate Treasury Management Team"
    prs.save(PPTX_OUT)
    return PPTX_OUT


if __name__=="__main__":
    build_model_workbook()
    print(make_deck())
    print(XLSX_OUT)
